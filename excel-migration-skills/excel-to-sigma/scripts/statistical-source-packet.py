#!/usr/bin/env python3
"""Extract the reviewed Statistical Demand Forecast inputs without changing Excel."""
from __future__ import annotations

import argparse
import csv
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import sys

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent / "lib"))
from statistical_forecast import SOURCE_SHA256, SKUS, calculate  # noqa: E402


def _date(value):
    if not isinstance(value, (date, datetime)):
        raise ValueError("SOURCE_DATE_MISSING")
    return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()


def extract(source):
    source = Path(source)
    if hashlib.sha256(source.read_bytes()).hexdigest() != SOURCE_SHA256:
        raise ValueError("SOURCE_HASH_MISMATCH")
    book = load_workbook(source, read_only=True, data_only=False)
    history = []
    for row in range(2, 38):
        month = _date(book["History"][f"A{row}"].value)
        for offset, sku in enumerate(SKUS, 2):
            cell = book["History"].cell(row, offset)
            value = cell.value
            if cell.data_type == "f" or isinstance(value, bool) or not isinstance(value, (int, float)):
                raise ValueError("SOURCE_HISTORY_LITERAL_MISSING")
            history.append({"Month": month, "SKU": sku, "Demand": value,
                            "sourceCells": {"Month": f"History!A{row}",
                                            "SKU": f"History!{cell.column_letter}1",
                                            "Demand": f"History!{cell.coordinate}"}})
    labels = (("Selected SKU", "SelectedSKU"), ("Seasonality (months)", "Seasonality"),
              ("Service level", "ServiceLevel"), ("Lead time (months)", "LeadTime"))
    parameters = {"Scope": book["Params"]["A1"].value}
    parameter_cells = {"Scope": "Params!A1"}
    for row, (label, field) in enumerate(labels, 3):
        if book["Params"][f"A{row}"].value != label:
            raise ValueError("SOURCE_PARAMETER_LABEL_MISMATCH")
        cell = book["Params"][f"B{row}"]
        if cell.data_type == "f" or cell.value is None:
            raise ValueError("SOURCE_PARAMETER_LITERAL_MISSING")
        parameters[field] = cell.value
        parameter_cells[field] = f"Params!B{row}"
    if parameters["Scope"] != "Parameters":
        raise ValueError("SOURCE_PARAMETER_SCOPE_MISMATCH")
    parameters["sourceCells"] = parameter_cells
    targets = []
    for period, row in enumerate(range(4, 16), 1):
        cell = book["Forecast"][f"A{row}"]
        if cell.data_type == "f":
            raise ValueError("SOURCE_TARGET_LITERAL_MISSING")
        targets.append({"Period": period, "Month": _date(cell.value),
                        "sourceCells": {"Month": f"Forecast!A{row}"},
                        "derivedFrom": {"Period": f"ordinal position of Forecast!A{row}"}})
    values = {role: [{key: value for key, value in row.items() if key not in ("sourceCells", "derivedFrom")}
                     for row in records]
              for role, records in (("history", history), ("parameters", [parameters]), ("targets", targets))}
    check = calculate(values["history"], parameters["SelectedSKU"],
                      parameters["Seasonality"], parameters["ServiceLevel"],
                      parameters["LeadTime"], [row["Month"] for row in values["targets"]])
    if check["Status"] != "ETS_UNAVAILABLE" or len(check["Periods"]) != 12:
        raise ValueError("SOURCE_VALIDATION_FAILED")
    return {"sourceHash": SOURCE_SHA256, "modelDecision": "CANDIDATE_REVIEW_ONLY",
            "roles": {"history": {"naturalKey": ["Month", "SKU"], "columns": {"Month": "datetime", "SKU": "text", "Demand": "number"}, "rows": history},
                      "parameters": {"naturalKey": ["Scope"], "columns": {"Scope": "text", "SelectedSKU": "text", "Seasonality": "number", "ServiceLevel": "number", "LeadTime": "number"}, "rows": [parameters]},
                      "targets": {"naturalKey": ["Period"], "columns": {"Period": "number", "Month": "datetime"}, "rows": targets}}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    packet = extract(args.source)
    args.out.mkdir(parents=True, exist_ok=True)
    (args.out / "source-inputs.json").write_text(json.dumps(packet, indent=2) + "\n")
    for role, spec in packet["roles"].items():
        with (args.out / f"{role}.csv").open("w", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(spec["columns"]))
            writer.writeheader()
            writer.writerows({name: row[name] for name in spec["columns"]} for row in spec["rows"])
    print(json.dumps({"sourceHash": packet["sourceHash"],
                      "rows": {role: len(spec["rows"]) for role, spec in packet["roles"].items()},
                      "out": str(args.out)}))


if __name__ == "__main__":
    main()
