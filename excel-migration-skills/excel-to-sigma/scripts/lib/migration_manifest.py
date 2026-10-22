"""Versioned, source-bound migration manifests and evidence gates."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

try:
    from .excel_expression import ExpressionError, extract_references, parse_expression
    from .formula_graph import FormulaGraph
except ImportError:  # standalone tests/scripts that place lib directly on sys.path
    from excel_expression import ExpressionError, extract_references, parse_expression
    from formula_graph import FormulaGraph


MANIFEST_VERSION = 1


class ManifestError(ValueError):
    pass


_ROOT_EXCLUDED = {"source", "sourcePath", "sourceSha256", "migrationManifest",
                  "evidence", "runtimeEvidence", "execution", "executionState",
                  "verification", "targetVersions", "targetVersion", "targetBinding",
                  "actualTarget", "runtime", "formulaCells", "sourceReviewEvidence"}
_RECORD_RUNTIME_FIELDS = {"evidence", "runtimeEvidence", "execution", "executionState",
                          "verification", "targetVersions", "targetVersion", "targetBinding",
                          "actualTarget", "runtime"}
_LOGICAL_TARGET_FIELDS = {"logicalElementId", "logicalColumnId", "logicalLineId",
                          "logicalCheckId", "logicalSourceFeatureId", "year",
                          "equivalentReference", "logicalLaneColumnIds"}
_AMBIGUOUS_TARGET_FIELDS = {"workbookId", "documentId", "elementId", "columnId", "lineId",
                            "checkId", "sourceFeatureId", "actualColumnId", "version"}


def _sha256(path):
    with Path(path).open("rb") as stream:
        return hashlib.sha256(stream.read()).hexdigest()


def _formula_projection(formulas):
    keys = ("id", "sheet", "address", "originalFormula", "dependencies",
            "mappingStatus", "targetMapping", "unresolvedReason", "reviewedReplacement", "nativeCandidate")
    return [{key: item.get(key) for key in keys if key in item}
            for item in sorted(formulas, key=lambda value: value.get("id", ""))]


def _plan_projection(plan):
    result = {}
    for key, value in sorted((plan or {}).items()):
        if key in _ROOT_EXCLUDED:
            continue
        if key in {"lines", "features", "auditChecks"} and isinstance(value, list):
            result[key] = [_record_projection(item) for item in value]
        elif key == "provenance" and isinstance(value, dict):
            result[key] = {name: _record_projection(item) for name, item in sorted(value.items())}
        else:
            result[key] = value
    return result


def _record_projection(record):
    if not isinstance(record, dict):
        return record
    _validate_target_mapping(record.get("targetMapping"), record.get("id", "plan record"))
    result = {}
    for key, value in sorted(record.items()):
        if key in _RECORD_RUNTIME_FIELDS:
            continue
        if key == "provenance" and isinstance(value, dict):
            result[key] = {name: _record_projection(item) for name, item in sorted(value.items())}
        else:
            result[key] = value
    return result


def _validate_target_mapping(mapping, formula_id):
    if mapping is None:
        return
    if not isinstance(mapping, dict):
        raise ManifestError(f"logical targetMapping must be an object for {formula_id}")
    lanes = mapping.get("logicalLaneColumnIds")
    if lanes is not None and (not isinstance(lanes, dict)
            or set(lanes) != {"number", "text", "logical", "kind", "error"}
            or any(not isinstance(value, str) or not value.strip() for value in lanes.values())):
        raise ManifestError(f"logical native lane identities are incomplete for {formula_id}")
    misplaced = sorted(set(mapping) & _AMBIGUOUS_TARGET_FIELDS)
    unknown = sorted(set(mapping) - _LOGICAL_TARGET_FIELDS)
    if misplaced or unknown:
        fields = ", ".join(sorted(set(misplaced + unknown)))
        raise ManifestError(
            f"migration targetMapping for {formula_id} contains ambiguous or runtime fields: "
            f"{fields}; use logical* keys for approved intent and targetBinding for assigned IDs"
        )


def _valid_reviewed_replacement(record, original_formula):
    replacement = record.get("reviewedReplacement")
    return (
        isinstance(original_formula, str)
        and record.get("mappingStatus") == "live"
        and isinstance(record.get("targetMapping"), dict)
        and isinstance(replacement, dict)
        and replacement.get("originalFormula") == original_formula
        and isinstance(replacement.get("replacementFormula"), str)
        and bool(replacement["replacementFormula"].strip())
        and isinstance(replacement.get("reviewer"), str)
        and bool(replacement["reviewer"].strip())
        and replacement.get("reviewStatus") == "approved"
    )


def _plan_hash(source_hash, projection, formulas):
    logical = {"sourceHash": source_hash, "plan": projection,
               "formulaCells": _formula_projection(formulas)}
    encoded = json.dumps(logical, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _native_records(source_hash, plan, source_path=None):
    """Opt-in reviewed compiler evidence; this is never runtime verification."""
    native = (plan or {}).get("nativeCalculationPlan")
    if native is None:
        return {}
    approval = (plan or {}).get("nativeReviewApproval") or {}
    canonical = {key: value for key, value in native.items() if key != "nativePlanHash"}
    digest = hashlib.sha256(json.dumps(canonical, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()).hexdigest()
    if (native.get("nativePlanVersion") != 2 or native.get("nativePlanHash") != digest
            or native.get("sourceSha256") != source_hash
            or approval.get("status") != "reviewed"
            or not isinstance(approval.get("owner"), str) or not approval["owner"].strip()
            or not isinstance(approval.get("reviewer"), str) or not approval["reviewer"].strip()
            or approval.get("sourceSha256") != source_hash
            or approval.get("nativePlanHash") != digest
            or native.get("inputPlanHash") != (plan or {}).get("inputPlanHash")):
        raise ManifestError("native bound-AST source/plan/review hash is stale")
    if source_path is not None:
        try:
            from .workbook_expression import WorkbookContext
            from .native_calculations import _ast
        except ImportError:
            from workbook_expression import WorkbookContext
            from native_calculations import _ast
        context = WorkbookContext.from_path(source_path)
        expected_cells = {f"{sheet}!{address}" for sheet, address in context.formulas}
        if set(native.get("formulas") or {}) != expected_cells:
            raise ManifestError("native evidence does not cover the exact source formula denominator")
    result = {}
    for cell, record in native.get("formulas", {}).items():
        if record.get("status") != "ready-to-bind":
            continue
        formula = record.get("originalFormula")
        ast = record.get("ast")
        ast_hash = hashlib.sha256(json.dumps(ast, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()).hexdigest()
        if (record.get("runtimeStatus") != "UNVERIFIED_CANDIDATE"
                or record.get("astHash") != ast_hash
                or not isinstance(formula, str)
                or record.get("formulaSha256") != hashlib.sha256(formula.encode()).hexdigest()):
            raise ManifestError("native bound-AST formula evidence is stale")
        if source_path is not None:
            sheet, address = cell.rsplit("!", 1)
            if context.formulas.get((sheet, address)) != formula:
                raise ManifestError("native evidence original source formula changed")
            try:
                rebound = context.bind_formula(formula, sheet=sheet, address=address)
                if _ast(rebound) != ast or not context._bound_evidence(rebound)[2]:
                    raise ManifestError("native bound-AST differs from source binding")
            except (ValueError, TypeError) as error:
                raise ManifestError("native source AST cannot be rebound") from error
        result[cell] = record
    return result


def build_manifest(source_path, plan=None) -> dict:
    source = Path(source_path).expanduser().resolve()
    if not source.is_file():
        raise ManifestError(f"source workbook not found: {source}")
    source_hash = _sha256(source)
    workbook = load_workbook(source, data_only=False, read_only=False)
    cached = load_workbook(source, data_only=True, read_only=False)
    graph = FormulaGraph.from_workbook(workbook)
    native_records = _native_records(source_hash, plan, source)
    cycle_members = {cell for component in graph.cycles() for cell in component}
    plan_cells = _plan_cell_index(plan or {})
    reviewed_roots = {
        identity for identity, formula in graph.formulas.items()
        if identity.id in native_records or identity in graph.unsupported_formulas
        and _valid_reviewed_replacement(plan_cells.get(identity.id, {}), formula)
    }
    unresolved_nodes = (
        set(graph.unresolved_references) | set(graph.unresolved_names)
        | (set(graph.unsupported_formulas) - reviewed_roots) | cycle_members)
    changed = True
    while changed:
        changed = False
        for node, dependencies in graph.dependencies.items():
            if node not in unresolved_nodes and any(dep in unresolved_nodes for dep in dependencies):
                unresolved_nodes.add(node)
                changed = True
    formulas = []
    for identity in sorted(graph.formulas):
        formula = graph.formulas[identity]
        record = plan_cells.get(identity.id, {})
        _validate_target_mapping(record.get("targetMapping"), identity.id)
        mapping_status = record.get("mappingStatus") or ("verification" if identity.sheet == "Audit" and (plan or {}).get("executionScope") != "input-tables-only" else "unresolved")
        if not isinstance(formula, str):
            refs, parse_error = [], graph.unsupported_formulas[identity]
        else:
            try:
                expression = parse_expression(formula, sheet=identity.sheet)
                named_refs = [_reference_id(ref) for ref in extract_references(expression) if ref.name]
                refs = sorted({cell.id for cell in graph.dependencies.get(identity, ())} | set(named_refs))
                parse_error = None
            except ExpressionError as error:
                refs, parse_error = [], str(error)
        item = {
            "id": identity.id,
            "sheet": identity.sheet,
            "address": identity.address,
            "originalFormula": formula,
            "dependencies": refs,
            "mappingStatus": mapping_status,
            "targetMapping": record.get("targetMapping"),
            "evidence": list(record.get("evidence") or []),
            "verification": record.get("verification"),
        }
        if record.get("reviewedReplacement") is not None:
            item["reviewedReplacement"] = record["reviewedReplacement"]
        cached_value = cached[identity.sheet][identity.address].value
        if record.get("mappingStatus") == "cached":
            item["cachedFallback"] = cached_value
        native = native_records.get(identity.id)
        if native is not None:
            item["dependencies"] = list(native["dependencies"])
            item["nativeCandidate"] = {"sourceSha256": source_hash,
                "nativePlanHash": plan["nativeCalculationPlan"]["nativePlanHash"],
                "formulaSha256": native["formulaSha256"], "astHash": native["astHash"],
                "sourceResultClass": native["sourceResultClass"],
                "runtimeStatus": "UNVERIFIED_CANDIDATE"}
            # Audit cells retain their distinct verification obligation.
            if identity.sheet == "Audit":
                item["mappingStatus"] = "verification"
        if parse_error and native is None and not _valid_reviewed_replacement(record, formula):
            item["mappingStatus"] = "unresolved"
            item["unresolvedReason"] = parse_error
        if parse_error and _valid_reviewed_replacement(record, formula):
            item["dependencies"] = list(record.get("dependencies") or [])
        elif parse_error:
            pass
        elif native is not None:
            pass
        elif identity in cycle_members:
            item["mappingStatus"] = "unresolved"
            item["unresolvedReason"] = "true cell dependency cycle"
        elif identity in unresolved_nodes:
            item["mappingStatus"] = "unresolved"
            item["unresolvedReason"] = "missing or transitively unresolved source reference"
        formulas.append(item)
    literals = []
    for identity in sorted(graph.literals):
        value = graph.literals[identity]
        if isinstance(value, (str, int, float, bool)):
            literals.append({"id": identity.id, "sheet": identity.sheet,
                             "address": identity.address, "value": value})
    workbook.close(); cached.close()
    return {
        "manifestVersion": MANIFEST_VERSION,
        "source": {"name": source.name, "path": str(source), "sha256": source_hash},
        "sourceHash": source_hash,
        "planProjection": _plan_projection(plan),
        "planHash": _plan_hash(source_hash, _plan_projection(plan), formulas),
        "formulaCells": formulas,
        "literalInputs": literals,
        "cycles": [[cell.id for cell in component] for component in graph.cycles()],
        "unsupported": {cell.id: reason for cell, reason in sorted(graph.unsupported_formulas.items())},
        "stagingApproval": None,
        "target": None,
    }


def _plan_cell_index(plan):
    result = {}
    for key in ("formulaCells", "cells", "provenance"):
        items = plan.get(key) or []
        if isinstance(items, dict):
            items = items.values()
        for item in items:
            if isinstance(item, dict) and item.get("id"):
                result[item["id"]] = item
    return result


def _reference_id(ref):
    if ref.name:
        return f"name:{ref.name}"
    return f"{ref.sheet}!{ref.address}"


def execution_scope(manifest):
    """The new scope is an enum, never an ignored annotation.

    Legacy/full plans omit executionScope. Only the explicit literal-input
    first phase currently defines a scoped plan. Unknown/null scope values
    fail closed rather than behaving as a full migration.
    """
    if not isinstance(manifest, dict):
        raise ManifestError("execution scope requires a manifest object")
    projection = manifest.get("planProjection") or {}
    if not isinstance(projection, dict):
        raise ManifestError("execution scope requires an object planProjection")
    if "executionScope" not in projection:
        return None
    if projection["executionScope"] != "input-tables-only":
        raise ManifestError("unsupported executionScope: expected input-tables-only or omitted full scope")
    return "input-tables-only"


def require_full_execution_scope(manifest):
    approval = manifest.get("stagingApproval") or {}
    if (execution_scope(manifest) == "input-tables-only"
            or approval.get("scope") == "input-tables-only"
            or approval.get("status") == "APPROVED_INPUT_STAGING"):
        raise ManifestError("input-tables-only scope cannot authorize full staging, native formulas, completion or production")
    return True


def validate_input_only_formula_obligations(manifest):
    if execution_scope(manifest) != "input-tables-only":
        raise ManifestError("explicit input-tables-only scope is required")
    records = manifest.get("formulaCells")
    if not isinstance(records, list) or any(not isinstance(record, dict) for record in records):
        raise ManifestError("input-tables-only formula obligations must be an array of records")
    for record in records:
        if record.get("mappingStatus") != "unresolved":
            raise ManifestError("input-tables-only formulas must remain unresolved: " + record.get("id", "unknown"))
        if any(record.get(field) for field in (
                "targetMapping", "targetBinding", "actualTarget", "verification", "evidence",
                "runtimeEvidence", "runtime", "runtimeStatus", "execution", "executionState")):
            raise ManifestError("input-tables-only formulas cannot carry target/runtime claims: " + record.get("id", "unknown"))
    return True


def validate_manifest(manifest, source_path=None, *, allow_input_only=False):
    if not isinstance(manifest, dict) or manifest.get("manifestVersion") != MANIFEST_VERSION:
        raise ManifestError("unsupported or malformed manifest version")
    if allow_input_only and execution_scope(manifest) == "input-tables-only":
        validate_input_only_formula_obligations(manifest)
    else:
        require_full_execution_scope(manifest)
    source_hash = manifest.get("sourceHash")
    plan_hash = manifest.get("planHash")
    if not _digest(source_hash) or not _digest(plan_hash):
        raise ManifestError("manifest requires valid sourceHash and planHash")
    formulas = manifest.get("formulaCells")
    literals = manifest.get("literalInputs")
    if not isinstance(formulas, list) or not isinstance(literals, list):
        raise ManifestError("manifest requires formulaCells and literalInputs arrays")
    formula_ids = [item.get("id") for item in formulas if isinstance(item, dict)]
    literal_ids = [item.get("id") for item in literals if isinstance(item, dict)]
    if len(formula_ids) != len(formulas) or len(set(formula_ids)) != len(formula_ids):
        raise ManifestError("each source formula must appear exactly once")
    if set(formula_ids) & set(literal_ids):
        raise ManifestError("literal inputs and formula cells must be distinct")
    if any(item.get("mappingStatus") not in {"live", "cached", "unresolved", "verification"}
           for item in formulas):
        raise ManifestError("formula mappingStatus is invalid")
    for item in formulas:
        _validate_target_mapping(item.get("targetMapping"), item["id"])
    projection = manifest.get("planProjection")
    if not isinstance(projection, dict):
        raise ManifestError("plan hash does not match canonical logical mappings")
    # Revisit designated records even when a supplied projection has a valid
    # digest; hash integrity alone does not enforce the logical mapping schema.
    _plan_projection(projection)
    if _plan_hash(source_hash, projection, formulas) != plan_hash:
        raise ManifestError("plan hash does not match canonical logical mappings")
    source_path = source_path or (manifest.get("source") or {}).get("path")
    if source_path is None or not Path(source_path).is_file():
        raise ManifestError("source workbook is required to validate formula completeness")
    if _sha256(source_path) != source_hash:
        raise ManifestError("source hash does not match manifest")
    workbook = load_workbook(source_path, data_only=False, read_only=False)
    graph = FormulaGraph.from_workbook(workbook)
    native_records = _native_records(source_hash, manifest.get("planProjection") or {}, source_path)
    workbook.close()
    actual = {cell.id: formula for cell, formula in graph.formulas.items()}
    _native_records(source_hash, manifest.get("planProjection") or {})
    recorded = {item["id"]: item.get("originalFormula") for item in formulas}
    if recorded != actual:
        raise ManifestError("manifest formula records are incomplete or stale")
    object_formulas = {cell.id: graph.unsupported_formulas[cell]
                       for cell, formula in graph.formulas.items()
                       if not isinstance(formula, str)}
    for item in formulas:
        if item["id"] in object_formulas:
            if (item.get("mappingStatus") != "unresolved"
                    or item.get("unresolvedReason") != object_formulas[item["id"]]):
                raise ManifestError(f"formula object requires unresolved review: {item['id']}")
    return True


def validate_staging_approval(manifest, approval):
    validate_manifest(manifest)
    if not isinstance(approval, dict):
        raise ManifestError("staging approval must be an object")
    if approval.get("scope") == "input-tables-only" or approval.get("status") == "APPROVED_INPUT_STAGING":
        raise ManifestError("input-tables-only approval cannot authorize full staging")
    if approval.get("sourceHash") != manifest["sourceHash"] or approval.get("planHash") != manifest["planHash"]:
        raise ManifestError("staging approval is stale")
    if not str(approval.get("owner") or "").strip():
        raise ManifestError("staging approval requires a named owner")
    mapped = {item["id"] for item in manifest["formulaCells"]}
    reviewed = approval.get("reviewItems")
    if not isinstance(reviewed, list) or set(reviewed) != mapped:
        raise ManifestError("staging approval must explicitly list every mapped review item")
    unresolved = [item["id"] for item in manifest["formulaCells"]
                  if item["mappingStatus"] in {"unresolved", "cached"}]
    if unresolved:
        raise ManifestError("staging approval cannot include unresolved formula mappings")
    return True


def validate_completion(manifest):
    validate_manifest(manifest)
    target = manifest.get("target") or {}
    if not target.get("documentId") or target.get("version") is None:
        raise ManifestError("completion requires an actual target document ID and version")
    for item in manifest["formulaCells"]:
        if item.get("mappingStatus") in {"cached", "unresolved"}:
            raise ManifestError(f"formula is not completion-ready: {item['id']}")
        evidence = item.get("verification") or {}
        binding = item.get("targetBinding") or item.get("actualTarget") or {}
        if evidence.get("status") != "passed" or not evidence.get("evidenceId"):
            raise ManifestError(f"verification evidence missing for {item['id']}")
        if evidence.get("sourceCellId") != item["id"]:
            raise ManifestError(f"verification evidence source binding is stale for {item['id']}")
        if not evidence.get("targetId"):
            raise ManifestError(f"verification evidence target binding is missing for {item['id']}")
        expected_target = (binding.get("columnId") or binding.get("elementId")
                           or binding.get("checkId"))
        if not expected_target or evidence.get("targetId") != expected_target:
            raise ManifestError(f"verification evidence target binding is stale for {item['id']}")
        if (not binding.get("elementId") or binding.get("documentId") != target.get("documentId")
                or binding.get("version") != target.get("version")):
            raise ManifestError(f"actual target binding is incomplete or stale for {item['id']}")
        if evidence.get("sourceHash") != manifest["sourceHash"]:
            raise ManifestError(f"verification evidence source hash is stale for {item['id']}")
        if evidence.get("planHash") != manifest["planHash"]:
            raise ManifestError(f"verification evidence plan hash is stale for {item['id']}")
        if target and (evidence.get("targetVersion") != target.get("version") or
                       evidence.get("targetDocumentId") != target.get("documentId")):
            raise ManifestError(f"verification evidence target version is stale for {item['id']}")
    _validate_source_issue_completion(manifest)
    if manifest.get("planProjection", {}).get("customPortArchitecture"):
        try:
            from .custom_ports import validate_custom_port_completion
        except ImportError:
            from custom_ports import validate_custom_port_completion
        validate_custom_port_completion(manifest)
    if manifest.get('planProjection', {}).get('workflowRepairArchitecture'):
        try:
            from .workflow_repairs import validate_workflow_completion
        except ImportError:
            from workflow_repairs import validate_workflow_completion
        validate_workflow_completion(manifest)
    return True


def _validate_source_issue_completion(manifest):
    # Import here to avoid the gate/manifest import cycle. Re-audit source bytes
    # rather than trusting a caller-editable staging receipt's pending list.
    try:
        from .migration_gate import gate_plan_staging_source, source_issue_hash, MigrationGateError
    except ImportError:
        from migration_gate import gate_plan_staging_source, source_issue_hash, MigrationGateError
    plan = {**manifest["planProjection"], "sourceSha256": manifest["sourceHash"]}
    try:
        _source, audit = gate_plan_staging_source(
            plan, source_arg=manifest["source"]["path"])
    except MigrationGateError as error:
        raise ManifestError(f"source issue audit prevents completion: {error}") from error
    issues = audit.get("issues") or []
    records = manifest.get("sourceReviewEvidence") or []
    if not isinstance(records, list) or any(not isinstance(record, dict) for record in records):
        raise ManifestError("source issue runtime evidence must be an array of records")
    indexed = {record.get("issueHash"): record for record in records}
    expected = {source_issue_hash(issue) for issue in issues}
    if len(indexed) != len(records) or set(indexed) != expected:
        raise ManifestError("source issue runtime evidence must exactly cover audited pending features")
    target = manifest["target"]
    for issue in issues:
        record = indexed[source_issue_hash(issue)]
        binding = record.get("targetBinding") or {}
        expected_target = binding.get("columnId") or binding.get("checkId") or binding.get("elementId")
        if (record.get("code") != issue["code"] or record.get("status") != "passed"
                or not record.get("evidenceId") or not record.get("targetId")
                or record.get("targetId") != expected_target or not binding.get("elementId")
                or binding.get("documentId") != target["documentId"]
                or binding.get("version") != target["version"]
                or record.get("sourceHash") != manifest["sourceHash"]
                or record.get("planHash") != manifest["planHash"]
                or record.get("targetDocumentId") != target["documentId"]
                or record.get("targetVersion") != target["version"]):
            raise ManifestError(f"source issue runtime evidence is incomplete or stale: {issue['code']}")


def _digest(value):
    return isinstance(value, str) and len(value) == 64 and all(ch in "0123456789abcdef" for ch in value)
