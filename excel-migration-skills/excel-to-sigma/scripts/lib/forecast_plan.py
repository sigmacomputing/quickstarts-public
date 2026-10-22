"""Draft a ``build-forecast-model.py`` plan from a cell-drawn P&L workbook.

The drafter reads the report sheet's formulas and Excel's cached values, the GL
actuals tab its SUMIFS read, and the driver tab its growth formulas read. It then
re-derives every line-month from the drafted plan and compares it with the
workbook's cached values, so a plan that leaves this module already reproduces
the spreadsheet.

Every precondition fails loudly with ``PlanDraftError`` and a precise cell
reference. An unknown shape is never coerced (SKILL.md Phase 0c): the builder
supports exactly the four sections in ``SECTIONS`` and the kinds growth / flat /
manual.
"""
from __future__ import annotations

import calendar
import importlib.util
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

try:
    from .migration_gate import source_sha256
except ImportError:  # direct import when scripts/lib itself is on sys.path
    from migration_gate import source_sha256

SECTIONS = ("Revenue", "Cost of Revenue", "Operating Expenses", "Below the Line")
DERIVED_LABELS = {"gross profit", "ebitda", "net income"}
TARGET_LABELS = {"total revenue": "Revenue", "gross profit": "Gross Profit",
                 "ebitda": "EBITDA", "net income": "Net Income"}
INTENTS = ("read-only", "editable", "what-if")
MONTH_ABBR = ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec")

MONTH_RE = re.compile(r"^\s*([A-Za-z]{3,9})[-\s/']*(\d{2}|\d{4})\s*$")
SUMIFS_RE = re.compile(r"\bSUMIFS\b", re.I)
SHEET_REF_RE = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_.]+))!")
GROWTH_RE = re.compile(
    r"^=\$?([A-Z]{1,3})\$?(\d+)\*\(1\+(?:'((?:[^']|'')+)'|([A-Za-z0-9_.]+))!\$?([A-Z]{1,3})\$?(\d+)\)$")
FLAT_RE = re.compile(r"^=\$?([A-Z]{1,3})\$?(\d+)$")
UNQUOTED_SPACE_RE = re.compile(r"\s+(?=(?:[^']*'[^']*')*[^']*$)")
FY_HEADER_RE = re.compile(
    r"^(?:fy(?:\s*\d{2,4}|\s+total)?|total|full[ -]?year|annual|year)$", re.I)


class PlanDraftError(RuntimeError):
    """The workbook does not have a shape this drafter supports."""


# ---------------------------------------------------------------- helpers
def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _text(value):
    return value.strip() if isinstance(value, str) and value.strip() else None


def _formula(value):
    return value if isinstance(value, str) and value.startswith("=") else None


def parse_month(value):
    """Return ``(year, month)`` for a month header (datetime or ``Jan-26`` text), else None."""
    if isinstance(value, (datetime, date)):
        return value.year, value.month
    if isinstance(value, str):
        found = MONTH_RE.match(value)
        if found and found.group(1)[:3].lower() in MONTH_ABBR:
            year = int(found.group(2))
            return (year + 2000 if year < 100 else year), MONTH_ABBR.index(found.group(1)[:3].lower()) + 1
    return None


def month_meta(year_month, num):
    year, month = year_month
    end = date(year, month, calendar.monthrange(year, month)[1])
    return {"month": f"{MONTH_ABBR[month - 1].capitalize()}-{year % 100:02d}", "num": num,
            "end": end.isoformat()}


def _ref(ws, col, row):
    return f"{ws.title}!{get_column_letter(col)}{row}"


def _load_discover():
    """Import xlsx-discover.py so report-sheet selection matches discovery exactly."""
    scripts = Path(__file__).resolve().parents[1]
    if str(scripts) not in sys.path:
        sys.path.insert(0, str(scripts))
    spec = importlib.util.spec_from_file_location("xlsx_discover_for_forecast_plan",
                                                  scripts / "xlsx-discover.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


# ---------------------------------------------------------------- locating
def find_report_sheet(wf, sheet=None):
    if sheet:
        if sheet not in wf.sheetnames:
            raise PlanDraftError(f"--sheet {sheet!r} not found; sheets: {', '.join(wf.sheetnames)}")
        return wf[sheet]
    sheetmap, _, _ = _load_discover().classify_report_cells(wf)
    scored = sorted(((counts.get("ACTUAL", 0) + counts.get("DRIVER_GROWTH", 0), name)
                     for name, counts in sheetmap.items()), reverse=True)
    if not scored or scored[0][0] == 0:
        raise PlanDraftError(
            "no sheet has SUMIFS actuals or driver-growth formulas; this is not an "
            "actuals + driver forecast model (pass --sheet if you know the report sheet)")
    if len(scored) > 1 and scored[1][0] == scored[0][0]:
        raise PlanDraftError(
            f"sheets {scored[0][1]!r} and {scored[1][1]!r} score equally as the report sheet; "
            "pass --sheet")
    return wf[scored[0][1]]


def find_month_header(ws):
    """First row with >=3 consecutive month labels -> (row, [(col, (year, month)), ...])."""
    for row in range(1, ws.max_row + 1):
        run = []
        best = []
        for col in range(1, ws.max_column + 1):
            parsed = parse_month(ws.cell(row, col).value)
            if parsed:
                run.append((col, parsed))
            else:
                best = max(best, run, key=len)
                run = []
        best = max(best, run, key=len)
        if len(best) < 3:
            continue
        for (_, prev), (col, cur) in zip(best, best[1:]):
            expected = (prev[0] + (prev[1] == 12), prev[1] % 12 + 1)
            if cur != expected:
                raise PlanDraftError(
                    f"{_ref(ws, col, row)}: month headers are not consecutive calendar months "
                    f"({prev[0]}-{prev[1]:02d} then {cur[0]}-{cur[1]:02d})")
        return row, best
    raise PlanDraftError(f"{ws.title!r}: no header row with 3 or more consecutive month labels")


def _walk_rows(ws, header_row, month_cols, fy_col):
    """Yield ('section', row, label) / ('data', row, label) in sheet order."""
    value_cols = list(month_cols) + ([fy_col] if fy_col else [])
    rows = []
    for row in range(header_row + 1, ws.max_row + 1):
        label = ws.cell(row, 1).value
        has_values = any(ws.cell(row, col).value is not None for col in value_cols)
        if label is None and not has_values:
            continue
        if label is not None and not _text(label) and not has_values:
            continue
        rows.append((row, label, has_values))
    for index, (row, label, has_values) in enumerate(rows):
        if has_values:
            yield "data", row, label
        elif _text(label) and index + 1 < len(rows) and rows[index + 1][2]:
            yield "section", row, _text(label)
        # else: a note / legend row with nothing under it


def _forecast_cell_kind(ws, col, row, value):
    """Return (kind, detail) for one forecast cell, or (None, repr) when unrecognised."""
    formula = _formula(value)
    if formula:
        # drop whitespace outside quoted sheet names ('Driver Assumptions'!B5 keeps its space)
        compact = UNQUOTED_SPACE_RE.sub("", formula)
        prev_col = get_column_letter(col - 1)
        grown = GROWTH_RE.match(compact)
        if grown and grown.group(1) == prev_col and int(grown.group(2)) == row:
            sheet = (grown.group(3) or grown.group(4)).replace("''", "'")
            return "growth", (sheet, grown.group(5), int(grown.group(6)))
        flat = FLAT_RE.match(compact)
        if flat and flat.group(1) == prev_col and int(flat.group(2)) == row:
            return "flat", None
        return None, repr(value)
    if _is_number(value):
        return "manual", None
    return None, repr(value)


# ---------------------------------------------------------------- drafting
def draft_plan(source, *, intent, connection, folder, name=None, sheet=None,
               currency_format="$,.0f", pct_format=".1%", tolerance=0.01):
    """Return ``(plan, summary)`` for ``source``; raise ``PlanDraftError`` on any unknown shape."""
    if intent not in INTENTS:
        raise PlanDraftError(f"intent must be one of {', '.join(INTENTS)}")
    src = Path(source).expanduser().resolve()
    if not src.is_file():
        raise PlanDraftError(f"source workbook not found: {src}")
    wf = load_workbook(src)
    wv = load_workbook(src, data_only=True)

    ws = find_report_sheet(wf, sheet)
    vs = wv[ws.title]
    header_row, months = find_month_header(ws)
    month_cols = [col for col, _ in months]
    next_header = ws.cell(header_row, month_cols[-1] + 1).value
    fy_col = month_cols[-1] + 1 if _text(next_header) and FY_HEADER_RE.match(_text(next_header)) else None

    # ---- sections and rows
    sections_seen, derived_rows, line_rows = [], [], []
    section = None
    for kind, row, label in _walk_rows(ws, header_row, month_cols, fy_col):
        if kind == "section":
            if label not in SECTIONS:
                found = [lab for k, _, lab in _walk_rows(ws, header_row, month_cols, fy_col)
                         if k == "section"]
                raise PlanDraftError(
                    f"{_ref(ws, 1, row)}: unsupported section {label!r}; this builder supports "
                    f"only {', '.join(SECTIONS)}. Sections found: {', '.join(found)}")
            if label in sections_seen:
                raise PlanDraftError(f"{_ref(ws, 1, row)}: section {label!r} appears twice")
            sections_seen.append(label)
            section = label
            continue
        text = _text(label)
        if text is None:
            raise PlanDraftError(f"{_ref(ws, 1, row)}: row has values but no text label")
        lowered = text.lower()
        if (lowered.startswith("total") or lowered in DERIVED_LABELS
                or "%" in text or "margin" in lowered):
            derived_rows.append((row, text))
            continue
        first = ws.cell(row, month_cols[0]).value
        if not (_formula(first) and SUMIFS_RE.search(first)):
            raise PlanDraftError(
                f"{_ref(ws, month_cols[0], row)} ({text!r}): row is neither a SUMIFS actuals line "
                "nor a recognised Total/Gross Profit/EBITDA/Net Income/margin row")
        if section is None:
            raise PlanDraftError(f"{_ref(ws, 1, row)} ({text!r}): line appears before any section label")
        line_rows.append((row, text, section))
    if not line_rows:
        raise PlanDraftError(f"{ws.title!r}: no SUMIFS actuals lines found under a section")
    names = Counter(text for _, text, _ in line_rows)
    dup = [n for n, c in names.items() if c > 1]
    if dup:
        raise PlanDraftError(f"duplicate line labels {dup}: SUMIFS keys on the label, so it must be unique")

    # ---- actual vs forecast columns
    n_act = 0
    first_row = line_rows[0][0]
    for col in month_cols:
        cell = ws.cell(first_row, col).value
        if _formula(cell) and SUMIFS_RE.search(cell):
            n_act += 1
        else:
            break
    if n_act == 0 or n_act == len(month_cols):
        raise PlanDraftError(
            f"{ws.title!r}: expected SUMIFS actual months followed by forecast months; "
            f"found {n_act} actual of {len(month_cols)} months")
    act_cols, fc_cols = month_cols[:n_act], month_cols[n_act:]
    actual_months = [month_meta(ym, i + 1) for i, (_, ym) in enumerate(months[:n_act])]
    forecast_months = [month_meta(ym, n_act + i + 1) for i, (_, ym) in enumerate(months[n_act:])]
    last_act = act_cols[-1]

    # ---- drivers
    lines, sumifs_sheets = [], set()
    driver_ranges = set()
    flat_lines = []
    for line_order, (row, text, sec) in enumerate(line_rows, 1):
        for col in act_cols:
            cell = ws.cell(row, col).value
            if not (_formula(cell) and SUMIFS_RE.search(cell)):
                raise PlanDraftError(
                    f"{_ref(ws, col, row)} ({text!r}): actual-month cell is not a SUMIFS "
                    f"formula: {cell!r}")
            for found in SHEET_REF_RE.finditer(cell):
                sumifs_sheets.add((found.group(1) or found.group(2)).replace("''", "'"))
        kinds = []
        for col in fc_cols:
            kind, detail = _forecast_cell_kind(ws, col, row, ws.cell(row, col).value)
            if kind is None:
                raise PlanDraftError(
                    f"{_ref(ws, col, row)} ({text!r}): unrecognised forecast cell {detail}; "
                    "expected =prev*(1+Drivers!$X$n), =prev, or a numeric literal")
            kinds.append((col, kind, detail))
        first_col, first_kind, first_detail = kinds[0]
        for col, kind, detail in kinds[1:]:
            if kind != first_kind or detail != first_detail:
                raise PlanDraftError(
                    f"{_ref(ws, col, row)} ({text!r}): forecast cells are not uniform "
                    f"({get_column_letter(first_col)}{row} is {first_kind}, "
                    f"{get_column_letter(col)}{row} is {kind}"
                    f"{'' if detail == first_detail else ' with a different driver'})")
        base = vs.cell(row, last_act).value
        entry = {"line": text, "section": sec, "secOrder": SECTIONS.index(sec) + 1,
                 "lineOrder": line_order, "sourceRow": row, "kind": first_kind}
        if first_kind == "growth":
            dsheet, dcol, drow = first_detail
            if dsheet not in wv.sheetnames:
                raise PlanDraftError(f"{_ref(ws, first_col, row)}: driver sheet {dsheet!r} not found")
            rate = wv[dsheet][f"{dcol}{drow}"].value
            if not _is_number(rate):
                raise PlanDraftError(
                    f"{dsheet}!{dcol}{drow}: driver for {text!r} is not a numeric rate ({rate!r})")
            driver_ranges.add((dsheet, dcol))
            drv_name = next((_text(c.value) for c in wv[dsheet][drow] if _text(c.value)), None)
            entry.update(rate=float(rate), juneBase=_number(base, ws, last_act, row, text),
                         driver=drv_name, driverCell=f"{dsheet}!{dcol}{drow}")
        elif first_kind == "flat":
            entry.update(juneBase=_number(base, ws, last_act, row, text))
            flat_lines.append(text)
        else:
            entry["manual"] = [float(ws.cell(row, col).value) for col in fc_cols]
        lines.append(entry)
    if len(driver_ranges) > 1:
        raise PlanDraftError("growth lines read rates from more than one driver sheet/column "
                             f"({sorted(driver_ranges)}); the builder supports one")

    assumptions = []
    if driver_ranges:
        dsheet, dcol = next(iter(driver_ranges))
        dws = wv[dsheet]
        rate_col = column_index_from_string(dcol)
        for drow in range(1, dws.max_row + 1):
            rate = dws.cell(drow, rate_col).value
            drv_name = next((_text(c.value) for c in dws[drow][:rate_col - 1] if _text(c.value)), None)
            if drv_name and _is_number(rate):
                basis = next((_text(c.value) for c in dws[drow][rate_col:] if _text(c.value)), "")
                assumptions.append([drv_name, float(rate), basis])

    # ---- actuals (GL)
    if len(sumifs_sheets - {ws.title}) != 1:
        raise PlanDraftError(
            f"expected the actual-month SUMIFS to read exactly one GL sheet, found "
            f"{sorted(sumifs_sheets) or 'none'}")
    gl_name = next(iter(sumifs_sheets - {ws.title}))
    actuals = read_actuals(wv[gl_name], lines, actual_months)

    # ---- verification: GL vs cached SUMIFS cells, then every line-month from the plan
    problems = verify(ws, vs, lines, actuals, actual_months, forecast_months,
                      act_cols, fc_cols, last_act, tolerance)
    if problems:
        shown = "\n  ".join(problems[:10])
        more = f"\n  ... and {len(problems) - 10} more" if len(problems) > 10 else ""
        raise PlanDraftError(
            f"drafted plan does not reproduce the workbook's cached values "
            f"({len(problems)} mismatch(es)):\n  {shown}{more}")

    # A header alone does not prove a full-year total: every line row's cached FY
    # value must equal the sum of its month cells, else targets come from the months.
    if fy_col and not _fy_is_month_sum(vs, line_rows, month_cols, fy_col, tolerance):
        fy_col = None
    targets = parity_targets(vs, derived_rows, month_cols, fy_col)
    charts, unmapped_charts = source_charts(wf, wv, lines, actuals, actual_months, forecast_months,
                                            tolerance)

    issues = []
    if flat_lines:
        issues.append({
            "code": "FLAT_CARRY_POLICY",
            "severity": "review",
            "route": "forecast-carry-forward-policy-review",
            "message": (f"{len(flat_lines)} forecast line(s) carry the prior month forward "
                        "(=prior); that syntax alone does not prove a flat carry-forward policy."),
            "evidence": flat_lines,
        })

    plan = {
        "name": name or (wf.properties.title or "").strip() or src.stem,
        "intent": intent,
        "connectionId": connection,
        "folderId": folder,
        "currencyFormat": currency_format,
        "pctFormat": pct_format,
        "sourcePath": str(src),
        "sourceSha256": source_sha256(src),
        "parityTargets": targets,
        "actualMonths": actual_months,
        "forecastMonths": forecast_months,
        "lines": lines,
        "assumptions": assumptions,
        "actuals": actuals,
        "sourceCharts": charts,
        "draftReviewIssues": issues,
    }
    summary = {
        "reportSheet": ws.title, "glSheet": gl_name,
        "kinds": dict(Counter(line["kind"] for line in lines)),
        "lines": len(lines), "glRows": len(actuals), "drivers": len(assumptions),
        "actualRange": (actual_months[0]["month"], actual_months[-1]["month"]),
        "forecastRange": (forecast_months[0]["month"], forecast_months[-1]["month"]),
        "fyColumn": get_column_letter(fy_col) if fy_col else None,
        "parityTargets": targets,
        "missingTargets": [m for m in TARGET_LABELS.values() if m not in targets],
        "reviewIssues": [issue["code"] for issue in issues],
        "charts": [f"{c['title']} ({c['grain']}, {len(c['series'])} series)" for c in charts],
        "unmappedCharts": unmapped_charts,
    }
    return plan, summary


def _number(value, ws, col, row, label):
    if not _is_number(value):
        raise PlanDraftError(
            f"{_ref(ws, col, row)} ({label!r}): no numeric cached value ({value!r}); open the "
            "workbook in Excel, recalculate and save so cached values exist")
    return float(value)


def read_actuals(gl, lines, actual_months):
    """Map the GL header by column name and return plan ``actuals`` rows."""
    by_line = {line["line"]: line for line in lines}
    months = {m["month"] for m in actual_months}
    roles = {
        "date": lambda h, t: "date" in t,
        "month": lambda h, t: h in ("month", "period", "month label"),
        "line": lambda h, t: h in ("line", "line item", "account", "gl account", "line item name"),
        "dept": lambda h, t: bool(t & {"dept", "department"}),
        "vendor": lambda h, t: "vendor" in t,
        "amt": lambda h, t: bool(t & {"amount", "amt"}),
    }
    header_row, cols, seen = None, {}, []
    for row in range(1, min(gl.max_row, 40) + 1):
        found = {}
        for col in range(1, gl.max_column + 1):
            value = gl.cell(row, col).value
            if not _text(value):
                continue
            norm = re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()
            tokens = set(norm.split())
            for role, match in roles.items():
                if match(norm, tokens):
                    found.setdefault(role, []).append(col)
        seen.append((len(found), row, found))
        if len(found) == len(roles):
            header_row, cols = row, found
            break
    if header_row is None:
        best = max(seen, key=lambda item: item[0])
        missing = [r for r in roles if r not in best[2]]
        raise PlanDraftError(
            f"{gl.title!r}: no header row maps all of date, month, line item, department, vendor, "
            f"amount; best candidate row {best[1]} is missing {', '.join(missing)}")
    ambiguous = [r for r, found in cols.items() if len(found) > 1]
    if ambiguous:
        raise PlanDraftError(f"{gl.title!r} row {header_row}: more than one column matches "
                             f"{', '.join(ambiguous)}")
    col = {role: found[0] for role, found in cols.items()}

    out = []
    for row in range(header_row + 1, gl.max_row + 1):
        values = {role: gl.cell(row, c).value for role, c in col.items()}
        if all(v is None for v in values.values()):
            continue
        where = f"{gl.title}!row {row}"
        label = _text(values["line"])
        if label not in by_line:
            raise PlanDraftError(f"{where}: line item {values['line']!r} is not a P&L line")
        parsed = parse_month(values["month"])
        month = (f"{MONTH_ABBR[parsed[1] - 1].capitalize()}-{parsed[0] % 100:02d}"
                 if parsed else None)
        if month not in months:
            raise PlanDraftError(f"{where}: month {values['month']!r} is not an actual month "
                                 f"({', '.join(sorted(months))})")
        if not _is_number(values["amt"]):
            raise PlanDraftError(f"{where}: amount {values['amt']!r} is not numeric")
        stamp = values["date"]
        if isinstance(stamp, datetime):
            stamp = stamp.date()
        if isinstance(stamp, date):
            stamp = stamp.isoformat()
        elif not (isinstance(stamp, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", stamp.strip())):
            raise PlanDraftError(f"{where}: date {values['date']!r} is not a date")
        line = by_line[label]
        out.append({"date": stamp.strip() if isinstance(stamp, str) else stamp, "month": month,
                    "line": label, "section": line["section"], "secOrder": line["secOrder"],
                    "lineOrder": line["lineOrder"], "dept": _text(values["dept"]) or "",
                    "vendor": _text(values["vendor"]) or "", "amt": float(values["amt"])})
    if not out:
        raise PlanDraftError(f"{gl.title!r}: no GL rows below the header")
    return out


def verify(ws, vs, lines, actuals, actual_months, forecast_months, act_cols, fc_cols, last_act,
           tolerance):
    """Recompute each line-month the way the builder does; list disagreements with the sheet."""
    gl = Counter()
    for r in actuals:
        gl[(r["line"], r["month"])] += r["amt"]
    problems = []

    def compare(coord, label, month, got, want):
        if not _is_number(want):
            problems.append(f"{coord} ({label!r}, {month}): sheet has no numeric cached value ({want!r})")
        elif abs(got - want) > tolerance:
            problems.append(f"{coord} ({label!r}, {month}): plan {got:,.2f} vs sheet {want:,.2f}")

    for ln in lines:
        row = ln["sourceRow"]
        for col, month in zip(act_cols, actual_months):
            compare(f"{ws.title}!{get_column_letter(col)}{row}", ln["line"], month["month"],
                    gl[(ln["line"], month["month"])], vs.cell(row, col).value)
        base = ln.get("juneBase", 0)
        for i, (col, month) in enumerate(zip(fc_cols, forecast_months)):
            n = month["num"] - actual_months[-1]["num"]
            if ln["kind"] == "manual":
                got = ln["manual"][i]
            elif ln["kind"] == "flat":
                got = base
            else:
                got = base * (1 + ln["rate"]) ** n
            compare(f"{ws.title}!{get_column_letter(col)}{row}", ln["line"], month["month"],
                    got, vs.cell(row, col).value)
    return problems


def _fy_is_month_sum(vs, line_rows, month_cols, fy_col, tolerance):
    for row, _, _ in line_rows:
        cells = [vs.cell(row, col).value for col in month_cols]
        total = vs.cell(row, fy_col).value
        if not (_is_number(total) and all(_is_number(c) for c in cells)
                and abs(sum(cells) - total) <= tolerance):
            return False
    return True


def parity_targets(vs, derived_rows, month_cols, fy_col):
    """Cached FY (else summed month) totals for Revenue / Gross Profit / EBITDA / Net Income."""
    targets = {}
    for row, label in derived_rows:
        metric = TARGET_LABELS.get(label.lower())
        if not metric or metric in targets:
            continue
        if fy_col:
            value = vs.cell(row, fy_col).value
        else:
            cells = [vs.cell(row, col).value for col in month_cols]
            value = sum(cells) if all(_is_number(c) for c in cells) else None
        if _is_number(value):
            targets[metric] = float(value)
    return targets


# ---------------------------------------------------------------- source charts
QUARTER_LABEL_RE = re.compile(r"^\s*Q([1-4])\b", re.I)


def _range_values(wv, formula):
    """Cached values of a chart reference such as ``Quarterly!$B$4:$E$4`` (row or column)."""
    found = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^!]+))!(\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)",
                         formula or "")
    if not found:
        return None
    sheet = (found.group(1) or found.group(2)).replace("''", "'")
    if sheet not in wv.sheetnames:
        return None
    min_col, min_row, max_col, max_row = range_boundaries(found.group(3).replace("$", ""))
    ws = wv[sheet]
    return [ws.cell(r, c).value for r in range(min_row, max_row + 1)
            for c in range(min_col, max_col + 1)]


def _metric_series(lines, actuals, actual_months, forecast_months):
    """Plan totals per month for each parity metric, computed the way the builder does.

    The actual/forecast boundary is the plan's ``actualMonths``: a final actual month
    with no GL postings is still an actual month.
    """
    months = actual_months + forecast_months
    by_section = {m["month"]: Counter() for m in months}
    for r in actuals:
        by_section[r["month"]][r["section"]] += r["amt"]
    last_actual = actual_months[-1]["num"]
    for ln in lines:
        base = ln.get("juneBase", 0)
        for i, m in enumerate(forecast_months):
            if ln["kind"] == "manual":
                value = ln["manual"][i]
            elif ln["kind"] == "flat":
                value = base
            else:
                value = base * (1 + ln["rate"]) ** (m["num"] - last_actual)
            by_section[m["month"]][ln["section"]] += value
    out = {}
    for m in months:
        s = by_section[m["month"]]
        rev = s["Revenue"]
        gp = rev - s["Cost of Revenue"]
        ebitda = gp - s["Operating Expenses"]
        out[m["month"]] = {"Revenue": rev, "Gross Profit": gp, "EBITDA": ebitda,
                           "Net Income": ebitda - s["Below the Line"]}
    return out


def source_charts(wf, wv, lines, actuals, actual_months, forecast_months, tolerance):
    """Map workbook line charts whose series are P&L metrics to plan ``sourceCharts``.

    A chart maps when every series title is Total Revenue / Gross Profit / EBITDA /
    Net Income and its categories are quarters (``Q1 ...``) or months. Mapped charts
    must reproduce their cached series values from the plan; a mismatch raises.
    Returns ``(charts, unmapped)`` where ``unmapped`` names charts left to review.
    """
    months = actual_months + forecast_months
    monthly = _metric_series(lines, actuals, actual_months, forecast_months)
    quarters = []
    for m in months:
        year, month = int(m["end"][:4]), int(m["end"][5:7])
        key = (year, (month - 1) // 3 + 1)
        if not quarters or quarters[-1][0] != key:
            quarters.append((key, []))
        quarters[-1][1].append(m["month"])

    charts, unmapped = [], []
    for ws in wf.worksheets:
        for chart in getattr(ws, "_charts", []):
            title = _chart_title(chart) or f"{type(chart).__name__} on {ws.title}"
            where = f"{title!r} on {ws.title!r}"
            if type(chart).__name__ != "LineChart" or not chart.series:
                unmapped.append(f"{where}: only line charts are mapped")
                continue
            metrics, categories = [], None
            for s in chart.series:
                ref = s.tx.strRef.f if s.tx is not None and s.tx.strRef is not None else None
                label = (_range_values(wv, ref) or [None])[0] if ref else None
                metric = TARGET_LABELS.get(str(label).strip().lower()) if label else None
                cat_source = None
                if s.cat is not None:
                    cat_source = s.cat.strRef if s.cat.strRef is not None else s.cat.numRef
                cat_ref = cat_source.f if cat_source is not None else None
                cats = _range_values(wv, cat_ref) if cat_ref else None
                vals = _range_values(wv, s.val.numRef.f) if s.val is not None and s.val.numRef else None
                if not metric or not cats or not vals:
                    metrics = None
                    break
                metrics.append((metric, vals))
                categories = cats
            if not metrics:
                unmapped.append(f"{where}: a series is not Total Revenue/Gross Profit/EBITDA/Net Income")
                continue
            if all(QUARTER_LABEL_RE.match(str(c or "")) for c in categories):
                grain = "quarter"
                buckets = [(key, members) for key, members in quarters]
                labels = [int(QUARTER_LABEL_RE.match(str(c)).group(1)) for c in categories]
                if [key[1] for key, _ in buckets[:len(labels)]] != labels:
                    unmapped.append(f"{where}: quarter categories {categories} do not follow the plan months")
                    continue
                buckets = buckets[:len(labels)]
            else:
                labels = [parse_month(c) for c in categories]
                if not all(labels):
                    unmapped.append(f"{where}: categories are neither quarters nor months")
                    continue
                expected = [(int(m["end"][:4]), int(m["end"][5:7]))
                            for m in months[:len(labels)]]
                if labels != expected:
                    unmapped.append(f"{where}: month categories {categories} do not follow the plan months")
                    continue
                grain = "month"
                buckets = [((m["month"],), [m["month"]]) for m in months][:len(labels)]
            for metric, vals in metrics:
                for (key, members), want in zip(buckets, vals):
                    got = sum(monthly[m][metric] for m in members)
                    if not _is_number(want) or abs(got - want) > tolerance:
                        raise PlanDraftError(
                            f"chart {where}: series {metric!r} at {key} is {want!r} in the workbook "
                            f"but {got:,.2f} from the plan")
            # The source chart's own period range, so the rebuilt chart shows the same subset.
            covered = [m for m in months if any(m["month"] in members for _, members in buckets)]
            charts.append({"title": title, "kind": "line", "grain": grain,
                           "series": [metric for metric, _ in metrics],
                           "categories": [str(c) for c in categories],
                           "periodStart": covered[0]["end"][:8] + "01", "periodEnd": covered[-1]["end"],
                           "sourceSheet": ws.title, "verifiedPoints": len(metrics) * len(buckets)})
    return charts, unmapped


def _chart_title(chart):
    try:
        return "".join(run.t for p in chart.title.tx.rich.p for run in (p.r or [])).strip() or None
    except AttributeError:
        return None
