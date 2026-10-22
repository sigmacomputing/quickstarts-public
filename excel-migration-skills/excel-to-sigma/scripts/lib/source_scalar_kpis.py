"""Source-scoped Sales Overview scalar owners, delivered by native SDK append.

Offline planning only. Reviewed original-row identity and fresh complete native
observations are inputs, never inferred from mutable business keys on retries.
Construction/readback remain separate from scalar runtime acceptance.
"""
from __future__ import annotations

import copy
import hashlib
import json
import math
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

try:
    from .input_sync import _digest
except ImportError:
    from input_sync import _digest


def extract_sales_overview_scalars(source_path):
    """Support the exact finite SUM and row-product family, without caches.

    This intentionally rejects arbitrary SUM syntax and partial/drifting ranges;
    unsupported workbooks need a separately reviewed scalar contract.
    """
    path = Path(source_path).resolve()
    workbook = load_workbook(path, data_only=False)
    try:
        pipeline, overview = workbook['Pipeline'], workbook['Overview']
        bounds = range_boundaries(pipeline.tables['tblSalesPipeline'].ref)
        if bounds != (1, 1, 14, 49):
            raise ValueError('unsupported Sales scalar source table scope')
        for address, name in [('A1', 'Opportunity ID'), ('I1', 'Probability'),
                              ('J1', 'ACV'), ('L1', 'Weighted ACV')]:
            if pipeline[address].value != name:
                raise ValueError('Sales scalar source column changed: ' + address)
        scope = []
        for row in range(2, 50):
            key = pipeline.cell(row, 1).value
            if not isinstance(key, str) or not key or pipeline.cell(row, 1).data_type == 'f':
                raise ValueError('Sales scalar source keys must be authored text')
            if pipeline.cell(row, 12).value != f'=I{row}*J{row}':
                raise ValueError('Sales weighted source row equation changed')
            for col in (9, 10):
                value = pipeline.cell(row, col).value
                if value is not None and (isinstance(value, bool) or not isinstance(value, (int, float))
                                          or not math.isfinite(value)):
                    raise ValueError('Sales scalar numeric source dependencies must be numeric/null literals')
            scope.append({'sourceRow': row, 'originalKey': key})
        if len({item['originalKey'] for item in scope}) != len(scope):
            raise ValueError('Sales scalar source identity is ambiguous')
        owners = []
        for address, column, name, logical, value_id, deps, frame in [
            ('B6', 'J', 'Total ACV', 'sales-overview-total-acv',
             'source_overview_b6_total_acv', ['ACV'], [1, 48, 12, 6]),
            ('B7', 'L', 'Weighted pipeline', 'sales-overview-weighted-pipeline',
             'source_overview_b7_weighted_pipeline', ['Probability', 'ACV'], [13, 48, 12, 6]),
        ]:
            formula = overview[address].value
            expected = '=SUM(' + ','.join(f"'Pipeline'!${column}${r}" for r in range(2, 50)) + ')'
            if formula != expected or overview[address].number_format != '$#,##0':
                raise ValueError('Sales scalar source formula/format/scope changed: ' + address)
            owners.append({
                'sourceId': 'Overview!' + address, 'originalFormula': formula,
                'sourceFormulaHash': _digest(formula),
                'sourceCells': [f'Pipeline!{column}{r}' for r in range(2, 50)],
                'numericDependencies': deps, 'name': name,
                'logicalElementId': logical, 'logicalColumnId': value_id,
                'kind': 'kpi-chart', 'pageId': 'sales-overview', 'frame': frame,
                'formatPattern': '$,.0f', 'emptySum': 0, 'groupings': [],
                'verification': 'UNRESOLVED'})
    except KeyError as error:
        raise ValueError('unsupported Sales scalar source structure') from error
    finally:
        workbook.close()
    contract = {'schema': 'source-scoped-sales-scalars-v1', 'sourcePath': str(path),
                'sourceHash': hashlib.sha256(path.read_bytes()).hexdigest(),
                'sourceRows': scope, 'owners': owners, 'formulaCachesUsed': False,
                'delivery': 'native-sdk-append-only', 'productionStatus': 'BLOCKED'}
    contract['contractHash'] = _digest(contract)
    return contract


def _validate_contract(contract):
    # Re-extraction checks both immutable bytes and the supported semantic shape.
    if contract != extract_sales_overview_scalars(contract['sourcePath']):
        raise ValueError('Sales scalar source/contract changed')


def _layout(contents):
    try:
        # Public layout may contain several Page roots after an XML declaration.
        layout = re.sub(r'\A<\?xml\s[^?]*\?>', '', contents['layout'], count=1)
        return ET.fromstring('<Root>' + layout + '</Root>')
    except (KeyError, ET.ParseError) as error:
        raise ValueError('complete current layout required') from error


def _elements(contents):
    result = {}
    for element in contents['elements']:
        if not element.get('id') or element['id'] in result:
            raise ValueError('missing/duplicate observed element ID')
        result[element['id']] = element
    return result


def _placement(contents, element_id):
    placements = []
    for page in _layout(contents).iter('Page'):
        for element in page.iter('Element'):
            if element.get('elementId') == element_id:
                try:
                    x, end_x = map(int, element.attrib['gridColumn'].split('/'))
                    y, end_y = map(int, element.attrib['gridRow'].split('/'))
                except (KeyError, ValueError) as error:
                    raise ValueError('observed grid placement required') from error
                placements.append((page.get('id'), [x, y, end_x-x, end_y-y]))
    return placements


def _scope(contract, identity, snapshot):
    target = snapshot['input']
    if (identity.get('sourceHash') != contract['sourceHash']
            or identity.get('documentId') != snapshot['documentId']
            or identity.get('inputElementId') != target.get('elementId')
            or identity.get('rowIdColumn') != 'ID' or target.get('rowIdColumn') != 'ID'
            or not identity.get('evidenceId') or not identity.get('reviewer')):
        raise ValueError('reviewed source-bound original physical-row identity receipt required')
    mapping = identity.get('rows', [])
    if (len(mapping) != len(contract['sourceRows'])
            or [{'sourceRow': item.get('sourceRow'), 'originalKey': item.get('originalKey')}
                for item in mapping] != contract['sourceRows']):
        raise ValueError('original physical-row receipt source scope changed')
    ids = [item.get('physicalId') for item in mapping]
    if any(not isinstance(value, str) or not value for value in ids) or len(set(ids)) != len(ids):
        raise ValueError('original physical-row IDs must be complete and unique')
    observed_ids = [row.get('ID') for row in target['rows']]
    if (any(not isinstance(value, str) or not value for value in observed_ids)
            or len(set(observed_ids)) != len(observed_ids)):
        raise ValueError('fresh input rows have missing/duplicate physical IDs')
    if not set(ids) <= set(observed_ids):
        raise ValueError('original physical rows missing; explicit identity reconciliation required')
    return ids, _digest({'sourceHash': contract['sourceHash'],
                        'documentId': snapshot['documentId'],
                        'inputElementId': target['elementId'], 'rows': mapping})


def _validate_snapshot(snapshot):
    if (not snapshot.get('documentId') or type(snapshot.get('version')) is not int
            or snapshot['version'] < 1 or snapshot.get('published') is not True
            or snapshot.get('hasDraft') is not False or not snapshot.get('nativeActions')):
        raise ValueError('fresh published document/version and complete native action inventory required')
    target = snapshot['input']
    elements = _elements(snapshot['contents'])
    element = elements.get(target.get('elementId'), {})
    observed_columns = element.get('columns')
    if not isinstance(observed_columns, list) or not observed_columns:
        raise ValueError('complete input column readback required')
    columns, column_ids = {}, set()
    for column in observed_columns:
        if (not isinstance(column, dict)
                or not isinstance(column.get('name'), str) or not column['name']
                or not isinstance(column.get('id'), str) or not column['id']
                or column['name'] in columns or column['id'] in column_ids):
            raise ValueError('malformed/duplicate input column names or IDs')
        columns[column['name']] = column['id']
        column_ids.add(column['id'])
    if (element.get('kind') != 'input-table' or element.get('name') != target.get('name')
            or any(columns.get(name) != target.get('columnIds', {}).get(name)
                   or not columns.get(name) for name in ('ID', 'ACV', 'Probability'))
            or target['columnIds']['ID'] != 'ID'
            or any(element.get(key) or target.get(key) for key in ('filters', 'groupings'))):
        raise ValueError('actual unfiltered input source/column binding required')
    native_id_formula = target.get('nativeIdFormula')
    if native_id_formula is not None and (not isinstance(native_id_formula, str)
            or not re.fullmatch(r'\[[^/\]]+/ID\]', native_id_formula)):
        raise ValueError('native source ID must be an explicitly observed direct ID projection')
    input_placements = _placement(snapshot['contents'], target['elementId'])
    if len(input_placements) != 1 or input_placements[0][0] != target.get('pageId'):
        raise ValueError('input page binding changed')
    pages = [page for page in snapshot['contents']['pages'] if page['id'] == 'sales-overview']
    layout_pages = [page for page in _layout(snapshot['contents']).iter('Page')
                    if page.get('id') == 'sales-overview']
    if len(pages) != 1 or len(layout_pages) != 1:
        raise ValueError('one existing Sales overview page required')
    return elements, layout_pages[0]


def _formula(owner, target, ids):
    reference = lambda name: '[' + target['name'] + '/' + name + ']'
    amount = ' * '.join(reference(name) for name in owner['numericDependencies'])
    membership = 'In(' + reference('ID') + ', ' + ', '.join(json.dumps(value) for value in ids) + ')'
    return f'Coalesce(SumIf({amount}, {membership}), 0)'


def _verify_owner(owner, binding, elements, snapshot, input_id):
    element = elements.get(binding.get('elementId'), {})
    value_id = binding.get('columnId')
    column = next((c for c in element.get('columns', []) if c.get('id') == value_id), {})
    if (element.get('kind') != 'kpi-chart' or element.get('name') != owner['name']
            or element.get('source') != {'kind': 'table', 'elementId': input_id}
            or element.get('value') != {'columnId': value_id}
            or not value_id or column.get('formula') != owner['formula']
            or column.get('name') != owner['name']
            or column.get('format') != {'kind': 'number', 'formatString': owner['formatPattern']}
            or any(element.get(key) for key in ('groupings', 'groups', 'filters', 'timeline',
                                                'comparison', 'dateGranularity', 'topN'))
            or _placement(snapshot['contents'], element.get('id')) != [('sales-overview', owner['frame'])]):
        raise ValueError('actual ungrouped scalar formula/source/value/placement readback mismatch')
    return element


def _program(owners, target, page_ids, plan_hash):
    if not owners:
        return ''
    lines = [f'# nativeScalarPlanHash={plan_hash}',
             '# Re-query and replan after interruption; native append only.',
             '# ADMISSION FENCE: the authorized external controller must admit the run and inverse first.',
             "raise RuntimeError('external run receipt and guarded native inverse admission required')",
             'from sigma.page import Page',
             'from sigma.inputTable import InputTable, InputTableIDColumn',
             'from sigma.column import FormulaColumn',
             "host = sigma.workbook.get_page('sales-overview')",
             "if not isinstance(host, Page): raise RuntimeError('expected overview page')",
             'observed_ids = []', 'for existing in host.list_elements():',
             '    observed_ids.append(existing.id())',
             f'if sorted(observed_ids) != {sorted(page_ids)!r}:',
             "    raise RuntimeError('overview inventory changed; re-query and replan')",
             f"source_host = sigma.workbook.get_page({target['pageId']!r})",
             f"source = source_host.get_element({target['elementId']!r})",
             "if not isinstance(source, InputTable): raise RuntimeError('expected existing input')",
             "source_id = source.get_column('ID')"]
    if target.get('nativeIdFormula'):
        lines += [f"if not isinstance(source_id, FormulaColumn) or source_id.formula() != {target['nativeIdFormula']!r}:",
                  "    raise RuntimeError('native source ID projection changed')"]
    else:
        lines += ["if not isinstance(source_id, InputTableIDColumn): raise RuntimeError('expected physical system ID')",
                  "if source_id.kind() != 'row_id': raise RuntimeError('expected system row identity')"]
    for name in ('ACV', 'Probability'):
        lines += [f"if source.get_column({target['columnIds'][name]!r}).name() != {name!r}:",
                  "    raise RuntimeError('input column binding changed')"]
    for n, owner in enumerate(owners):
        var = f'k{n}'
        lines += [f'{var} = host.add_kpi(source.to_source())',
                  f"print({(owner['sourceId'] + ':elementId=')!r} + {var}.id())",
                  f'{var}.set_name({owner["name"]!r})',
                  f'{var}.set_frame({", ".join(map(str, owner["frame"]))})',
                  f"if {var}.has_column('ID'):",
                  f"    id{n} = {var}.get_column_by_name('ID')", 'else:',
                  f"    id{n} = {var}.add_source_column('ID')",
                  f"print({(owner['sourceId'] + ':projectionColumnId=')!r} + id{n}.id())",
                  f"if id{n}.formula() != {('[' + target['name'] + '/ID]')!r}:",
                  "    raise RuntimeError('inherited physical ID source changed')"]
        operands = [f"sigma.formula.as_num({var}.get_column_by_name({name!r}))"
                    for name in owner['numericDependencies']]
        value = operands[0] if len(operands) == 1 else f'sigma.formula.multiply({", ".join(operands)})'
        membership = f'sigma.formula.in_(id{n}, ' + ', '.join(
            f'sigma.formula.const_text({value!r})' for value in owner['physicalIds']) + ')'
        expression = f'sigma.formula.coalesce(sigma.formula.sum_if({value}, {membership}), sigma.formula.const_int(0))'
        lines += [f'{var}.set_value(({expression}, {owner["name"]!r}))',
                  f'v{n} = {var}.value()',
                  f"if v{n} is None: raise RuntimeError('assigned scalar value unavailable')",
                  f"print({(owner['sourceId'] + ':columnId=')!r} + v{n}.id())",
                  f'v{n}.set_format_pattern({owner["formatPattern"]!r})']
    return '\n'.join(lines) + '\n'


def plan_native_scalar_append(contract, identity, snapshot, *, ownership=None):
    """Plan missing KPI appends, or verify exact owned nodes and emit a no-op.

    A same-name node without a prior assigned readback is a collision. Persist
    the plan and capture returned IDs, then bind against a fresh full readback.
    Inventory guards prevent replay from appending duplicates. They are not CAS.
    """
    _validate_contract(contract)
    elements, page = _validate_snapshot(snapshot)
    ids, scope_hash = _scope(contract, identity, snapshot)
    target = snapshot['input']
    owners = [{**copy.deepcopy(item), 'formula': _formula(item, target, ids), 'physicalIds': ids}
              for item in contract['owners']]
    prior = (ownership or {}).get('targets', {})
    if ownership and (ownership.get('sourceHash') != contract['sourceHash']
                      or ownership.get('contractHash') != contract['contractHash']
                      or ownership.get('scopeHash') != scope_hash
                      or ownership.get('documentId') != snapshot['documentId']
                      or set(prior) - {owner['sourceId'] for owner in owners}):
        raise ValueError('stale/foreign scalar ownership receipt')
    if len({binding.get('elementId') for binding in prior.values()}) != len(prior):
        raise ValueError('scalar ownership elements are not distinct')
    missing = []
    for owner in owners:
        binding = prior.get(owner['sourceId'])
        if binding:
            if (binding.get('sourceFormulaHash') != owner['sourceFormulaHash']
                    or binding.get('documentId') != snapshot['documentId']
                    or type(binding.get('version')) is not int
                    or binding['version'] > snapshot['version']):
                raise ValueError('stale scalar ownership binding')
            element = _verify_owner(owner, binding, elements, snapshot, target['elementId'])
            if binding.get('elementHash') != _digest(element):
                raise ValueError('owned scalar element hash changed')
        elif any(element.get('name') == owner['name'] or element['id'] == owner['logicalElementId']
                 for element in elements.values()):
            raise ValueError('scalar collision: assigned native ownership readback required')
        else:
            missing.append(owner)
    # Native append has no authority to move unrelated elements out of the way.
    for owner in missing:
        x, y, width, height = owner['frame']
        for existing in page.iter('Element'):
            for _, frame in _placement(snapshot['contents'], existing.get('elementId')):
                ex, ey, ew, eh = frame
                if x < ex+ew and ex < x+width and y < ey+eh and ey < y+height:
                    raise ValueError('scalar append placement collides with existing content')
    result = {'schema': 'native-source-scalar-append-v1',
              'mode': 'append-native-sdk' if missing else 'reconcile-existing',
              'sourceHash': contract['sourceHash'], 'contractHash': contract['contractHash'],
              'documentId': snapshot['documentId'], 'version': snapshot['version'],
              'scopeHash': scope_hash, 'owners': owners, 'missingSourceIds': [o['sourceId'] for o in missing],
              'emittedSourceIds': [o['sourceId'] for o in missing[:1]],
              'priorTargets': copy.deepcopy(prior), 'before': copy.deepcopy(snapshot),
              'executionStatus': 'BLOCKED_EXTERNAL_RUN_RECEIPT_AND_INVERSE_REQUIRED',
              'verification': 'UNRESOLVED', 'productionStatus': 'BLOCKED'}
    result['planHash'] = _digest(result)
    result['program'] = _program(missing[:1], target, [e.get('elementId') for e in page.iter('Element')],
                                 result['planHash'])
    return result


def _without_added(contents, added):
    result = copy.deepcopy(contents)
    result['elements'] = [element for element in result['elements'] if element['id'] not in added]
    root = _layout(result)
    if root.text is not None and root.text.strip(' \t\r\n'):
        raise ValueError('scalar layout has authored text before its first node')
    for parent in root.iter():
        for child in list(parent):
            if child.tag == 'Element' and child.get('elementId') in added:
                if child.tail is not None and child.tail.strip(' \t\r\n'):
                    raise ValueError('scalar append added authored layout text')
                parent.remove(child)
    # Indentation changes carry no layout meaning; retain all authored text.
    for node in root.iter():
        if node.text is not None and not node.text.strip(' \t\r\n'):
            node.text = None
        if node.tail is not None and not node.tail.strip(' \t\r\n'):
            node.tail = None
    result['layout'] = ''.join(ET.tostring(child, encoding='unicode') for child in root)
    return result


def bind_native_scalar_owners(plan, snapshot, assignments):
    """Accept observed assigned IDs, never logical IDs or runtime acceptance.

    Partial assigned readback is retained for a safe next plan. Unknown outcomes
    without returned IDs fail closed on the name collision rather than guessing.
    """
    payload = {k: v for k, v in plan.items() if k not in ('planHash', 'program')}
    if plan.get('planHash') != _digest(payload):
        raise ValueError('native scalar append plan changed')
    elements, _ = _validate_snapshot(snapshot)
    before = plan['before']
    if (snapshot['documentId'] != plan['documentId'] or snapshot['version'] <= plan['version']
            or snapshot['nativeActions'] != before['nativeActions']
            or snapshot['input'] != before['input']):
        raise ValueError('scalar readback changed unrelated document/input/native action state')
    allowed = set(plan['emittedSourceIds'])
    if not assignments or set(assignments) - allowed:
        raise ValueError('assigned scalar readback scope changed')
    prior = copy.deepcopy(plan['priorTargets'])
    ids = [binding.get('elementId') for binding in [*prior.values(), *assignments.values()]]
    if len(set(ids)) != len(ids) or any(not value for value in ids):
        raise ValueError('assigned scalar element IDs must be distinct')
    added = {binding['elementId'] for binding in assignments.values()}
    if added & set(_elements(before['contents'])):
        raise ValueError('assigned scalar ID was already present before append')
    if _without_added(snapshot['contents'], added) != _without_added(before['contents'], set()):
        raise ValueError('scalar append changed unrelated contents or has unassigned additions')
    for owner in plan['owners']:
        binding = assignments.get(owner['sourceId']) or prior.get(owner['sourceId'])
        if binding is None:
            continue
        element = _verify_owner(owner, binding, elements, snapshot, before['input']['elementId'])
        prior[owner['sourceId']] = {
            'kind': 'native-formula', 'documentId': snapshot['documentId'], 'version': snapshot['version'],
            'elementId': binding['elementId'], 'columnId': binding['columnId'],
            'sourceFormulaHash': owner['sourceFormulaHash'], 'elementHash': _digest(element)}
    return {'sourceHash': plan['sourceHash'], 'contractHash': plan['contractHash'],
            'scopeHash': plan['scopeHash'], 'documentId': plan['documentId'], 'targets': prior,
            'verification': 'UNRESOLVED', 'productionStatus': 'BLOCKED'}
