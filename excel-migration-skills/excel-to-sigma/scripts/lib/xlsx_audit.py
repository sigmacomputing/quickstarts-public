"""Fail-loud capability inventory for OOXML Excel workbooks.

The audit does not claim that every inventoried feature is converted. It
classifies each feature into an automatic path or an explicit review route so a
migration cannot silently discard workbook behavior.
"""
from __future__ import annotations

import json
import re
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}
VOLATILE_FUNCTIONS = {
    "CELL", "INDIRECT", "NOW", "OFFSET", "RAND", "RANDBETWEEN", "TODAY",
}
KNOWN_FUNCTIONS = {
    "ABS", "AND", "AVERAGE", "AVERAGEIF", "AVERAGEIFS", "CEILING", "CHOOSE",
    "COLUMN", "COLUMNS", "CONCAT", "CONCATENATE", "COUNT", "COUNTA", "COUNTBLANK",
    "COUNTIF", "COUNTIFS", "DATE", "DATEVALUE", "DAY", "EDATE", "EOMONTH", "ERROR.TYPE",
    "FILTER", "FIND", "FLOOR", "HLOOKUP", "IF", "IFERROR", "IFS", "INDEX", "INT",
    "ISBLANK", "ISERROR", "ISNA", "LARGE", "LEFT", "LEN", "LET", "LOOKUP", "LOWER",
    "MATCH", "MAX", "MID", "MIN", "MOD", "MONTH", "NOT", "OR", "PROPER", "RANK",
    "REPLACE", "RIGHT", "ROUND", "ROUNDDOWN", "ROUNDUP", "ROW", "ROWS", "SEARCH",
    "SEQUENCE", "SMALL", "SORT", "SORTBY", "SUBSTITUTE", "SUM", "SUMIF", "SUMIFS",
    "SUMPRODUCT", "TEXT", "TEXTJOIN", "TIME", "TRIM", "UNIQUE", "UPPER", "VALUE",
    "VLOOKUP", "WEEKDAY", "XLOOKUP", "XMATCH", "YEAR",
}
DYNAMIC_FUNCTIONS = {
    "CHOOSECOLS", "CHOOSEROWS", "DROP", "EXPAND", "FILTER", "HSTACK",
    "RANDARRAY", "SEQUENCE", "SORT", "SORTBY", "TAKE", "TOCOL", "TOROW",
    "UNIQUE", "VSTACK", "WRAPCOLS", "WRAPROWS",
}
EXTERNAL_REF_RE = re.compile(r"\[[^\]]+\.(?:xlsx|xlsm|xlsb|xls)\]", re.I)
INDEXED_EXTERNAL_REF_RE = re.compile(r"\[\d+\][^!]*!", re.I)
ERROR_LITERAL_RE = re.compile(
    r"#(?:N/A|DIV/0!|GETTING_DATA|[A-Z][A-Z0-9_.]*(?:!|\?))",
    re.I,
)


def _formula_code(expression):
    """Mask literals, quoted identifiers, and bracketed references in place.

    Apostrophes inside structured references escape the following character;
    they do not open a quoted sheet identifier. Keeping the original offsets
    lets the tokenizer normalization remove only actual spill operators.
    """
    result = []
    quote = None
    bracket_depth = 0
    index = 0
    while index < len(expression):
        char = expression[index]
        if quote:
            if char == quote:
                if index + 1 < len(expression) and expression[index + 1] == quote:
                    result.extend((" ", " "))
                    index += 2
                    continue
                quote = None
            result.append(" ")
        elif bracket_depth:
            if char == "'" and index + 1 < len(expression):
                result.extend((" ", " "))
                index += 2
                continue
            if char == "[":
                bracket_depth += 1
            elif char == "]":
                bracket_depth -= 1
            result.append(" ")
        elif char in {"'", '"'}:
            quote = char
            result.append(" ")
        elif char == "[":
            bracket_depth = 1
            result.append(" ")
        else:
            result.append(char)
        index += 1
    return "".join(result)


def _has_spill_reference(expression):
    """Return whether code contains Excel's postfix spill operator."""
    return "#" in ERROR_LITERAL_RE.sub("", _formula_code(expression))


def _without_spill_operators(expression):
    """Blank postfix spill operators so openpyxl's legacy tokenizer can parse."""
    code = ERROR_LITERAL_RE.sub(
        lambda match: " " * len(match.group()), _formula_code(expression),
    )
    return "".join(
        " " if code[index] == "#" else char
        for index, char in enumerate(expression)
    )


def _formula_tokens(expression):
    formula = expression if expression.startswith("=") else f"={expression}"
    formula = ERROR_LITERAL_RE.sub("0", formula)
    formula = _without_spill_operators(formula)
    try:
        return Tokenizer(formula).items
    except TokenizerError:
        return None


def _function_name(token_value):
    parts = token_value.rstrip("(").split(".")
    while parts and parts[0].lower() in {"_xlfn", "_xlws"}:
        parts.pop(0)
    return ".".join(parts).upper()


def _is_external_range(range_operand):
    return "!" in range_operand and bool(
        EXTERNAL_REF_RE.search(range_operand)
        or INDEXED_EXTERNAL_REF_RE.search(range_operand)
    )


def _analyze_expression(expression):
    tokens = _formula_tokens(expression)
    tokenization_failed = tokens is None
    tokens = tokens or []
    functions = [
        _function_name(token.value)
        for token in tokens
        if token.type == "FUNC" and token.subtype == "OPEN"
    ]
    if not tokens:
        # The legacy tokenizer rejects some valid structured-reference escapes.
        # Inventory functions from masked code even when it cannot emit tokens.
        functions = [
            _function_name(match.group(1))
            for match in re.finditer(
                r"(?<![\w.])([A-Z_][A-Z0-9_.]*)\s*\(",
                _formula_code(expression),
                re.I,
            )
        ]
    range_operands = [
        token.value
        for token in tokens
        if token.type == "OPERAND" and token.subtype == "RANGE"
    ]
    error_code = _formula_code(expression)
    return {
        "tokenizationFailed": tokenization_failed,
        "functions": functions,
        "rangeOperands": range_operands,
        "external": any(_is_external_range(item) for item in range_operands),
        "dynamic": bool(set(functions) & DYNAMIC_FUNCTIONS)
        or _has_spill_reference(expression),
        "errors": ERROR_LITERAL_RE.findall(error_code),
    }


def _formula_text(value):
    if isinstance(value, ArrayFormula):
        return value.text or ""
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def _truthy_ooxml_flag(value):
    return value is True or value == 1 or value == "1"


def _defined_name_type(item):
    try:
        return item.type
    except TokenizerError:
        return "FORMULA"


def _issue(code, message, route, evidence=None, severity="review"):
    item = {
        "code": code,
        "severity": severity,
        "route": route,
        "message": message,
    }
    if evidence:
        item["evidence"] = evidence
    return item


def _zip_inventory(path):
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    lowered = [name.lower() for name in names]

    def any_path(*needles):
        return any(any(needle in name for needle in needles) for name in lowered)

    return {
        "vbaProject": any(name.endswith("vbaproject.bin") for name in lowered),
        "externalLinks": sum("/externallinks/" in name for name in lowered),
        "connections": any(name.endswith("/connections.xml") for name in lowered),
        "queryTables": sum("/querytables/" in name for name in lowered),
        "powerQuery": any_path("/queries/", "datamashup"),
        "pivotParts": sum("/pivottables/" in name for name in lowered),
        "slicerParts": sum(
            "/slicers/" in name or "/slicercaches/" in name for name in lowered
        ),
        "timelineParts": sum(
            "/timelines/" in name or "/timelinecaches/" in name for name in lowered
        ),
        "activeXParts": sum("/activex/" in name for name in lowered),
        "embeddedObjects": sum("/embeddings/" in name for name in lowered),
        "customXmlParts": sum(name.startswith("customxml/") and name.endswith(".xml")
                              for name in lowered),
        "webExtensions": sum("/webextensions/" in name for name in lowered),
    }


def _workbook_inventory(path):
    suffix = Path(path).suffix.lower()
    wb = openpyxl.load_workbook(
        path,
        data_only=False,
        keep_vba=suffix == ".xlsm",
        read_only=False,
    )
    functions = Counter()
    formulas = []
    formula_expressions = []
    array_formula_cells = []
    data_table_formula_cells = []
    external_formula_cells = []
    dynamic_formula_cells = []
    formula_error_cells = []
    unparsed_formula_cells = []
    external_defined_names = []
    dynamic_defined_names = []
    formula_error_defined_names = []
    unparsed_defined_names = []
    tables = []
    charts = []
    validations = 0
    conditional_formats = 0
    pivots = 0
    protected_sheets = []
    comments = 0
    images = 0

    for ws in wb.worksheets:
        tables.extend({
            "sheet": ws.title,
            "name": table.name,
            "ref": table.ref,
        } for table in ws.tables.values())
        charts.extend({
            "sheet": ws.title,
            "type": type(chart).__name__,
            "title": str(getattr(chart, "title", "") or ""),
        } for chart in getattr(ws, "_charts", []))
        validations += len(getattr(getattr(ws, "data_validations", None), "dataValidation", []) or [])
        conditional_formats += len(getattr(ws, "conditional_formatting", []) or [])
        pivots += len(getattr(ws, "_pivots", []) or [])
        if getattr(ws.protection, "sheet", False):
            protected_sheets.append(ws.title)
        images += len(getattr(ws, "_images", []) or [])

        for row in ws.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    comments += 1
                value = cell.value
                cell_name = f"{ws.title}!{cell.coordinate}"
                if isinstance(value, DataTableFormula):
                    formulas.append(cell_name)
                    data_table_formula_cells.append({
                        "cell": cell_name,
                        "ref": value.ref,
                        "twoVariable": _truthy_ooxml_flag(value.dt2D),
                        "rowInput": value.r1,
                        "columnInput": value.r2,
                    })
                    continue
                if cell.data_type == "e":
                    formula_error_cells.append(cell_name)
                    continue
                expression = _formula_text(value)
                if expression is None:
                    continue
                formulas.append(cell_name)
                if isinstance(value, ArrayFormula):
                    array_formula_cells.append({
                        "cell": cell_name,
                        "ref": value.ref,
                    })
                analysis = _analyze_expression(expression)
                formula_expressions.append(
                    (ws.title, cell_name, analysis["rangeOperands"])
                )
                functions.update(analysis["functions"])
                if analysis["tokenizationFailed"]:
                    unparsed_formula_cells.append(cell_name)
                if analysis["external"]:
                    external_formula_cells.append(cell_name)
                if analysis["dynamic"]:
                    dynamic_formula_cells.append(cell_name)
                if analysis["errors"]:
                    formula_error_cells.append(cell_name)

    defined_names = []
    global_defined_names = {}
    local_defined_names = {}
    scoped_defined_names = [(item, None) for item in wb.defined_names.values()]
    for ws in wb.worksheets:
        scoped_defined_names.extend(
            (item, ws.title) for item in ws.defined_names.values()
        )
    for item, worksheet_scope in scoped_defined_names:
        local_sheet_id = getattr(item, "localSheetId", None)
        if worksheet_scope is not None:
            scope = worksheet_scope
            label = f"{scope}!{item.name}"
        elif local_sheet_id is None:
            scope = "workbook"
            label = item.name
        elif 0 <= local_sheet_id < len(wb.worksheets):
            scope = wb.worksheets[local_sheet_id].title
            label = f"{scope}!{item.name}"
        else:
            scope = f"worksheet-index-{local_sheet_id}"
            label = f"{scope}!{item.name}"
        expression = getattr(item, "attr_text", None) or ""
        analysis = _analyze_expression(expression)
        functions.update(analysis["functions"])
        if analysis["tokenizationFailed"]:
            unparsed_defined_names.append(label)
        if analysis["external"]:
            external_defined_names.append(label)
        if analysis["dynamic"]:
            dynamic_defined_names.append(label)
        if analysis["errors"]:
            formula_error_defined_names.append(label)
        inventory_item = {
            "name": item.name,
            "type": _defined_name_type(item),
            "hidden": bool(getattr(item, "hidden", False)),
            "scope": scope,
            "attrText": expression,
            "referencedBy": [],
        }
        defined_names.append(inventory_item)
        name_key = item.name.casefold()
        if scope == "workbook":
            global_defined_names[name_key] = inventory_item
        else:
            local_defined_names[(scope.casefold(), name_key)] = inventory_item

    for sheet_name, cell_name, range_operands in formula_expressions:
        for range_operand in range_operands:
            reference = range_operand.lstrip("@")
            target = None
            if "!" in reference:
                qualifier, name = reference.rsplit("!", 1)
                if "[" not in qualifier:
                    qualifier = qualifier.strip("'").replace("''", "'")
                    target = local_defined_names.get(
                        (qualifier.casefold(), name.casefold())
                    )
            else:
                name_key = reference.casefold()
                target = local_defined_names.get(
                    (sheet_name.casefold(), name_key)
                ) or global_defined_names.get(name_key)
            if target is not None and cell_name not in target["referencedBy"]:
                target["referencedBy"].append(cell_name)

    workbook_protected = bool(
        getattr(wb.security, "lockStructure", False)
        or getattr(wb.security, "workbookPassword", None)
        or getattr(wb.security, "revisionsPassword", None)
    )
    return {
        "sheets": [
            {"name": ws.title, "state": ws.sheet_state, "rows": ws.max_row, "columns": ws.max_column}
            for ws in wb.worksheets
        ],
        "tables": tables,
        "formulaCells": len(formulas),
        "arrayFormulaCells": array_formula_cells,
        "dataTableFormulaCells": data_table_formula_cells,
        "functions": dict(sorted(functions.items())),
        "charts": charts,
        "pivotTables": pivots,
        "dataValidations": validations,
        "conditionalFormattingRules": conditional_formats,
        "definedNames": defined_names,
        "protectedSheets": protected_sheets,
        "workbookProtected": workbook_protected,
        "comments": comments,
        "images": images,
        "externalFormulaCells": external_formula_cells,
        "dynamicFormulaCells": dynamic_formula_cells,
        "formulaErrorCells": formula_error_cells,
        "unparsedFormulaCells": unparsed_formula_cells,
        "externalDefinedNames": external_defined_names,
        "dynamicDefinedNames": dynamic_defined_names,
        "formulaErrorDefinedNames": formula_error_defined_names,
        "unparsedDefinedNames": unparsed_defined_names,
    }


def audit_workbook(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        return {
            "source": str(path),
            "status": "BLOCKED",
            "inventory": {},
            "issues": [_issue(
                "LEGACY_OR_UNSUPPORTED_FORMAT",
                f"{suffix or 'extensionless'} files are not OOXML workbooks",
                "convert-to-xlsx-or-xlsm",
                severity="blocker",
            )],
            "routing": {"automatic": [], "review": [], "blocked": ["convert-to-xlsx-or-xlsm"]},
        }
    if not path.exists():
        raise FileNotFoundError(path)

    package = _zip_inventory(path)
    workbook = _workbook_inventory(path)
    inventory = {**workbook, "package": package}
    issues = []

    function_names = set(workbook["functions"])
    unknown = sorted(function_names - KNOWN_FUNCTIONS - VOLATILE_FUNCTIONS)
    volatile = sorted(function_names & VOLATILE_FUNCTIONS)
    if unknown:
        issues.append(_issue(
            "UNKNOWN_FORMULAS",
            "Formula functions outside the translator allowlist require a disposition.",
            "formula-review-or-python",
            unknown,
        ))
    if volatile:
        issues.append(_issue(
            "VOLATILE_FORMULAS",
            "Volatile/location-dependent formulas need an explicit Sigma replacement.",
            "formula-review-or-python",
            volatile,
        ))
    if workbook["unparsedFormulaCells"] or workbook["unparsedDefinedNames"]:
        evidence = list(workbook["unparsedFormulaCells"])
        evidence.extend(
            f"defined-name:{name}" for name in workbook["unparsedDefinedNames"]
        )
        issues.append(_issue(
            "UNPARSED_FORMULAS",
            "Formula syntax unsupported by the audit tokenizer requires review; "
            "recovered function names do not establish complete reference coverage.",
            "formula-review-or-python",
            evidence,
        ))
    if (workbook["externalFormulaCells"] or workbook["externalDefinedNames"]
            or package["externalLinks"]):
        evidence = list(workbook["externalFormulaCells"])
        evidence.extend(
            f"defined-name:{name}" for name in workbook["externalDefinedNames"]
        )
        if not evidence:
            evidence.append(f"{package['externalLinks']} OOXML link part(s)")
        issues.append(_issue(
            "EXTERNAL_WORKBOOK_LINKS",
            "External workbook references must be landed as governed Sigma sources.",
            "source-mapping-review",
            evidence,
        ))
    if workbook["dynamicFormulaCells"] or workbook["dynamicDefinedNames"]:
        evidence = list(workbook["dynamicFormulaCells"])
        evidence.extend(
            f"defined-name:{name}" for name in workbook["dynamicDefinedNames"]
        )
        issues.append(_issue(
            "DYNAMIC_ARRAY_FORMULAS",
            "Spill formulas require a grain-aware table or Python translation.",
            "formula-review-or-python",
            evidence,
        ))
    if workbook["formulaErrorCells"] or workbook["formulaErrorDefinedNames"]:
        evidence = list(workbook["formulaErrorCells"])
        evidence.extend(
            f"defined-name:{name}" for name in workbook["formulaErrorDefinedNames"]
        )
        issues.append(_issue(
            "FORMULA_ERRORS",
            "Formula error literals require correction or an explicit business disposition.",
            "formula-review-or-python",
            evidence,
        ))
    if workbook["arrayFormulaCells"]:
        issues.append(_issue(
            "ARRAY_FORMULAS_PRESENT",
            "Array formula ranges require an explicit translation and parity check.",
            "array-formula-migration-review",
            [item["cell"] for item in workbook["arrayFormulaCells"]],
        ))
    if workbook["dataTableFormulaCells"]:
        issues.append(_issue(
            "WHAT_IF_DATA_TABLES_PRESENT",
            "What-if data tables require explicit input and scenario-model mapping.",
            "what-if-analysis-review",
            [item["cell"] for item in workbook["dataTableFormulaCells"]],
        ))
    if package["connections"] or package["queryTables"] or package["powerQuery"]:
        issues.append(_issue(
            "POWER_QUERY_OR_CONNECTIONS",
            "Power Query and workbook connections are source pipelines, not cell formulas.",
            "warehouse-or-api-source-review",
        ))
    if package["vbaProject"]:
        issues.append(_issue(
            "VBA_PROJECT",
            "VBA must pass macro-classify.py; no procedure may be silently dropped.",
            "macro-classifier",
        ))
    if workbook["charts"]:
        issues.append(_issue(
            "CHARTS_PRESENT",
            "Chart types, encodings, filters, and aggregation must be mapped and rendered.",
            "chart-migration-and-render-check",
            sorted({chart["type"] for chart in workbook["charts"]}),
        ))
    if workbook["pivotTables"] or package["pivotParts"]:
        issues.append(_issue(
            "PIVOT_TABLES_PRESENT",
            "Pivot rows, columns, values, filters, subtotals, and sorts require mapping.",
            "pivot-migration-and-parity-check",
        ))
    if package["slicerParts"] or package["timelineParts"]:
        issues.append(_issue(
            "SLICERS_OR_TIMELINES",
            "Slicers and timelines must become Sigma controls with equivalent targets.",
            "control-migration-review",
        ))
    if workbook["dataValidations"]:
        issues.append(_issue(
            "DATA_VALIDATION_PRESENT",
            "Translate supported list/range/file constraints and review custom formulas.",
            "input-validation-migration",
        ))
    if workbook["conditionalFormattingRules"]:
        issues.append(_issue(
            "CONDITIONAL_FORMATTING_PRESENT",
            "Conditional formatting semantics need a visual parity check.",
            "formatting-migration-review",
        ))
    if workbook["protectedSheets"] or workbook["workbookProtected"]:
        issues.append(_issue(
            "WORKBOOK_PROTECTION",
            "Excel protection is not RLS/CLS; capture its intent and map permissions explicitly.",
            "security-review",
        ))
    if package["activeXParts"] or package["webExtensions"]:
        issues.append(_issue(
            "ADDINS_OR_ACTIVEX",
            "ActiveX and web add-ins cannot be inferred as workbook formulas.",
            "external-application-review",
            severity="blocker",
        ))
    if package["embeddedObjects"]:
        issues.append(_issue(
            "EMBEDDED_OBJECTS",
            "Embedded OLE objects require extraction and an explicit destination.",
            "content-extraction-review",
        ))
    if package["customXmlParts"]:
        issues.append(_issue(
            "CUSTOM_XML",
            "Custom XML may contain add-in or business metadata and must be dispositioned.",
            "package-metadata-review",
        ))
    if workbook["comments"] or workbook["images"]:
        issues.append(_issue(
            "ANNOTATIONS_OR_IMAGES",
            "Comments, notes, and images require content-level migration review.",
            "content-migration-review",
            [f"comments={workbook['comments']}", f"images={workbook['images']}"],
        ))

    blockers = [item for item in issues if item["severity"] == "blocker"]
    status = "BLOCKED" if blockers else ("REVIEW_REQUIRED" if issues else "AUTO")
    automatic = [
        "formal-table-to-input-or-data-model" if workbook["tables"] else
        "cell-model-or-research-template-routing"
    ]
    routes = sorted({item["route"] for item in issues})
    return {
        "source": str(path),
        "status": status,
        "inventory": inventory,
        "issues": issues,
        "routing": {
            "automatic": automatic,
            "review": routes if not blockers else [
                item["route"] for item in issues if item["severity"] != "blocker"
            ],
            "blocked": sorted({item["route"] for item in blockers}),
        },
    }


def dumps(report):
    return json.dumps(report, indent=2, sort_keys=False)
