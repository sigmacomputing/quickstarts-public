"""Approved SDK display plans; code-rep input-column exports are not evidence."""
from __future__ import annotations

import ast
import copy
import hashlib
import json
import re

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

try:
    from .input_sync import _approved_feature
except ImportError:
    from input_sync import _approved_feature


def _hash(value):
    return hashlib.sha256(json.dumps(value,sort_keys=True,separators=(',',':')).encode()).hexdigest()


def source_format_intent(raw, column_type):
    """Translate bounded Excel semantics to documented raw d3/strftime formats."""
    if column_type in {'text','boolean'} or raw in {'General','@'}:
        return {'kind':'automatic'},'EXACT',None
    if column_type=='datetime':
        pattern={'yyyy-mm-dd':'%Y-%m-%d','yyyy-mm-dd hh:mm:ss':'%Y-%m-%d %H:%M:%S'}.get(raw)
        return (({'pattern':pattern},'EXACT',None) if pattern else
                (None,'UNRESOLVED','source date pattern has no verified exact SDK mapping'))
    sections=raw.split(';')
    positive=sections[0]
    if not re.fullmatch(r'\$?(?:#,##)?0(?:\.0+)?%?',positive):
        return None,'UNRESOLVED','source numeric pattern has no verified exact SDK mapping'
    decimals=len(positive.split('.')[1].rstrip('%')) if '.' in positive else 0
    negative='(' if len(sections)>1 and '(' in sections[1] else ''
    pattern=negative+('$' if '$' in positive else '')+(',' if ',' in positive else '')+f'.{decimals}'+('%' if positive.endswith('%') else 'f')
    if len(sections)>1:
        return {'pattern':pattern},'PARTIAL','Excel accounting sections and zero-as-dash fidelity are unresolved; only supported currency/grouping/precision/negative-parentheses semantics are configured'
    return {'pattern':pattern},'EXACT',None


def currency_pattern(raw, precision):
    """Use a real d3 currency marker; currencySymbol metadata alone does not render it."""
    intent,_,_=source_format_intent(raw,'number')
    if intent and 'pattern' in intent and '$' in intent['pattern']:
        return intent['pattern']
    return f'$,.{precision}f'


def _system_column_contract(target):
    """Default Row ID, evidenced native ID, or an evidenced post-REST alias.

    FormulaColumn exposes formula()/element_id(), not a documented type() or
    system kind getter. Its row-ID identity/type comes from independent reviewed
    evidence bound into the complete approved targetBinding.
    """
    alias=target.get('systemColumn')
    if alias is None:
        return {'name':'Row ID','columnClass':'InputTableIDColumn','type':None,
                'systemKind':'row_id','formula':None,'elementId':target['elementId']}
    if isinstance(alias,dict) and alias.get('columnClass')=='InputTableIDColumn':
        fields={'id','name','columnClass','type','systemKind','formula','elementId','sdkEvidenceHash'}
        if (set(alias)!=fields or alias.get('name')!='ID'
                or not isinstance(alias.get('id'),str) or not alias['id']
                or alias['id']!=target.get('rowIdColumnId')
                or alias['id'] in target.get('columnIds',{}).values()
                or alias.get('type') is not None or alias.get('systemKind')!='row_id'
                or alias.get('formula') is not None or alias.get('elementId')!=target['elementId']
                or not re.fullmatch(r'[0-9a-f]{64}',str(alias.get('sdkEvidenceHash') or ''))):
            raise ValueError('native ID requires exact reviewed ID/name/class/kind/type/owner and SDK evidence')
        # Provenance is receipt-bound; it is not an SDK column getter field.
        return {k:copy.deepcopy(v) for k,v in alias.items() if k!='sdkEvidenceHash'}
    keys={'id','name','columnClass','type','systemKind','formula','elementId','rowIdBinding'}
    identity=alias.get('rowIdBinding') if isinstance(alias,dict) else None
    if (not isinstance(alias,dict) or set(alias)!=keys
            or alias['id']!='ID' or alias['name']!='ID'
            or alias['columnClass']!='FormulaColumn' or alias['type'] is not None
            or alias['systemKind'] is not None or alias['elementId']!=target['elementId']
            or alias['formula']!='[New Input Table/ID]'
            or not isinstance(identity,dict)
            or set(identity)!={'kind','sourceColumn','valueType','evidenceHash'}
            or identity.get('kind')!='row_id' or identity.get('sourceColumn')!='ID'
            or identity.get('valueType')!='text'
            or not re.fullmatch(r'[0-9a-f]{64}',str(identity.get('evidenceHash') or ''))
            or alias['id'] in target['columnIds'].values()):
        raise ValueError('approved system alias requires the exact observed formula/owner/ID and reviewed row-ID identity binding')
    return copy.deepcopy(alias)


def _editable_column_contract(manifest, target, column):
    """A list conversion is legal only with reviewed SDK + source-list identity.

    sdkEvidenceHash names the canonical raw SDK capture (or column observation)
    reviewed by the controller. The receipt hashes the entire descriptor; this
    function verifies its source semantics, not server authenticity.
    """
    contract=(target.get('columnContracts') or {}).get(column['name'])
    if contract is None:
        return {'id':target['columnIds'][column['name']],
                'columnClass':'LiteralEditableColumn','type':column['type']}
    fields={'id','columnClass','type','allowedValues','sourceValidationHash','sdkEvidenceHash'}
    if (not isinstance(contract,dict) or set(contract)!=fields
            or contract['id']!=target['columnIds'][column['name']]
            or contract['columnClass']!='SingleSelectEditableColumn'
            or contract['type']!=column['type']
            or not re.fullmatch('[0-9a-f]{64}',str(contract.get('sdkEvidenceHash') or ''))):
        raise ValueError('select conversion requires exact reviewed class/type/ID/evidence descriptor')
    source=column.get('sourceColumn') or ''
    if '!' not in source:raise ValueError('select conversion has no literal source column')
    sheet,letter=source.rsplit('!',1)
    workbook=load_workbook(manifest['source']['path'],data_only=False)
    try:
        try:
            from .presentation import _extract_validations
        except ImportError:
            from presentation import _extract_validations
        rules=_extract_validations(workbook, manifest['source']['path'])
    finally:workbook.close()
    matches=[]
    for rule in rules:
        low,_,high,_=range_boundaries(rule['range'])
        from openpyxl.utils.cell import column_index_from_string
        if (rule['sheet']==sheet and low<=column_index_from_string(letter)<=high
                and column['name'] in rule['fields'] and rule['status']=='supported'
                and rule['emission'] is not None and 'values' in rule['emission']):
            matches.append(rule)
    if len(matches)!=1 or _hash(matches[0])!=contract['sourceValidationHash']:
        raise ValueError('select conversion source validation hash/column is stale or ambiguous')
    expected=matches[0]['emission']['values']
    if contract['allowedValues']!=expected or any(type(a) is not type(b) for a,b in zip(contract['allowedValues'],expected)):
        raise ValueError('select conversion allowed values differ from parsed source validation')
    return copy.deepcopy(contract)


def build_input_display(plan,manifest,receipt,feature=None):
    if manifest.get('sourceHash')!=plan['source']['sha256']:
        raise ValueError('input display source hash is stale')
    architecture=manifest.get('planProjection',{}).get('sharedInputArchitecture',{})
    feature=feature or architecture.get('presentationInputFeature')
    spec,target,artifact=_approved_feature(manifest,receipt,feature,require_assigned=True)
    if manifest.get('target')!={'documentId':target['documentId'],'version':target['version']}:
        raise ValueError('input display target/version binding is stale')
    system_column=_system_column_contract(target)
    workbook=load_workbook(plan['source']['path'],data_only=False)
    try:
        records=[]
        prominent=['Scope']+[f'i_model_inputs_c{i}' for i in range(2,8)] if plan['workflow']=='research' else [c['name'] for c in spec['columns']]
        for column in spec['columns']:
            name=column['name']; source_cell=None; raw=None
            provenance=column.get('provenance','')
            if '!' in provenance and not provenance.startswith('source:'):
                source_cell=provenance
                sheet,address=source_cell.rsplit('!',1)
                raw=workbook[sheet][address].number_format
            elif column.get('sourceColumn'):
                sheet,letter=column['sourceColumn'].rsplit('!',1)
                _,first,_,last=range_boundaries(spec['sourceRange'])
                source_cell=f'{sheet}!{letter}{first+1}'
                formats={workbook[sheet][f'{letter}{row}'].number_format for row in range(first+1,last+1)}
                raw=next(iter(formats)) if len(formats)==1 else None
            if raw is None and column['type'] in {'text','boolean'}:
                raw='General'
            if raw is None:
                fmt,fidelity,reason=None,'UNRESOLVED','source column has missing or heterogeneous format bindings'
            else:
                fmt,fidelity,reason=source_format_intent(raw,column['type'])
            records.append({'id':target['columnIds'][name],'name':name,'visible':name in prominent,
                            **{k:v for k,v in _editable_column_contract(manifest,target,column).items() if k in {'columnClass','type','allowedValues'}},
                            'sourceCell':source_cell,'sourceFormat':raw,'format':fmt,
                            'fidelity':fidelity,'unresolvedReason':reason})
    finally:
        workbook.close()
    # Keep the intended prominent order independent of source-column enumeration.
    records.sort(key=lambda c:(0,prominent.index(c['name'])) if c['name'] in prominent else (1,c['name']))
    contract={'sourceHash':manifest['sourceHash'],'planHash':manifest['planHash'],
              'receiptHash':_hash(receipt),'artifactHash':artifact['artifactHash'],'feature':feature,
              'target':copy.deepcopy(target),'tableName':spec['name'],'columns':records,'systemColumn':system_column}
    contract_hash=_hash(contract)
    header=['from sigma.page import Page','from sigma.inputTable import InputTable, LiteralEditableColumn, InputTableIDColumn',
            'from sigma.column import FormulaColumn',f'# sourceHash={manifest["sourceHash"]}',f'# planHash={manifest["planHash"]}',
            f'# displayContractHash={contract_hash}',f'# execute only in document={target["documentId"]} version={target["version"]}',
            f'host = sigma.workbook.get_page({target["pageId"]!r})',
            "if not isinstance(host, Page): raise RuntimeError('expected approved page')",
            f'table = host.get_element({target["elementId"]!r})',
            "if not isinstance(table, InputTable): raise RuntimeError('expected approved input table')",
            f"if table.name() != {spec['name']!r}: raise RuntimeError('input table name is stale')"]
    if any(c['columnClass']=='SingleSelectEditableColumn' for c in records):
        header.insert(0,'from sigma.inputTable import SingleSelectEditableColumn')
    for i,c in enumerate(records):
        header += [f'c{i} = table.get_column({c["id"]!r})',
                   f"if not isinstance(c{i}, {c['columnClass']}): raise RuntimeError('expected approved editable column class')",
                   f"if c{i}.type() != {c['type']!r}: raise RuntimeError('approved column type is stale')",
                   f"if c{i}.name() != {c['name']!r}: raise RuntimeError('approved column name is stale')",
                   f"if c{i}.element_id() != {target['elementId']!r}: raise RuntimeError('foreign editable column owner')"]
    for i,c in enumerate(records):
        if c['columnClass']=='SingleSelectEditableColumn':
            header.append(f"if c{i}.allowed_values() != {c['allowedValues']!r}: raise RuntimeError('source-list options changed')")
    # Validate the complete set before any setter, in both SDK programs.
    header += [f'approved_ids = {[c["id"] for c in records]!r}',
               'all_columns = table.list_columns()',
               'observed_ids = {}', 'system_columns = []',
               'for column in all_columns:',
               "    if column.id() in observed_ids: raise RuntimeError('duplicate input columns')",
               '    observed_ids[column.id()] = True',
               '    if column.id() not in approved_ids: system_columns.append(column)',
               'for approved_id in approved_ids:',
               "    if approved_id not in observed_ids: raise RuntimeError('missing approved input columns')",
               "if len(system_columns) != 1: raise RuntimeError('expected exactly one system row ID column')",
               'system_column = system_columns[0]',
               f"if system_column.element_id() != {target['elementId']!r}: raise RuntimeError('foreign system column owner')",
               f"if system_column.name() != {system_column['name']!r}: raise RuntimeError('expected approved system column name')"]
    if system_column['columnClass']=='FormulaColumn':
        header += ["if not isinstance(system_column, FormulaColumn): raise RuntimeError('expected explicitly approved system alias class')",
                   f"if system_column.id() != {system_column['id']!r}: raise RuntimeError('foreign system alias ID')",
                   f"if system_column.formula() != {system_column['formula']!r}: raise RuntimeError('foreign system alias formula')"]
    else:
        header += ["if not isinstance(system_column, InputTableIDColumn): raise RuntimeError('expected system row ID column class')",
                   "if system_column.kind() != 'row_id': raise RuntimeError('expected system row_id kind')"]
    if system_column['columnClass']=='InputTableIDColumn' and 'id' in system_column:
        header.append(f"if system_column.id() != {system_column['id']!r}: raise RuntimeError('foreign native row ID column')")
    program=list(header)
    program.append('table.set_visible_columns(['+', '.join(f'c{i}' for i,c in enumerate(records) if c['visible'])+'])')
    for i,c in enumerate(records):
        fmt=c['format']
        if not fmt: continue
        if 'pattern' in fmt: program.append(f'c{i}.set_format_pattern({fmt["pattern"]!r})')
        elif 'decimals' in fmt: program.append(f'c{i}.set_format({fmt["kind"]!r}, decimals={fmt["decimals"]})')
        else: program.append(f'c{i}.set_format({fmt["kind"]!r})')
    program.append(f"print({{'source': 'sdk-display-execution', 'contractHash': {contract_hash!r}, 'elementId': table.id()}})")
    readback=[line for line in header if not line.startswith('# execute only')]+['observations = []','for column in table.list_columns():','    fmt = column.format()']
    has_select=any(c['columnClass']=='SingleSelectEditableColumn' for c in records)
    classes=(['SingleSelectEditableColumn'] if has_select else [])+['LiteralEditableColumn','FormulaColumn','InputTableIDColumn']
    for index,cls in enumerate(classes):
        fields=["'id': column.id()", "'name': column.name()", "'visible': column.is_visible()", "'elementId': column.element_id()",
            f"'columnClass': {cls!r}", "'type': "+('column.type()' if cls in {'LiteralEditableColumn','SingleSelectEditableColumn'} else 'None'),
            "'formula': "+('column.formula()' if cls=='FormulaColumn' else 'None'),
            "'systemKind': "+('column.kind()' if cls=='InputTableIDColumn' else 'None'),
            "'format': None if fmt is None else {'kind': fmt.kind(), 'decimals': fmt.decimal_places(), 'currency': fmt.currency(), 'pattern': fmt.pattern()}"]
        if has_select:fields.append("'allowedValues': "+('column.allowed_values()' if cls=='SingleSelectEditableColumn' else 'None'))
        readback += [('    if' if index==0 else '    elif')+f' isinstance(column, {cls}):', '        observations.append({'+', '.join(fields)+'})']
    readback += ["    else: raise RuntimeError('unreviewed input column class')",
        f"print({{'source': 'sdk-column-readback', 'contractHash': {contract_hash!r}, 'elementId': table.id(), 'columns': observations}})"]
    return {**contract,'contractHash':contract_hash,'status':'PENDING_SDK_READBACK',
            'executionBinding':{'workbookId':target['documentId'],'documentVersion':target['version']},
            'limitations':['SDK system aliases require explicit reviewed row-ID identity evidence in the receipt; no formula type getter or automatic alias-by-name inference is used.',
                          'SDK has no documented workbook-version CAS; caller must bind the execution context to this document/version and capture a fresh read-only SDK session after execution.',
                          'REST code-rep exports omit input display metadata and are not display readback evidence.',
                          'LiteralEditableColumn remains the default. SingleSelectEditableColumn requires an exact source-validation and SDK-evidence descriptor; other post-validation classes remain unresolved.'],
            'program':'\n'.join(program)+'\n','readbackProgram':'\n'.join(readback)+'\n'}


def _validate_contract(display):
    keys=('sourceHash','planHash','receiptHash','artifactHash','feature','target','tableName','columns','systemColumn')
    if _hash({key:display[key] for key in keys})!=display.get('contractHash'):
        raise ValueError('display contract hash is stale')


def _version(value):
    return isinstance(value,int) and not isinstance(value,bool) and value>=0


def _current_version(capture,document_id):
    if (not isinstance(capture,dict) or capture.get('request',{}).get('workbookId')!=document_id
            or capture.get('result',{}).get('workbookId')!=document_id
            or not _version(capture['result'].get('documentVersion'))):
        raise ValueError('controller current-version capture has a foreign target or invalid version')
    return capture['result']['documentVersion']


def _sdk_output(capture,display,program,version,expected_type,persisted):
    request=capture.get('request',{})
    raw_result=capture.get('result',{})
    result=raw_result.get('structuredContent',{})
    if raw_result.get('isError') is True:
        raise ValueError('raw SDK tool call failed')
    if not _version(request.get('documentVersion')) or request!={'workbookId':display['target']['documentId'],'documentVersion':version,'code':display[program]}:
        raise ValueError('raw SDK request context/code/target/version does not match approved display program')
    if result.get('type')!=expected_type or result.get('persisted') is not persisted:
        raise ValueError('raw SDK result has no successful expected execution/read-only status')
    if program=='readbackProgram' and (result.get('workbookId',display['target']['documentId'])!=display['target']['documentId']
            or result.get('documentVersion',version)!=version):
        raise ValueError('raw SDK observation result target/version is stale or foreign')
    try:
        output=result['output']
        if not isinstance(output,str): raise ValueError('SDK output must be raw text')
        # SDK print output can be JSON or a Python literal; never execute its text.
        try: output=json.loads(output)
        except json.JSONDecodeError: output=ast.literal_eval(output)
    except (KeyError,ValueError,SyntaxError,TypeError) as exc:
        raise ValueError('raw SDK output must contain one structured display observation') from exc
    expected_source='sdk-display-execution' if program=='program' else 'sdk-column-readback'
    if (not isinstance(output,dict) or output.get('source')!=expected_source
            or output.get('contractHash')!=display['contractHash']
            or output.get('elementId')!=display['target']['elementId']):
        raise ValueError('raw SDK output target/contract/source is stale or foreign')
    return output,result


def _capture_binding(display,execution,observation,current):
    _validate_contract(display)
    target=display['target']
    _,result=_sdk_output(execution,display,'program',target['version'],'success',True)
    version=result.get('documentVersion')
    if (result.get('workbookId')!=target['documentId'] or not _version(version)
            or version<=target['version']):
        raise ValueError('raw SDK execution result target/version is stale or foreign')
    output,_=_sdk_output(observation,display,'readbackProgram',version,'success-no-edits',False)
    if _current_version(current,target['documentId'])!=version:
        raise ValueError('current workbook version differs from SDK observation execution version')
    return {'contractHash':display['contractHash'],'sourceHash':display['sourceHash'],
            'planHash':display['planHash'],'documentId':target['documentId'],'documentVersion':version,
            'executionHash':_hash(execution),'observationHash':_hash(observation),'currentVersionHash':_hash(current)},output


def checkpoint_input_display(display,execution,observation,current,state=None):
    """Controller records exact submitted SDK calls/results, never a caller-invented readback ID.

    The returned authoritative ledger must be persisted before verification. This is a
    local controller evidence boundary, not authentication of arbitrary supplied JSON.
    """
    binding,_=_capture_binding(display,execution,observation,current)
    state=copy.deepcopy(state) if state is not None else {'schemaVersion':1,'attempts':[]}
    _validate_ledger(state)
    for attempt in state['attempts']:
        if (attempt['binding']['executionHash']==binding['executionHash']
                or attempt['binding']['observationHash']==binding['observationHash']):
            raise ValueError('raw SDK execution or observation already recorded; replay rejected')
        if (attempt['binding']['documentId']==binding['documentId']
                and attempt['binding']['documentVersion']>=binding['documentVersion']):
            raise ValueError('display checkpoint version is stale relative to durable ledger')
    for attempt in state['attempts']:
        if attempt['binding']['documentId']==binding['documentId'] and attempt['status']=='PENDING':
            attempt.update(status='SUPERSEDED',supersededBy=_hash(binding))
    state['attempts'].append({'binding':binding,'proofHash':_hash(binding),'execution':copy.deepcopy(execution),
        'observation':copy.deepcopy(observation),'currentVersion':copy.deepcopy(current),'status':'PENDING'})
    return state


def _validate_ledger(state):
    if not isinstance(state,dict) or state.get('schemaVersion')!=1 or not isinstance(state.get('attempts'),list):
        raise ValueError('durable controller display checkpoint ledger is required')
    proofs=set()
    for attempt in state['attempts']:
        binding=attempt.get('binding',{})
        if (attempt.get('status') not in {'PENDING','CONSUMED','SUPERSEDED'} or attempt.get('proofHash')!=_hash(binding)
                or attempt['proofHash'] in proofs
                or binding.get('executionHash')!=_hash(attempt.get('execution'))
                or binding.get('observationHash')!=_hash(attempt.get('observation'))
                or binding.get('currentVersionHash')!=_hash(attempt.get('currentVersion'))):
            raise ValueError('durable display capture/proof hash is stale or invalid')
        proofs.add(attempt['proofHash'])


def verify_input_display(display,state,current):
    """Consume a persisted pending checkpoint in-place, even when metadata fails.

    Caller must durably save this ledger before using the verdict. The public CLI
    locks the authoritative ledger and atomically persists consumption. Current is
    a fresh controller-recorded workbook GET/context capture; REST columns are not
    display evidence. SDK offers no CAS or signed provenance: controller capture and
    freshness are trusted, and ledger rollback/deletion must not be used to retry.
    """
    _validate_contract(display)
    _validate_ledger(state)
    pending=[a for a in state['attempts'] if a['binding']['contractHash']==display['contractHash'] and a['status']=='PENDING']
    if len(pending)!=1: raise ValueError('exactly one pending unconsumed display checkpoint is required')
    attempt=pending[0]
    binding,readback=_capture_binding(display,attempt['execution'],attempt['observation'],attempt['currentVersion'])
    if binding!=attempt['binding']: raise ValueError('display checkpoint binding hash is stale')
    if _current_version(current,binding['documentId'])!=binding['documentVersion']:
        raise ValueError('fresh current workbook version differs from checkpoint version')
    records=readback.get('columns')
    if not isinstance(records,list) or any(not isinstance(c,dict) for c in records):
        raise ValueError('SDK display columns are missing')
    observed={c.get('id'):c for c in records}
    if len(observed)!=len(records): raise ValueError('duplicate SDK display columns')
    mismatches=[]
    for c in display['columns']:
        actual=observed.get(c['id']) or {}
        for field in ('name','columnClass','type'):
            if actual.get(field)!=c[field]: mismatches.append({'columnId':c['id'],'field':field})
        if (actual.get('systemKind') is not None or actual.get('visible') is not c['visible']
                or actual.get('elementId')!=display['target']['elementId'] or actual.get('formula') is not None):
            mismatches.append({'columnId':c['id'],'field':'visibility/systemKind'})
        if c.get('allowedValues') is not None and actual.get('allowedValues')!=c['allowedValues']:
            mismatches.append({'columnId':c['id'],'field':'allowedValues'})
        intended=c['format']; fmt=actual.get('format')
        if fmt is not None and not isinstance(fmt,dict):
            raise ValueError('SDK format observation must be a structured record or null')
        if intended:
            matches=(fmt is None or fmt.get('kind')=='automatic') if intended=={'kind':'automatic'} else (
                isinstance(fmt,dict) and all(fmt.get(key)==value for key,value in intended.items()))
            if not matches: mismatches.append({'columnId':c['id'],'field':'format','expected':intended,'actual':fmt})
    expected={c['id'] for c in display['columns']}
    extra=[c for c in records if c.get('id') not in expected]
    system=display['systemColumn']
    fields=('name','columnClass','type','systemKind','formula','elementId')
    if (len(extra)!=1 or any(extra[0].get(field)!=system.get(field) for field in fields)
            or ('id' in system and extra[0].get('id')!=system['id'])):
        mismatches.append({'field':'complete approved system ID/alias identity readback required'})
    if any(c.get('visible') is not False for c in extra):
        mismatches.append({'field':'unexpected visible columns'})
    unresolved=[{'columnId':c['id'],'sourceFormat':c['sourceFormat'],'reason':c['unresolvedReason']}
                for c in display['columns'] if c['unresolvedReason']]
    result={'status':'FAILED' if mismatches else ('PARTIAL' if unresolved else 'PASSED'),
            'supportedDisplayStatus':'FAILED' if mismatches else 'PASSED',
            'contractHash':display['contractHash'],'proofHash':attempt['proofHash'],
            'documentId':binding['documentId'],'documentVersion':binding['documentVersion'],
            'mismatches':mismatches,'unresolved':unresolved,
            'verificationScope':'SDK column visibility/format metadata only; rendered output, data invariants and UI validation require separate acceptance'}
    attempt.update(status='CONSUMED',verification=result,currentVerificationCapture=copy.deepcopy(current),
                   currentVerificationHash=_hash(current))
    return result
