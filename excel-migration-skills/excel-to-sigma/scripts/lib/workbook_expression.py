"""Bind the shared Excel AST to one immutable source workbook and evaluate locally.

No cached formula values or external workbook files participate in calculation.
The source workbook is loaded with ``data_only=False`` and never saved.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
import hashlib
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.datetime import to_excel

try:
    from .excel_expression import (Binary, Boolean, BoundIndirect, BoundRange, CellReference,
        ColumnRange, ExcelCalcError, ExpressionError, Function, Group, Number, Percent, Range,
        Reference, StructuredReference, Text, Unary, UnsupportedExpression, evaluate,
        expand_range, parse_expression, _contains_whole_column)
except ImportError:
    from excel_expression import (Binary, Boolean, BoundIndirect, BoundRange, CellReference,
        ColumnRange, ExcelCalcError, ExpressionError, Function, Group, Number, Percent, Range,
        Reference, StructuredReference, Text, Unary, UnsupportedExpression, evaluate,
        expand_range, parse_expression, _contains_whole_column)


class UnresolvedBinding(UnsupportedExpression):
    pass


class FormulaCycle(UnresolvedBinding):
    def __init__(self, members):
        self.members = frozenset(members)
        super().__init__("formula cycle: " + ", ".join(sorted(f"{sheet}!{address}" for sheet, address in members)))


_MAX_BOUND_CELLS = 100_000


@dataclass(frozen=True)
class _Table:
    sheet: str
    name: str
    first_col: int
    last_col: int
    header_first: int
    header_last: int
    body_first: int
    body_last: int
    totals_first: int | None
    totals_last: int | None
    columns: tuple[str, ...]


def _address(address):
    return address.replace("$", "").upper()


def _identity(ref):
    return ref.sheet, ref.address


def _cell_id(ref):
    if ref.workbook:
        return f"[{ref.workbook}]{ref.sheet}!{ref.address}"
    return f"{ref.sheet}!{ref.address}"


def _checked_range(bound):
    if bound.empty:
        return bound
    left, top, _, _ = range_boundaries(bound.start.address)
    right, bottom, _, _ = range_boundaries(bound.end.address)
    if (abs(right - left) + 1) * (abs(bottom - top) + 1) > _MAX_BOUND_CELLS:
        raise UnresolvedBinding(f"bound range exceeds {_MAX_BOUND_CELLS} cells")
    return bound


class WorkbookContext:
    def __init__(self, source_path, workbook, source_hash, external_inputs=None, indirect_bindings=None):
        self.source_path = str(source_path)
        self.source_sha256 = source_hash
        self.workbook = workbook
        self.external_inputs = {(_workbook, _sheet, _address(address)): value
                                for (_workbook, _sheet, address), value in (external_inputs or {}).items()}
        self.indirect_bindings = indirect_bindings or {}
        self.formulas = {}
        self.literals = {}
        self.numeric_blanks = set()
        self.tables = {}
        self.names = {}
        self._bound = {}
        for worksheet in workbook.worksheets:
            for cell in worksheet._cells.values():
                key = worksheet.title, cell.coordinate
                if cell.data_type == "f":
                    self.formulas[key] = cell.value
                elif cell.data_type == "inlineStr" and cell.value is None:
                    self.literals[key] = ""
                elif cell.value is not None:
                    self.literals[key] = cell.value
                elif cell.data_type == "n":
                    self.numeric_blanks.add(key)
            for table in worksheet.tables.values():
                first_col, first_row, last_col, last_row = range_boundaries(table.ref)
                headers = table.headerRowCount if table.headerRowCount is not None else 1
                totals = table.totalsRowCount or 0
                column_names = tuple(column.name for column in table.tableColumns)
                if not column_names:
                    column_names = tuple(str(worksheet.cell(first_row, col).value or "")
                                         for col in range(first_col, last_col + 1))
                self.tables[table.name.casefold()] = _Table(
                    worksheet.title, table.name, first_col, last_col, first_row,
                    first_row + headers - 1, first_row + headers, last_row - totals,
                    last_row - totals + 1 if totals else None, last_row if totals else None,
                    column_names)
        for definition in workbook.defined_names.values():
            self.names[(None, definition.name.casefold())] = definition.attr_text
        for worksheet in workbook.worksheets:
            for definition in worksheet.defined_names.values():
                self.names[(worksheet.title, definition.name.casefold())] = definition.attr_text
        self.active_last_row = max((int(re.search(r"\d+$", address).group())
                                    for _, address in self.formulas.keys() | self.literals.keys()),
                                   default=1)

    @classmethod
    def from_path(cls, source_path, *, external_inputs=None, indirect_bindings=None):
        path = Path(source_path)
        source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        workbook = load_workbook(path, data_only=False, read_only=False,
                                 keep_vba=path.suffix.lower() == ".xlsm", keep_links=True)
        return cls(path, workbook, source_hash, external_inputs, indirect_bindings)

    def _name(self, name, sheet, address, stack):
        key = (sheet, name.casefold()) if (sheet, name.casefold()) in self.names else (None, name.casefold())
        if key not in self.names:
            raise UnresolvedBinding(f"unresolved defined name {name!r}")
        if key in stack:
            raise FormulaCycle({(sheet, address)})
        expression = self.names[key]
        if not isinstance(expression, str) or not expression:
            raise UnresolvedBinding(f"defined name {name!r} has no expression")
        return self._bind(parse_expression(expression if expression.startswith("=") else "=" + expression,
                                           sheet=sheet, workbook_mode=True),
                          sheet, address, stack | {key})

    def _table_range(self, table, specification, sheet, address):
        tokens = [token.strip() for token in re.findall(r"\[([^\[\]]+)\]", specification)]
        selector = "#This Row" if "@" in specification else None
        columns = []
        for token in tokens:
            if token.startswith("#"):
                selector = token
            elif token.startswith("@"):
                selector = "#This Row"
                if token[1:]: columns.append(token[1:])
            else:
                columns.append(token)
        if specification and not tokens:
            raise UnresolvedBinding(f"invalid structured reference {specification!r}")
        selector = (selector or "#Data").casefold()
        if selector in ("#data", "#all"):
            first_row = table.header_first if selector == "#all" else table.body_first
            last_row = table.totals_last or table.body_last if selector == "#all" else table.body_last
        elif selector in ("#this row", "#thisrow"):
            if sheet != table.sheet or not address:
                raise UnresolvedBinding(f"{table.name} row context is outside table")
            row = int(re.search(r"\d+$", address).group())
            if row < table.body_first or row > table.body_last:
                raise UnresolvedBinding(f"{table.name} row context is outside body")
            first_row = last_row = row
        elif selector == "#headers":
            first_row, last_row = table.header_first, table.header_last
        elif selector == "#totals":
            if table.totals_first is None:
                raise UnresolvedBinding(f"{table.name} has no totals row")
            first_row, last_row = table.totals_first, table.totals_last
        else:
            raise UnresolvedBinding(f"unsupported table selector {selector}")
        if len(columns) > 2:
            raise UnresolvedBinding("noncontiguous table column selection requires review")
        indexes = []
        for name in columns:
            matching = [index for index, header in enumerate(table.columns)
                        if header.casefold() == name.casefold()]
            if len(matching) != 1:
                raise UnresolvedBinding(f"unknown or ambiguous table column {table.name}[{name}]")
            indexes.append(matching[0])
        first_col = table.first_col + min(indexes) if indexes else table.first_col
        last_col = table.first_col + max(indexes) if indexes else table.last_col
        empty = first_row > last_row
        start = CellReference(table.sheet, f"{get_column_letter(first_col)}{first_row}")
        end = CellReference(table.sheet, f"{get_column_letter(last_col)}{first_row if empty else last_row}")
        return _checked_range(BoundRange(start, end, f"table:{table.name}:{selector}",
                                         f"{table.header_first}:{table.totals_last or table.body_last}",
                                         "empty" if empty else f"{first_row}:{last_row}", empty))

    def _bind_indirect(self, argument, sheet, address, stack):
        selector = self._bind(argument, sheet, address, stack)
        if isinstance(argument, Text):
            keys = [argument.value]
        elif (isinstance(argument, Binary) and argument.operator == "&"
              and isinstance(argument.left, Text)):
            prefix = argument.left.value.casefold()
            keys = sorted({name for scope, name in self.names if name.startswith(prefix)
                           and (scope is None or scope == sheet)})
        else:
            raise UnresolvedBinding("INDIRECT requires a finite reviewed name selector")
        if not keys:
            raise UnresolvedBinding("INDIRECT selector has no finite named-range alternatives")
        override = self.indirect_bindings.get((sheet, address))
        if override is not None:
            if not isinstance(override, dict) or not override:
                raise UnresolvedBinding("INDIRECT reviewed bindings must map selectors to source references")
            alternatives = []
            for name, target in override.items():
                if not isinstance(name, str) or not name:
                    raise UnresolvedBinding("INDIRECT selector name must be nonempty text")
                if isinstance(target, tuple) and len(target) == 2:
                    bound_target = self._bind(Reference(CellReference(target[0], _address(target[1]))),
                                              sheet, address, stack)
                elif isinstance(target, str):
                    bound_target = self._bind(parse_expression(
                        target if target.startswith("=") else "=" + target,
                        sheet=sheet, workbook_mode=True), sheet, address, stack)
                else:
                    raise UnresolvedBinding("INDIRECT target must be a source reference expression")
                if not isinstance(bound_target, (Reference, BoundRange)):
                    raise UnresolvedBinding("INDIRECT reviewed target must bind to source cells")
                alternatives.append((name, bound_target))
            return BoundIndirect(selector, tuple(alternatives))
        alternatives = tuple((name, self._name(name, sheet, address, stack)) for name in keys)
        if any(not isinstance(target, (Reference, BoundRange)) for _, target in alternatives):
            raise UnresolvedBinding("INDIRECT named target must bind to source cells")
        return BoundIndirect(selector, alternatives)

    def _bind(self, node, sheet, address, stack=frozenset()):
        if isinstance(node, (Number, Text, Boolean, BoundRange, BoundIndirect)):
            return node
        if isinstance(node, StructuredReference):
            if node.table is None:
                candidates = [table for table in self.tables.values() if table.sheet == sheet and address
                              and table.body_first <= int(re.search(r"\d+$", address).group()) <= table.body_last]
                if len(candidates) != 1:
                    raise UnresolvedBinding("unqualified row table reference is ambiguous")
                table = candidates[0]
            else:
                table = self.tables.get(node.table.casefold())
                if table is None: raise UnresolvedBinding(f"missing table {node.table!r}")
            return self._table_range(table, node.specification, sheet, address)
        if isinstance(node, ColumnRange):
            target_sheet = node.sheet or sheet
            if target_sheet not in self.workbook.sheetnames:
                raise UnresolvedBinding(f"missing sheet {target_sheet!r}")
            first_col, last_col = sorted((column_index_from_string(node.start_column),
                                          column_index_from_string(node.end_column)))
            top = CellReference(target_sheet, f"{get_column_letter(first_col)}1")
            bottom = CellReference(target_sheet, f"{get_column_letter(last_col)}{self.active_last_row}")
            return _checked_range(BoundRange(top, bottom, "whole_column", "1:1048576", f"1:{self.active_last_row}"))
        if isinstance(node, Reference):
            ref = node.value
            if ref.name:
                table = self.tables.get(ref.name.casefold())
                if table is not None: return self._table_range(table, "", sheet, address)
                return self._name(ref.name, sheet, address, stack)
            if ref.workbook:
                return node
            target = ref.sheet or sheet
            if target not in self.workbook.sheetnames:
                raise UnresolvedBinding(f"missing sheet {target!r}")
            return Reference(CellReference(target, ref.address, ref.column_absolute, ref.row_absolute))
        if isinstance(node, Range):
            start = self._bind(Reference(node.start), sheet, address, stack)
            end = self._bind(Reference(node.end), sheet, address, stack)
            if not isinstance(start, Reference) or not isinstance(end, Reference):
                raise UnresolvedBinding("range endpoints must be cells")
            if start.value.sheet != end.value.sheet or start.value.workbook or end.value.workbook:
                raise UnresolvedBinding("cross-sheet/external ranges require review")
            return _checked_range(BoundRange(start.value, end.value, "cell_range", f"{start.value.address}:{end.value.address}",
                                             f"{start.value.address}:{end.value.address}"))
        if isinstance(node, Group): return Group(self._bind(node.expression, sheet, address, stack))
        if isinstance(node, Unary): return Unary(node.operator, self._bind(node.operand, sheet, address, stack))
        if isinstance(node, Percent): return Percent(self._bind(node.operand, sheet, address, stack))
        if isinstance(node, Binary):
            return Binary(node.operator, self._bind(node.left, sheet, address, stack),
                          self._bind(node.right, sheet, address, stack))
        if isinstance(node, Function):
            if node.name == "INDIRECT":
                if len(node.arguments) != 1:
                    raise UnresolvedBinding("INDIRECT reference mode requires review")
                return self._bind_indirect(node.arguments[0], sheet, address, stack)
            arguments = tuple(self._bind(arg, sheet, address, stack) for arg in node.arguments)
            return Function(node.name, arguments)
        raise UnsupportedExpression(f"cannot bind {type(node).__name__}")

    def bind_formula(self, formula, *, sheet, address):
        parsed = parse_expression(formula, sheet=sheet, workbook_mode=True)
        return self._bind(parsed, sheet, _address(address))

    def _bound_formula(self, key):
        if key not in self._bound:
            formula = self.formulas[key]
            if not isinstance(formula, str):
                raise UnresolvedBinding("array/shared formula object needs scalar source expression")
            self._bound[key] = self.bind_formula(formula, sheet=key[0], address=key[1])
        return self._bound[key]

    def _normal_value(self, value):
        if isinstance(value, (datetime, date, time)):
            return to_excel(value, self.workbook.epoch)
        return value

    def evaluate_cell(self, sheet, address, *, overrides=None):
        key = sheet, _address(address)
        replacement = {(target_sheet, _address(target_address)): value
                       for (target_sheet, target_address), value in (overrides or {}).items()}
        if any(target in self.formulas for target in replacement):
            raise ValueError("overrides replace source literals, not formulas")
        memo, active = {}, []

        def calculate(target):
            if target in memo: return memo[target]
            if target in active:
                raise FormulaCycle(active[active.index(target):])
            if target in replacement:
                value = self._normal_value(replacement[target]); memo[target] = value; return value
            if target in self.literals:
                value = self.literals[target]
                if value is not None and isinstance(value, str) and value.startswith("#") and self.workbook[target[0]][target[1]].data_type == "e":
                    raise ExcelCalcError(value)
                value = self._normal_value(value)
                memo[target] = value
                return value
            if target not in self.formulas:
                return None
            active.append(target)
            try:
                expression = self._bound_formula(target)

                def resolve(ref):
                    if ref.workbook:
                        identity = ref.workbook, ref.sheet, ref.address
                        if identity not in self.external_inputs:
                            raise UnresolvedBinding(f"missing external binding [{ref.workbook}]{ref.sheet}!{ref.address}")
                        return self._normal_value(self.external_inputs[identity])
                    return calculate(_identity(ref))

                def arithmetic(ref, value):
                    if value is None and _identity(ref) in self.numeric_blanks:
                        return 0
                    return value

                value = evaluate(expression, resolve, arithmetic_reference_resolver=arithmetic,
                                 arithmetic_context=True, date_epoch=self.workbook.epoch,
                                 excel_errors=True)
                memo[target] = value
                return value
            finally:
                active.pop()

        return calculate(key)

    def _bound_evidence(self, expression):
        dependencies, domains = set(), []
        complete = True

        def visit(node):
            nonlocal complete
            if isinstance(node, Reference):
                dependencies.add(_cell_id(node.value))
            elif isinstance(node, BoundRange):
                domains.append({"kind": node.source_kind, "declaredDomain": node.declared_domain,
                                "activeDomain": node.active_domain})
                dependencies.update(_cell_id(ref) for ref in expand_range(node))
            elif isinstance(node, BoundIndirect):
                visit(node.selector)
                for _, alternative in node.alternatives: visit(alternative)
            elif isinstance(node, (Group, Unary, Percent)):
                visit(node.expression if isinstance(node, Group) else node.operand)
            elif isinstance(node, Binary):
                visit(node.left); visit(node.right)
            elif isinstance(node, Function):
                if node.name == "OFFSET":
                    complete = False
                if node.name == "COUNTIFS" and any(
                        _contains_whole_column(node.arguments[index])
                        for index in range(0, len(node.arguments), 2)):
                    complete = False
                for argument in node.arguments: visit(argument)
        visit(expression)
        return sorted(dependencies), domains, complete

    def formula_report(self, *, overrides=None):
        formulas = {}
        counts = {"verified": 0, "unresolved": 0, "error": 0, "cycle": 0}
        for key in sorted(self.formulas):
            identity = f"{key[0]}!{key[1]}"
            record = {"formula": str(self.formulas[key]), "status": "unresolved", "value": None,
                      "error": None, "dependencies": [], "dependencyStatus": "incomplete",
                      "bindingEvidence": []}
            try:
                bound = self._bound_formula(key)
                record["dependencies"], record["bindingEvidence"], complete = self._bound_evidence(bound)
                record["dependencyStatus"] = "complete" if complete else "incomplete"
                record["value"] = self.evaluate_cell(*key, overrides=overrides)
                record["status"] = "verified"
            except FormulaCycle as error:
                record["status"] = "cycle" if key in error.members else "unresolved"
                record["error"] = str(error)
            except ExcelCalcError as error:
                record["status"] = "error"
                record["error"] = error.code
            except (ExpressionError, ArithmeticError, TypeError, ValueError) as error:
                record["error"] = str(error)
            counts[record["status"]] += 1
            formulas[identity] = record
        return {"sourceSha256": self.source_sha256, "formulaCount": len(self.formulas),
                "counts": counts, "formulas": formulas}
