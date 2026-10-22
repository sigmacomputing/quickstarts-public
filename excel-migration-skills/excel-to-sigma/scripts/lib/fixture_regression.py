"""Independent explicit equations for the five designated synthetic fixtures.

No migration translator, emitted expressions or cached formula values are used.
Edits are source-cell-keyed literal replacements. These expected results need
independent review and actual Sigma captures; computing them is not a test pass.
"""
import datetime as dt
import hashlib
import math
from pathlib import Path
from decimal import Decimal,ROUND_HALF_UP
from openpyxl import load_workbook


class FixtureOracle:
    def __init__(self,path,edits=None):
        self.workbook=load_workbook(path,data_only=False);self.edits=edits or {}
        self.source_hash=hashlib.sha256(Path(path).read_bytes()).hexdigest()
    def literal(self,sheet,address):
        cell=self.workbook[sheet][address];identity=sheet+'!'+address
        if cell.data_type=='f':raise ValueError('independent oracle may only read source literals: '+identity)
        return self.edits.get(identity,cell.value)
    def primary(self,sheet,number):
        ws=self.workbook[sheet];headers=[c.value for c in ws[1]]
        return {name:self.literal(sheet,ws.cell(number,i+1).coordinate) for i,name in enumerate(headers) if ws.cell(number,i+1).data_type!='f'}
    def close(self):self.workbook.close()


def _day(value):
    if isinstance(value,dt.datetime):return value.date()
    if isinstance(value,dt.date):return value
    return dt.datetime.fromisoformat(value.replace('Z','+00:00')).date()


def _timestamp(value):
    if isinstance(value,dt.datetime):return value
    if isinstance(value,dt.date):return dt.datetime.combine(value,dt.time())
    return dt.datetime.fromisoformat(value.replace('Z','+00:00'))


def sales(source,edits=None):
    oracle=FixtureOracle(source,edits)
    try:
        asof=oracle.literal('Assumptions','F2');rows=[]
        for i in range(2,50):
            r=oracle.primary('Pipeline',i);r['Weighted ACV']=r['Probability']*r['ACV']
            r['Cycle Days']=(_timestamp(r['Expected Close'])-_timestamp(r['Created Date'])).total_seconds()/86400
            r['Timing Status']='Closed' if r['Stage'] in ('Closed Won','Closed Lost') else 'Slipped' if _day(r['Expected Close'])<_day(asof) else 'Open'
            rows.append(r)
        targets=[]
        for i in range(2,6):
            rep=oracle.literal('Rep Targets','A'+str(i));target=oracle.literal('Rep Targets','B'+str(i))
            weighted=sum(r['Weighted ACV'] for r in rows if r['Rep']==rep)
            targets.append({'Rep':rep,'Annual Target':target,'Weighted Pipeline':weighted,'Coverage':weighted/target})
        return {'rows':rows,'repTargets':targets,'asOfDate':asof}
    finally:oracle.close()


def inventory(source,edits=None):
    oracle=FixtureOracle(source,edits)
    try:
        rows=[];purchase=[]
        for i in range(2,38):
            r=oracle.primary('Inventory',i);r['Available']=r['On Hand']-r['Allocated']
            r['Reorder Point']=r['Avg Daily Use']*r['Lead Days']+r['Safety Stock']
            r['Days Cover']=r['Available']/r['Avg Daily Use']
            r['Reorder Qty']=math.ceil(max(0,r['Reorder Point']+r['Avg Daily Use']*30-r['Available']))
            r['Status']='REORDER' if r['Available']<=r['Reorder Point'] else 'OK';rows.append(r)
            purchase.append({'SKU':r['SKU'],'Product':r['Product'],'Supplier':r['Supplier'],'Status':r['Status'],
                'Recommended Qty':r['Reorder Qty'],'Unit Cost':r['Unit Cost'],'Order Cost':r['Reorder Qty']*r['Unit Cost'],
                'Buyer Review':oracle.literal('Purchase Plan','H'+str(i+3))})
        return {'rows':rows,'purchasePlan':purchase}
    finally:oracle.close()


def expense(source,edits=None):
    oracle=FixtureOracle(source,edits)
    def cents(value):return int((Decimal(str(value))*100).quantize(Decimal(1),rounding=ROUND_HALF_UP))
    try:
        limits={oracle.literal('Policy','A'+str(i)):oracle.literal('Policy','B'+str(i)) for i in range(2,8)};rows=[];reconciliation=[]
        for i in range(2,50):
            r=oracle.primary('Expense Register',i);r['Policy Limit']=limits[r['Category']]
            r['Recon Variance']=r['Card Amount']-r['Ledger Amount'];r['Policy Excess']=max(0,r['Amount']-r['Policy Limit'])
            r['Exception']='RECON' if abs(cents(r['Card Amount'])-cents(r['Ledger Amount']))>1 else 'OVER LIMIT' if r['Policy Excess']>0 else 'MISSING RECEIPT' if r['Receipt Status']=='Missing' else 'PENDING' if r['Approval Status']!='Approved' else 'CLEAR'
            rows.append(r);reconciliation.append({**{k:r[k] for k in ('Expense ID','Employee','Department','Category','Amount','Exception')},
                'Approval':r['Approval Status'],'Variance':r['Recon Variance'],'Required Action':{'CLEAR':'None','RECON':'Match card and ledger','OVER LIMIT':'Manager review','MISSING RECEIPT':'Attach receipt'}.get(r['Exception'],'Complete approval')})
        return {'rows':rows,'reconciliation':reconciliation,'dashboardStaticExceptionExpression':7+31+1+5}
    finally:oracle.close()


def forecast(source,edits=None):
    oracle=FixtureOracle(source,edits)
    try:
        months=[oracle.literal('Monthly Model',chr(67+i)+'5') for i in range(12)];actuals={}
        for i in range(2,oracle.workbook['Actuals Detail'].max_row+1):
            r=oracle.primary('Actuals Detail',i);key=(r['Line Item'],r['Month']);actuals[key]=actuals.get(key,0)+r['Amount']
        rows=[];line_rows=[]
        for i in range(8,20):
            line=oracle.literal('Monthly Model','B'+str(i));section=oracle.literal('Monthly Model','A'+str(i));input_row=i-2
            method=oracle.literal('Forecast Inputs','B'+str(input_row));rate=oracle.literal('Forecast Inputs','C'+str(input_row))
            r={'Source Row':i,'Section':section,'Line Item':line}
            for n,month in enumerate(months):
                if n<6:value=actuals[(line,month)]
                elif method=='growth':value=actuals[(line,months[5])]*(1+rate)**(n-5)
                elif method=='flat':value=actuals[(line,months[5])]
                elif method=='manual':value=oracle.literal('Forecast Inputs',chr(68+n-6)+str(input_row))
                else:raise ValueError('unknown source method')
                r[month]=value
            r['FY Total']=sum(r[m] for m in months);line_rows.append(r)
        rows.extend(line_rows)
        for i in (22,23,24,25,27,28):
            label=oracle.literal('Monthly Model','B'+str(i));r={'Source Row':i,'Section':'Derived','Line Item':label}
            for month in months:
                totals={s:sum(v[month] for v in line_rows if v['Section']==s) for s in ('Revenue','Cost of Revenue','Operating Expenses','Below the Line')}
                r[month]=totals[label] if i<26 else totals['Revenue']-totals['Cost of Revenue']-totals['Operating Expenses']-(totals['Below the Line'] if i==28 else 0)
            r['FY Total']=sum(r[m] for m in months);rows.append(r)
        return {'rows':rows,'months':months,'monthlyValues':{(r['Line Item'],m):r[m] for r in line_rows for m in months}}
    finally:oracle.close()


def research(source,edits=None):
    oracle=FixtureOracle(source,edits);out={};previous={}
    rows=[8,9,11,13,14,16,18,19,20,22,23,25,27,28,29,31,32,36,37,38,39,43,44,46,47,48,53,54,55,56,57,61,62,63,64]
    try:
        for i,col in enumerate('BCDEFGHIJKLM'):
            literal=lambda row:oracle.literal('FY results',col+str(row))
            r={};r[8]=literal(8) if i<6 else previous[8]*(1+oracle.literal('Model Inputs','C'+str(i-4)))
            r[9]=None if i==0 else r[8]/previous[8]-1;r[11]=literal(11) if i<6 else literal(13)-r[8]
            r[13]=r[8]+r[11] if i<6 else literal(13);r[14]=r[13]/r[8]
            r[16]=literal(16) if i<6 else previous[16]*1.045;r[18]=r[13]+r[16];r[19]=r[18]/r[8];r[20]=literal(20)
            r[22]=r[18]+r[20];r[23]=r[22]/r[8];r[25]=literal(25) if i==0 else previous[25]
            r[27]=r[22]+r[25];r[28]=literal(28);r[29]=-r[27]*r[28];r[31]=r[27]+r[29];r[32]=None if i==0 else r[31]/previous[31]-1
            r[36]=literal(36) if i<6 else previous[36]-1;r[37]=r[31]/r[36];r[38]=literal(38);r[39]=r[38]/r[37]
            r[43]=literal(43);r[44]=r[43]/r[37];r[46]=r[43]*r[36];r[53]=literal(53);r[54]=literal(54);r[55]=r[54]-r[53]
            r[56]=literal(56);r[47]=r[46]+r[55];r[48]=r[47]/r[18];r[57]=r[55]/r[18]
            r[61]=literal(61);r[62]=literal(62);r[63]=r[61]+r[62];r[64]=r[63]/r[31]
            for row in rows:out['FY results!'+col+str(row)]=r[row]
            previous=r
        return {'cells':out}
    finally:oracle.close()
