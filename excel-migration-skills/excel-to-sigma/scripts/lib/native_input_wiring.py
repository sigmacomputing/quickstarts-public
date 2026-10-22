"""Source-owned native calculation graphs for ordinary business inputs.

All returned code-rep fragments are immutable architecture material. Input IDs
are logical @input references until an approved assigned receipt resolves them.
"""
from __future__ import annotations

import calendar
import copy
import datetime as dt
import json
import math

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


def _column_id(name):
    import re
    return re.sub(r'[^a-z0-9]+', '_', name.lower()).strip('_')


def _table(workbook, required):
    matches = []
    for ws in workbook:
        for table in ws.tables.values():
            c1, r1, c2, r2 = range_boundaries(table.ref)
            headers = [ws.cell(r1, c).value for c in range(c1, c2 + 1)]
            if set(required) <= set(headers):
                rows = []
                for row in ws.iter_rows(min_row=r1 + 1, max_row=r2, min_col=c1, max_col=c2):
                    if any(cell.data_type == 'f' for cell in row):
                        raise ValueError('forecast source tables must contain authored literals')
                    rows.append(dict(zip(headers, [cell.value for cell in row])))
                matches.append((rows, f'{ws.title}!{table.ref}'))
    if len(matches) != 1:
        raise ValueError('forecast requires one unambiguous source table with: ' + ', '.join(required))
    return matches[0]


def extract_forecast_model(source_path, input_spec):
    """Read a table-based P&L family; never read a cached formula result.

    Periods are the wide input's month headers; base amounts are aggregated by
    live SQL from the last actual month, including when a manual row switches to
    growth or flat. Each source line has exactly one wide editable row.
    """
    workbook = load_workbook(source_path, data_only=False)
    try:
        actuals, actual_range = _table(workbook, ['Date', 'Month', 'Line Item', 'Section',
                                                 'Section Order', 'Line Order', 'Amount'])
        lines, line_range = _table(workbook, ['Line Item', 'Section', 'Section Order', 'Line Order',
                                             'Forecast Method'])
    finally:
        workbook.close()
    keys = [row['Line Item'] for row in lines]
    if len(set(keys)) != len(keys) or set(keys) != {row['Line Item'] for row in input_spec['rows']}:
        raise ValueError('forecast line definitions and wide input keys must agree uniquely')
    if any(row['Method'] not in {'growth', 'manual', 'flat'} for row in input_spec['rows']):
        raise ValueError('forecast contains an unknown authored method')
    actual_months = {}
    rows = []
    for row in actuals:
        date = row['Date']
        if not isinstance(date, (dt.date, dt.datetime)):
            raise ValueError('forecast actual dates must be typed source dates')
        if row['Line Item'] not in keys or not isinstance(row['Amount'], (int, float)) or not math.isfinite(row['Amount']):
            raise ValueError('forecast actual amounts/line keys are outside the supported source contract')
        period = date.year * 12 + date.month
        end = dt.date(date.year, date.month, calendar.monthrange(date.year, date.month)[1]).isoformat()
        if row['Month'] in actual_months and actual_months[row['Month']]['num'] != period:
            raise ValueError('forecast source month label is ambiguous')
        actual_months[row['Month']] = {'month': row['Month'], 'num': period, 'end': end}
        rows.append({'Section': row['Section'], 'Section Order': row['Section Order'],
                     'Line Item': row['Line Item'], 'Line Order': row['Line Order'],
                     'Month': row['Month'], 'Month Num': period, 'Month End': end,
                     'Amount': row['Amount'], 'Period': 'Actual'})
    if not actual_months:
        raise ValueError('forecast requires actuals to establish its base period')
    base = max(month['num'] for month in actual_months.values())
    months = []
    for column in input_spec['columns']:
        name = column['name']
        if name in {'Line Item', 'Method', 'Monthly Growth'}:
            continue
        try:
            parsed = dt.datetime.strptime(name, '%b-%y')
        except ValueError as error:
            raise ValueError('unsupported forecast month header: ' + name) from error
        num = parsed.year * 12 + parsed.month
        end = dt.date(parsed.year, parsed.month, calendar.monthrange(parsed.year, parsed.month)[1])
        months.append({'month': name, 'num': num, 'end': end.isoformat(), 'n': num - base})
    if not months or any(m['n'] != i + 1 for i, m in enumerate(months)):
        raise ValueError('forecast months must follow the final actual month in source order')
    return {'actualRows': rows, 'lines': lines, 'months': months, 'baseMonthNum': base,
            'sourceRanges': [actual_range, line_range, input_spec['sourceRange']]}


def _literal(value):
    if value is None:
        return 'NULL'
    if isinstance(value, str):
        return "'" + value.replace("'", "''") + "'"
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError('unsupported SQL source literal')
    return str(value)


def _select_rows(rows, names):
    if not rows:
        raise ValueError('native source rows must not be empty')
    return '\nUNION ALL\n'.join(
        'SELECT ' + ', '.join(_literal(row[name]) + (f' AS C{i}' if index == 0 else '')
                             for i, name in enumerate(names))
        for index, row in enumerate(rows))


def forecast_elements(model, input_spec, connection_id, input_id='@input:forecastBusinessInput'):
    """Native actuals → base/spine → wide-input join → forecast → P&L union."""
    fields = ['Section', 'Section Order', 'Line Item', 'Line Order', 'Month', 'Month Num',
              'Month End', 'Amount', 'Period']
    actual_sql = _select_rows(model['actualRows'], fields)
    def sql_columns(names):
        return [{'id': _column_id(name), 'name': name,
                 'formula': f'[Custom SQL/C{i}]' if name != 'Month End' else f'Date([Custom SQL/C{i}])'}
                for i, name in enumerate(names)]
    actuals = {'id': 'actuals', 'kind': 'table', 'name': 'Actuals',
               'source': {'kind': 'sql', 'connectionId': connection_id, 'statement': actual_sql},
               'columns': sql_columns(fields)}
    spine_rows = []
    for line in model['lines']:
        for month in model['months']:
            spine_rows.append({**{name: line[name] for name in fields[:4]},
                               'Month': month['month'], 'Month Num': month['num'],
                               'Month End': month['end'], 'N': month['n']})
    spine_fields = fields[:7] + ['N']
    # Base is a query-time aggregate of actual source literals, never a formula cache.
    spine_sql = (f'WITH actual_source AS ({actual_sql}),\n'
                 f'spine_source AS ({_select_rows(spine_rows, spine_fields)})\n'
                 'SELECT s.C0, s.C1, s.C2, s.C3, s.C4, s.C5, s.C6, s.C7, '
                 'b.BASE AS C8 FROM spine_source s LEFT JOIN '
                 f'(SELECT C2, SUM(C7) AS BASE FROM actual_source WHERE C5 = {model["baseMonthNum"]} '
                 'GROUP BY C2) b ON s.C2 = b.C2')
    spine = {'id': 'spineWrap', 'kind': 'table', 'name': 'Forecast Spine',
             'source': {'kind': 'sql', 'connectionId': connection_id, 'statement': spine_sql},
             'columns': sql_columns(spine_fields + ['June Base'])}
    # Explicit source period matching also works for another year/month span.
    manual = 'Switch([FJ/Month], ' + ', '.join(
        json.dumps(month['month']) + ', [Inputs/' + month['month'] + ']'
        for month in model['months']) + ', Null)'
    forecast = {'id': 'forecast', 'kind': 'table', 'name': 'Forecast',
                'source': {'kind': 'join', 'name': 'FJ',
                           'primarySource': {'kind': 'table', 'elementId': 'spineWrap'},
                           'joins': [{'name': 'Inputs', 'joinType': 'left-outer',
                                      'left': {'kind': 'table', 'elementId': 'spineWrap'},
                                      'right': {'kind': 'table', 'elementId': input_id},
                                      'columns': [{'left': '[Line Item]', 'right': '[Line Item]'}]}]},
                'columns': [{'id': _column_id(name), 'name': name, 'formula': f'[FJ/{name}]'}
                            for name in spine_fields + ['June Base']] + [
                    {'id': 'kind', 'name': 'Kind', 'formula': '[Inputs/Method]'},
                    {'id': 'eff_rate', 'name': 'Eff Rate', 'formula': '[Inputs/Monthly Growth]'},
                    {'id': 'eff_manual', 'name': 'Eff Manual', 'formula': manual},
                    {'id': 'amount', 'name': 'Amount', 'formula':
                     'If([Kind] = "manual", [Eff Manual], [Kind] = "flat", [June Base], '
                     '[Kind] = "growth", [June Base] * (1 + [Eff Rate]) ^ [N], Null)'},
                    {'id': 'period', 'name': 'Period', 'formula': '"Forecast"'}]}
    combined = {'id': 'pnlCombined', 'kind': 'table', 'name': 'PnL Combined',
                'source': {'kind': 'union', 'sources': [{'kind': 'table', 'elementId': 'actuals'},
                                                       {'kind': 'table', 'elementId': 'forecast'}],
                           'matches': [{'outputColumnName': name, 'sourceColumns': [f'[{name}]', f'[{name}]']}
                                       for name in fields]},
                'columns': [{'id': _column_id(name), 'name': name,
                             'formula': f'[Union of 2 Sources/{name}]'} for name in fields] + [
                    {'id': _column_id(name), 'name': name,
                     'formula': f'If([Union of 2 Sources/Section] = "{section}", [Union of 2 Sources/Amount], 0)'}
                    for name, section in [('Revenue Amt', 'Revenue'), ('COGS Amt', 'Cost of Revenue'),
                                           ('OpEx Amt', 'Operating Expenses'), ('Below Amt', 'Below the Line')]]}
    metadata_fields = ['Line Item', 'Section', 'Section Order', 'Line Order', 'Owner']
    metadata_rows = [{name: row.get(name) for name in metadata_fields} for row in model['lines']]
    definitions = {'id': 'forecastLineDefinitions', 'kind': 'table', 'name': 'Forecast Line Definitions',
                   'source': {'kind': 'sql', 'connectionId': connection_id,
                              'statement': _select_rows(metadata_rows, metadata_fields)},
                   'columns': sql_columns(metadata_fields)}
    view = {'id': 'forecastInputsView', 'kind': 'table', 'name': 'Forecast Inputs View',
            'source': {'kind': 'join', 'name': 'Wide',
                       'primarySource': {'kind': 'table', 'elementId': input_id},
                       'joins': [{'name': 'Definition', 'joinType': 'left-outer',
                                  'left': {'kind': 'table', 'elementId': input_id},
                                  'right': {'kind': 'table', 'elementId': 'forecastLineDefinitions'},
                                  'columns': [{'left': '[Line Item]', 'right': '[Line Item]'}]}]},
            'columns': [{'id': _column_id(c['name']), 'name': c['name'],
                         'formula': f'[Wide/{c["name"]}]'} for c in input_spec['columns']] + [
                             {'id': _column_id(name), 'name': name, 'formula': f'[Definition/{name}]'}
                             for name in metadata_fields[1:]]}
    return [actuals, spine, forecast, combined, definitions, view]


def business_native_wiring(architecture, workflow, extracted, legacy_fragments):
    """Own native targets in the approved plan, including chart/view aliases."""
    result = copy.deepcopy(legacy_fragments)
    result.setdefault('checks', [])
    # Companion inputs are constructed separately through the approved SDK path.
    result['elements'] = [e for e in result['elements'] if e['kind'] != 'input-table']
    for element in result['elements']:
        for join in (element.get('source') or {}).get('joins', []):
            feature = join['right'].get('elementId')
            if feature in architecture['inputs']:
                join['right']['elementId'] = '@input:' + feature
    feature = workflow + 'BusinessInput'
    spec = architecture['inputs'][feature]
    if workflow == 'sales':
        base = next(e for e in result['elements'] if e['id'] == 'salesInputBase')
        for column in base['columns']:
            if column['name'] == 'Weighted ACV':
                column['formula'] = '[Probability] * [ACV]'
            elif column['name'] == 'Cycle Days':
                # Excel =C-B retains time fractions; DateDiff returns integer units.
                column['formula'] = 'DateDiff("millisecond", [Created Date], [Close Date]) * 0.000000011574074074074074'
        result['elements'].append({
            'id': 'salesPipeline', 'name': 'Pipeline', 'kind': 'table',
            'source': {'kind': 'table', 'elementId': 'salesWithSettings'},
            'columns': [{'id': _column_id(c['name']),
                         'name': {'Close Date': 'Expected Close', 'TimingStatus': 'Timing Status'}.get(c['name'], c['name']),
                         'formula': f'[Sales With Settings/{c["name"]}]'}
                        for c in next(e for e in result['elements'] if e['id'] == 'salesWithSettings')['columns']]})
    elif workflow == 'inventory':
        raw = '[Reorder Point] + [Avg Daily Use] * 30 - [Available]'
        formulas = [('Available', '[On Hand] - [Allocated]'),
                    ('Reorder Point', '[Avg Daily Use] * [Lead Days] + [Safety Stock]'),
                    ('Days Cover', '[Available] / [Avg Daily Use]'),
                    ('Reorder Qty', f'If({raw} > 0, RoundUp({raw}, 0), 0)'),
                    ('Status', 'If([Available] <= [Reorder Point], "REORDER", "OK")')]
        result['elements'].append({
            'id': 'inventory', 'kind': 'table', 'name': 'Inventory',
            'source': {'kind': 'table', 'elementId': '@input:' + feature},
            'columns': [{'id': _column_id(c['name']), 'name': c['name'],
                         'formula': f'[{spec["name"]}/{c["name"]}]'} for c in spec['columns']] + [
                             {'id': _column_id(name), 'name': name, 'formula': formula}
                             for name, formula in formulas]})
    elif workflow == 'expense':
        calc = next(e for e in result['elements'] if e['id'] == 'expensesWithPolicy')
        result['elements'].extend([
            {'id': 'expenseRegister', 'kind': 'table', 'name': 'Expense Register',
             'source': {'kind': 'table', 'elementId': 'expensesWithPolicy'},
             'columns': [{'id': _column_id(c['name']),
                          'name': {'Per-item Limit': 'Policy Limit', 'PolicyExcess': 'Policy Excess'}.get(c['name'], c['name']),
                          'formula': f'[Expenses With Policy/{c["name"]}]'} for c in calc['columns']]},
            {'id': 'expenseNumericAmountCount', 'kind': 'kpi-chart', 'name': 'Numeric Amount Count',
             'source': {'kind': 'table', 'elementId': '@input:' + feature},
             'columns': [{'id': 'numeric_amount_count', 'name': 'Numeric Amount Count',
                          'formula': f'Count([{spec["name"]}/Amount])'}],
             'value': {'columnId': 'numeric_amount_count'}}])
        result['checks'].append({'logicalCheckId': 'expense-numeric-amount-count',
                                 'sourceCellId': 'Dashboard!B5',
                                 'elementId': 'expenseNumericAmountCount',
                                 'columnId': 'numeric_amount_count',
                                 'queryColumns': ['numeric_amount_count'],
                                 'status': 'pending-runtime-query',
                                 'domain': 'Amount is numeric; Count excludes null amounts'})
    elif workflow == 'forecast':
        result['elements'] = forecast_elements(architecture['forecastModel'], spec,
                                               architecture['targetScope']['connectionId'])
        result['preJoinValidation'][feature] = {'uniqueKeys': ['Line Item'],
                                               'expectedKeys': [r['Line Item'] for r in spec['rows']]}
    result['consumerSources'][workflow] = {'forecast': 'pnlCombined', 'sales': 'salesPipeline',
                                           'inventory': 'inventory', 'expense': 'expenseRegister'}[workflow]
    return result


def integrate_native_wiring(current, manifest, receipt, emitted):
    """Prepare a guarded full-contents PUT while preserving every foreign node."""
    try:
        from . import code_rep
    except ImportError:
        import code_rep
    scope = manifest.get('target') or {}
    if (current.get('workbookId') != scope.get('documentId')
            or current.get('documentVersion') != scope.get('version')):
        raise ValueError('native wiring current target/version is stale')
    result = copy.deepcopy(code_rep.document(current))
    if not isinstance(result.get('elements'), list) or not isinstance(result.get('pages'), list):
        raise ValueError('native wiring requires complete flat GET-back contents')
    existing = {e['id']: e for e in result['elements']}
    for feature, binding in receipt['targetBindings'].items():
        element = existing.get(binding.get('elementId')) or {}
        spec = manifest['planProjection']['sharedInputArchitecture']['inputs'][feature]
        if element.get('kind') != 'input-table' or element.get('name') != spec['name']:
            raise ValueError('native wiring assigned input name/type is stale')
        expected = {c['name'] for c in spec['columns']}
        if set(binding.get('columnIds') or {}) != expected:
            raise ValueError('native wiring requires a complete assigned input binding')
        observed_columns = {c.get('id'): c for c in element.get('columns') or []}
        for column in spec['columns']:
            observed = observed_columns.get(binding['columnIds'][column['name']]) or {}
            if observed.get('name') != column['name'] or observed.get('type') != column['type']:
                raise ValueError('native wiring assigned input column ID/name/type is stale')
    workflow = manifest['planProjection']['sharedInputArchitecture']['workflows']
    if len(workflow) != 1:
        raise ValueError('native wiring integration requires a single business workflow')
    marker = 'excel-to-sigma:native:' + workflow[0] + ':' + manifest['sourceHash']
    nodes = copy.deepcopy(emitted['wiring']['elements'])
    ids = {e['id'] for e in nodes}
    if len(ids) != len(nodes):
        raise ValueError('native wiring has duplicate logical targets')
    for node in nodes:
        observed = existing.get(node['id'])
        if observed and (observed.get('description') or {}).get('text') != marker:
            raise ValueError('native wiring target collides with foreign contents: ' + node['id'])
        node['description'] = {'text': marker, 'visibility': 'hidden'}
        def references(value):
            if isinstance(value, dict):
                if value.get('kind') == 'table' and value.get('elementId'):
                    yield value['elementId']
                for item in value.values():
                    yield from references(item)
            elif isinstance(value, list):
                for item in value:
                    yield from references(item)
        if set(references(node.get('source'))) - ids - set(existing):
            raise ValueError('native wiring references missing runtime elements')
    import re
    import xml.etree.ElementTree as ET
    page_id = workflow[0] + '-native-model'
    expected_page = {'id': page_id, 'name': workflow[0].title() + ' Model', 'visibility': 'hidden'}
    pages = [p for p in result['pages'] if p['id'] == page_id]
    layout = result.get('layout') or ''
    try:
        layout_root = ET.fromstring('<Layout>' + re.sub(r'<\?xml[^?]*\?>', '', layout) + '</Layout>')
    except ET.ParseError as error:
        raise ValueError('native wiring cannot establish layout ownership') from error
    layout_pages = [p for p in layout_root.iter('Page') if p.get('id') == page_id]
    page_xml = ET.Element('Page', {'id': page_id, 'type': 'grid',
                                 'gridTemplateColumns': 'repeat(24, 1fr)', 'gridTemplateRows': 'auto'})
    if pages:
        # Page records have no supported description marker. Require the exact
        # generated metadata AND its complete owned placement contract instead.
        page = pages[0]
        if (len(pages) != 1 or len(layout_pages) != 1
                or set(page) - {'id', 'name', 'visibility', 'type'}
                or any(page.get(key) != value for key, value in expected_page.items())
                or page.get('type', 'page') != 'page'
                or not ids.issubset(existing)
                or layout_pages[0].attrib != page_xml.attrib):
            raise ValueError('native wiring page collides with foreign metadata or ownership')
        placed = list(layout_pages[0])
        expected_placements = [
            {'elementId': node['id'], 'gridColumn': '1 / 25',
             'gridRow': f'{index*12+1} / {index*12+13}'} for index, node in enumerate(nodes)]
        if (len(placed) != len(expected_placements)
                or (layout_pages[0].text or '').strip()
                or any(child.tag != 'Element' or child.attrib != expected
                       or list(child) or (child.text or '').strip() or (child.tail or '').strip()
                       for child, expected in zip(placed, expected_placements))):
            raise ValueError('native wiring page contains foreign or unowned placements')
        expected_page = copy.deepcopy(page)
    elif layout_pages:
        raise ValueError('native wiring page collides with foreign layout ownership')
    owned_layout_page = layout_pages[0] if pages else None
    owned_placements = set(owned_layout_page.iter()) if owned_layout_page is not None else set()
    for placement in layout_root.iter():
        if placement.get('elementId') in ids and placement not in owned_placements:
            raise ValueError('native wiring cannot move owned elements from foreign page placements')
    result['elements'] = [e for e in result['elements'] if e['id'] not in ids] + nodes
    result['pages'] = [p for p in result['pages'] if p['id'] != page_id] + [expected_page]
    # Remove the one separator inserted with this owned page, preserving all
    # surrounding foreign-layout whitespace byte for byte.
    layout = re.sub(r"\n?<Page\b[^>]*\bid\s*=\s*[\"']" + re.escape(page_id)
                    + r"[\"'][^>]*>.*?</Page>", '', layout, flags=re.S)
    for index, node in enumerate(nodes):
        ET.SubElement(page_xml, 'Element', {'elementId': node['id'], 'gridColumn': '1 / 25',
                                          'gridRow': f'{index*12+1} / {index*12+13}'})
    result['layout'] = layout + '\n' + ET.tostring(page_xml, encoding='unicode')
    return {'documentVersion': scope['version'], 'contents': result,
            'sourceHash': manifest['sourceHash'], 'planHash': manifest['planHash'],
            'wiringArtifactHash': next(iter(receipt['artifacts'].values()))['workbookWiringHash']}
