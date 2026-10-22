"""Generic, source-bound presentation plans and guarded chart construction.

The inventory is source evidence. A chart element is constructed only from an
explicit point-by-point live binding; cached Excel values are pending parity
expectations, never chart data. Actual emission reuses the shared staging gates.
"""
from __future__ import annotations

import copy
import hashlib
import json
import math
import re
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile

from openpyxl import load_workbook

try:
    from . import presentation
    from .input_orchestration import gate_input_staging
    from .migration_gate import gate_plan_staging_source
    from .migration_manifest import validate_manifest
except ImportError:
    import presentation
    from input_orchestration import gate_input_staging
    from migration_gate import gate_plan_staging_source
    from migration_manifest import validate_manifest


PLAN_VERSION = 1
SUPPORTED_CHART_TYPES = {"bar", "line", "doughnut"}


def _table_column_width(column, role, workbook):
    """Return a readable pixel width, retaining an explicit native width."""
    style = column["style"] if "style" in column else {}
    if not isinstance(style, dict):
        raise ValueError("table column style must be an object")
    if "width" in style:
        width = style["width"]
        if (isinstance(width, bool) or not isinstance(width, (int, float))
                or not math.isfinite(width) or width <= 0):
            raise ValueError("table column style.width must be a positive finite pixel number")
        return width

    name = str(column.get("name") or "")
    lower = name.lower()
    fmt = column.get("format") or {}
    kind = fmt.get("kind") if isinstance(fmt, dict) else None
    if kind == "datetime" or re.search(r"\bdate\b|\btime\b", lower):
        minimum = 135
    elif (kind == "number" and ("$" in str(fmt.get("formatString") or "")
                                or fmt.get("currencySymbol"))
          or re.search(r"\bamount\b|\bcost\b|\bprice\b|\brevenue\b|\bexpense\b", lower)):
        minimum = 145
    elif re.search(r"\bstatus\b|\bstate\b|\bapproval\b", lower):
        minimum = 150
    elif kind == "number" and "%" in str(fmt.get("formatString") or "") or re.search(r"\bpercent\b|\brate\b|\bratio\b", lower):
        minimum = 120
    elif re.search(r"\bid\b|\bcode\b|\bsku\b|\bkey\b", lower):
        minimum = 120
    elif re.search(r"\bcount\b|\bqty\b|\bquantity\b|\bdays\b", lower):
        minimum = 110
    elif re.search(r"\bnotes?\b|\bdescription\b|\bcomment\b", lower):
        minimum = 170
    else:
        minimum = 135 if role == "main" else 150
    header_width = math.ceil(len(name) * 7.2 + 28)
    source_width = 0
    reference = column.get("sourceCell")
    if workbook is not None and isinstance(reference, str) and "!" in reference:
        sheet, address = reference.rsplit("!", 1)
        sheet = sheet.strip("'").replace("''", "'")
        match = re.fullmatch(r"\$?([A-Z]+)\$?\d+", address)
        if match and sheet in workbook:
            dimension = workbook[sheet].column_dimensions.get(match.group(1))
            if dimension is not None and dimension.customWidth:
                chars = dimension.width
                if isinstance(chars, (int, float)) and math.isfinite(chars) and chars > 0:
                    source_width = round(chars * 7 + 5)
    ceiling = 280 if role == "main" else 360
    return min(ceiling, max(minimum, header_width, source_width))


def _sized_table_columns(columns, role, workbook):
    result = []
    seen = set()
    for column in columns:
        if not isinstance(column, dict) or not column.get("id") or not column.get("name"):
            raise ValueError(f"{role} table columns require unique ID and name")
        if column["id"] in seen:
            raise ValueError(f"{role} table columns require unique ID and name")
        seen.add(column["id"])
        emitted = copy.deepcopy(column)
        emitted.pop("sourceCell", None)
        width = _table_column_width(column, role, workbook)
        emitted["style"] = {**(emitted["style"] if "style" in emitted else {}),
                            "width": width}
        result.append(emitted)
    return result


def _digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":"),
                                     ensure_ascii=False, default=str).encode()).hexdigest()


def _source_label(path):
    return re.sub(r"^\d+\s+", "", Path(path).stem).strip()


def _chart_xml(path, part):
    with ZipFile(path) as archive:
        return ET.fromstring(archive.read(part))


def _single_cell_value(workbook, reference):
    if not isinstance(reference, str) or "!" not in reference:
        return None
    sheet, address = reference.rsplit("!", 1)
    sheet = sheet.strip("'").replace("''", "'")
    address = address.replace("$", "")
    if not re.fullmatch(r"[A-Z]+\d+", address) or sheet not in workbook:
        return None
    return workbook[sheet][address].value


def _series_name(workbook, node, fallback):
    title_ref = node.findtext("c:tx/c:strRef/c:f", namespaces=presentation.NS)
    value = _single_cell_value(workbook, title_ref)
    if value not in (None, "") and not str(value).startswith("="):
        return str(value)
    value = node.findtext("c:tx/c:v", namespaces=presentation.NS)
    return str(value) if value not in (None, "") else fallback


def _uniform_source_format(workbook, cells):
    formats = {workbook[sheet][address].number_format
               for identity in cells for sheet, address in [identity.rsplit("!", 1)]}
    return formats.pop() if len(formats) == 1 else None


def _complete_chart_inventory(source, chart, formulas, cached):
    root = _chart_xml(source, chart["sourceLocation"]["chartPart"])
    plot_area = root.find(".//c:plotArea", presentation.NS)
    plot_nodes = [] if plot_area is None else [node for node in plot_area
                                               if node.tag.rsplit("}", 1)[-1].endswith("Chart")]
    plot_kinds = [node.tag.rsplit("}", 1)[-1][:-5] for node in plot_nodes]
    kind = plot_kinds[0] if len(plot_nodes) == 1 else "combo" if plot_nodes else "unsupported"
    geometry = None
    if kind == "doughnut":
        node = plot_nodes[0]
        hole = node.find("c:holeSize", presentation.NS)
        angle = node.find("c:firstSliceAng", presentation.NS)
        geometry = {"holePercent": int(hole.attrib["val"]) if hole is not None else None,
                    "firstSliceAngle": int(angle.attrib["val"]) if angle is not None else None}
    plots, series, category_refs = [], [], []
    for plot_index, (plot_node, plot_kind) in enumerate(zip(plot_nodes, plot_kinds), 1):
        plot_series = []
        for node in plot_node.findall("c:ser", presentation.NS):
            references = {}
            for child in node:
                formula = child.findtext(".//c:f", namespaces=presentation.NS)
                if formula is not None:
                    references[child.tag.rsplit("}", 1)[-1]] = formula
            category_ref = references.get("cat") or references.get("xVal")
            value_ref = references.get("val") or references.get("yVal")
            category_refs.append(category_ref)
            cells = presentation._range_cells(value_ref)
            source_formulas = presentation._range_values(formulas, value_ref)
            source_format = presentation._direct_reference_format(formulas, source_formulas, cells)
            if source_format is None:
                source_format = _uniform_source_format(formulas, cells)
            if source_format is None:
                value_tag = "yVal" if plot_kind == "scatter" else "val"
                source_format = node.findtext(
                    f"c:{value_tag}/c:numRef/c:numCache/c:formatCode", namespaces=presentation.NS)
            _intent, fidelity, format_reason = presentation.source_format_intent(
                source_format or "General", "number")
            item = {
                "name": _series_name(formulas, node, f"Series {len(series) + 1}"),
                "order": len(series), "plotIndex": plot_index,
                "plotType": plot_kind, "categoryRef": category_ref,
                "categoryCells": presentation._range_cells(category_ref),
                "valueRef": value_ref, "sourceCells": cells,
                "sourceReferences": references,
                "sourceReferenceCells": {role: presentation._range_cells(ref)
                                         for role, ref in references.items()},
                "sourceFormulas": source_formulas,
                "cachedExpectations": presentation._range_values(cached, value_ref),
                "aggregation": presentation._aggregation(source_formulas),
                "numberFormat": source_format,
                "unit": presentation._unit(source_format, chart["title"]),
                "formatFidelity": fidelity, "formatReason": format_reason,
            }
            series.append(item)
            plot_series.append(item["order"])
        plots.append({"plotIndex": plot_index, "chartType": plot_kind,
                      "seriesIndexes": plot_series,
                      "seriesCount": len(plot_series),
                      "pointCounts": [len(series[index]["sourceCells"])
                                      for index in plot_series]})
    category_ref = category_refs[0] if category_refs else chart.get("categoryRef")
    category_cells = presentation._range_cells(category_ref)
    point_count = max([len(category_cells)] + [len(item["sourceCells"]) for item in series])
    consistent = (len(plots) == 1 and bool(series) and bool(category_cells)
                  and all(ref == category_ref for ref in category_refs)
                  and all(len(item["sourceCells"]) == point_count for item in series))
    return {
        **chart, "chartType": kind, "plots": plots, "categoryRef": category_ref,
        "categoryCells": category_cells,
        "categorySourceValues": presentation._range_values(formulas, category_ref),
        "cachedCategoryExpectations": presentation._range_values(cached, category_ref),
        "categoryOrder": list(range(1, point_count + 1)),
        "pointCount": point_count, "series": series,
        "sourceGeometry": geometry,
        "status": "supported" if consistent and kind in SUPPORTED_CHART_TYPES else "unresolved",
        "unresolvedReason": (None if consistent and kind in SUPPORTED_CHART_TYPES
                             else "combination chart requires complete native plot mapping" if len(plots) > 1
                             else "unsupported source chart type" if kind not in SUPPORTED_CHART_TYPES
                             else "source chart ranges are missing or inconsistent"),
    }


def build_source_plan(source_path):
    """Inventory charts and presentation rules for any supported OOXML workbook."""
    source = Path(source_path).resolve()
    extracted = presentation.extract_presentation(source)
    formulas = load_workbook(source, data_only=False, read_only=False)
    cached = load_workbook(source, data_only=True, read_only=False)
    try:
        charts = [_complete_chart_inventory(source, chart, formulas, cached)
                  for chart in extracted["charts"]]
    finally:
        formulas.close()
        cached.close()
    result = {
        "version": PLAN_VERSION,
        "source": copy.deepcopy(extracted["source"]),
        "sourceHash": extracted["source"]["sha256"],
        "sourceName": _source_label(source),
        "charts": charts,
        "validations": copy.deepcopy(extracted["validations"]),
        "conditionalFormats": copy.deepcopy(extracted["conditionalFormats"]),
        "semanticFormats": copy.deepcopy(extracted["semanticFormats"]),
        "formulaProvenance": copy.deepcopy(extracted["formulaProvenance"]),
        "extractionHash": extracted["extractionHash"],
    }
    result["planHash"] = _digest(result)
    return result


def _live_expression(value):
    return (isinstance(value, str) and bool(value.strip())
            and "[" in value and "]" in value
            and not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", value.strip()))


def _verify_chart_binding(plan, chart, binding, reviewed, context):
    if binding.get("categorySourceCells") != chart["categoryCells"]:
        raise ValueError(f"chart {chart['id']} category point coverage is incomplete")
    if not (_live_expression(binding.get("categoryFormula"))
            and _live_expression(binding.get("orderFormula"))):
        raise ValueError(f"chart {chart['id']} requires live expression for category and order")
    supplied_series = binding.get("series")
    if not isinstance(supplied_series, list) or len(supplied_series) != len(chart["series"]):
        raise ValueError(f"chart {chart['id']} series point coverage is incomplete")
    for source, target in zip(chart["series"], supplied_series):
        if target.get("sourceCells") != source["sourceCells"]:
            raise ValueError(f"chart {chart['id']} series point coverage is incomplete")
        if not _live_expression(target.get("formula")):
            raise ValueError(f"chart {chart['id']} requires live expression for every series")
    points = binding.get("points")
    if not isinstance(points, list) or len(points) != chart["pointCount"]:
        raise ValueError(f"chart {chart['id']} point coverage is incomplete")
    keys = []
    for index, point in enumerate(points):
        source_values = [series["sourceCells"][index] for series in chart["series"]]
        if (point.get("sourceCategoryCell") != chart["categoryCells"][index]
                or point.get("sourceValueCells") != source_values
                or point.get("targetOrder") != index + 1):
            raise ValueError(f"chart {chart['id']} point coverage or source ordering is incomplete")
        key = point.get("targetCategoryKey")
        if not isinstance(key, str) or not key.strip() or key in keys:
            raise ValueError(f"chart {chart['id']} point coverage requires unique target category keys")
        keys.append(key)
        if (point.get("targetCategoryExpression") != binding["categoryFormula"]
                or point.get("targetValueExpressions") != [item["formula"] for item in supplied_series]):
            raise ValueError(f"chart {chart['id']} point coverage does not match live expressions")
    if not isinstance(binding.get("sourceElementId"), str) or not binding["sourceElementId"].strip():
        raise ValueError(f"chart {chart['id']} requires live sourceElementId")
    if reviewed:
        review = binding.get("review") or {}
        facts = {key: value for key, value in binding.items() if key != "review"}
        if (review.get("status") != "approved" or not str(review.get("reviewer") or "").strip()
                or not str(review.get("evidenceId") or "").strip()
                or review.get("sourceHash") != plan["sourceHash"]
                or review.get("presentationPlanHash") != plan["planHash"]
                or review.get("manifestPlanHash") != context.get("manifestPlanHash")
                or review.get("stagingReceiptHash") != context.get("stagingReceiptHash")
                or review.get("documentId") != context.get("documentId")
                or review.get("documentVersion") != context.get("documentVersion")
                or review.get("bindingHash") != _digest(facts)):
            raise ValueError(f"chart {chart['id']} requires unchanged reviewed binding facts")


def _target_list(plan, bindings, key, rule, *, reviewed):
    if key not in bindings:
        return []
    mapping = bindings[key]
    if not isinstance(mapping, dict):
        raise ValueError(f"{key} must be a rule ID to target list mapping")
    if rule["id"] not in mapping:
        return []
    targets = mapping[rule["id"]]
    if isinstance(targets, dict):
        targets = [targets]
    if not isinstance(targets, list):
        raise ValueError(f"{key} {rule['id']} has malformed target list")
    for target in targets:
        if (not isinstance(target, dict) or not target.get("elementId")
                or not target.get("columnId") or target.get("sourceRange") != rule["range"]
                or target.get("sourceSheet") != rule["sheet"]):
            raise ValueError(f"{key} {rule['id']} has incomplete source-to-target facts")
        if reviewed:
            review = target.get("review") or {}
            facts = {name: value for name, value in target.items() if name != "review"}
            if (review.get("status") != "approved" or not str(review.get("reviewer") or "").strip()
                    or not str(review.get("evidenceId") or "").strip()
                    or review.get("sourceHash") != plan["sourceHash"]
                    or review.get("presentationPlanHash") != plan["planHash"]
                    or review.get("manifestPlanHash") != bindings.get("manifestPlanHash")
                    or review.get("stagingReceiptHash") != bindings.get("stagingReceiptHash")
                    or review.get("documentId") != bindings.get("documentId")
                    or review.get("documentVersion") != bindings.get("documentVersion")
                    or review.get("targetBindingHash") != _digest(facts)):
                raise ValueError(f"{key} {rule['id']} requires unchanged reviewed target facts")
    return targets


def _category_format(plan, chart):
    semantics = [plan["semanticFormats"].get(cell) for cell in chart["categoryCells"]]
    if not semantics or any(item is None for item in semantics):
        return None, None
    source_formats = {item["sourceFormat"] for item in semantics}
    kinds = {item["semantic"] for item in semantics}
    if len(source_formats) != 1 or len(kinds) != 1:
        return None, "mixed source category formats require render review"
    raw = source_formats.pop()
    kind = kinds.pop()
    if kind == "date" and raw == "mmm-yy":
        return ({"kind": "datetime", "formatString": "%b-%y"},
                "monthly date label requires native render readback")
    if kind in {"date", "number", "currency", "percentage"}:
        return copy.deepcopy(semantics[0]["emission"]), None
    return None, "source category style has no verified native mapping"


def _layout(page_ids, title_ids, chart_ids, main_id, edit_id, detail_id):
    overview_items = [(title_ids[0], 1, 25, 1, 4)]
    for index, chart_id in enumerate(chart_ids):
        row = 4 + (index // 2) * 15
        col = 1 if index % 2 == 0 else 13
        overview_items.append((chart_id, col, col + 12, row, row + 15))
    main_row = 4 + ((len(chart_ids) + 1) // 2) * 15
    overview_items.append((main_id, 1, 25, main_row, main_row + 18))
    page_items = [overview_items, [(title_ids[1], 1, 25, 1, 4), (edit_id, 1, 25, 4, 28)],
                  [(title_ids[2], 1, 25, 1, 4), (detail_id, 1, 25, 4, 42)]]
    pages = []
    for page_id, items in zip(page_ids, page_items):
        lines = [f'<Page type="grid" gridTemplateColumns="repeat(24, 1fr)" gridTemplateRows="auto" id="{page_id}">']
        lines += [f'  <Element elementId="{eid}" gridColumn="{c1} / {c2}" gridRow="{r1} / {r2}"/>'
                  for eid, c1, c2, r1, r2 in items]
        pages.append("\n".join(lines + ["</Page>"]))
    return '<?xml version="1.0" encoding="utf-8"?>\n' + "\n".join(pages)


def _render(plan, bindings, *, staging_gate, manifest=None):
    if (bindings.get("sourceHash") != plan["sourceHash"]
            or bindings.get("presentationPlanHash") != plan["planHash"]):
        raise ValueError("presentation bindings are stale for source or plan")
    for name in ("documentId", "documentVersion", "viewSourceElementId", "mainElementId",
                 "detailElementId", "editElementId"):
        if not bindings.get(name):
            raise ValueError(f"presentation binding requires {name}")
    for name in ("mainViewColumns", "detailViewColumns"):
        if not isinstance(bindings.get(name), list) or not bindings[name]:
            raise ValueError(f"presentation binding requires {name}")
    prefix = f"source-{plan['sourceHash'][:10]}"
    marker = f"excel-to-sigma:source-presentation:{plan['sourceHash']}"
    page_ids = [f"{prefix}-{role}" for role in ("overview", "edit", "detail")]
    title_ids = [f"{page_id}-title" for page_id in page_ids]
    pages = [{"id": page_id, "name": f"{plan['sourceName']} {role.title()}"}
             for page_id, role in zip(page_ids, ("overview", "edit", "detail"))]
    elements = [{"id": title_id, "kind": "text", "body": f"# {escape(page['name'])}",
                 "verticalAlign": "center", "description": {"text": marker, "visibility": "hidden"}}
                for title_id, page in zip(title_ids, pages)]
    emitted_ids, emitted_sources, checks, unresolved = [], [], [], []
    checks_by_cell = defaultdict(list)
    for chart_index, chart in enumerate(plan["charts"], 1):
        binding = (bindings.get("charts") or {}).get(chart["id"])
        if chart["status"] != "supported":
            unresolved.append({"kind": "chart", "sourceId": chart["id"],
                               "title": chart["title"], "reason": chart["unresolvedReason"],
                               "sourcePointCount": chart["pointCount"]})
            continue
        if chart["chartType"] == "doughnut" and len(chart["series"]) != 1:
            unresolved.append({"kind": "chart", "sourceId": chart["id"],
                               "title": chart["title"],
                               "reason": "native donut mapping requires exactly one source series",
                               "sourcePointCount": chart["pointCount"]})
            continue
        if binding is None:
            unresolved.append({"kind": "chart", "sourceId": chart["id"],
                               "title": chart["title"], "reason": "missing reviewed live binding",
                               "sourcePointCount": chart["pointCount"]})
            continue
        _verify_chart_binding(plan, chart, binding,
                              reviewed=staging_gate["status"] != "LOGICAL_FIXTURE_ONLY",
                              context=bindings)
        element_id = f"{prefix}-chart-{chart_index}"
        category_id, order_id = f"chart_{chart_index}_category", f"chart_{chart_index}_order"
        columns = [{"id": category_id, "name": "Category", "formula": binding["categoryFormula"]}]
        category_format, category_reason = _category_format(plan, chart)
        if category_format is not None:
            columns[0]["format"] = category_format
        if category_reason:
            unresolved.append({"kind": "chart-category-format", "sourceId": chart["id"],
                               "reason": category_reason})
        y_ids = []
        for series_index, (series, target) in enumerate(zip(chart["series"], binding["series"]), 1):
            col_id = f"chart_{chart_index}_series_{series_index}"
            column = {"id": col_id, "name": series["name"], "formula": target["formula"]}
            if series.get("numberFormat") not in (None, "General"):
                column["format"] = presentation._series_format(series)
            columns.append(column)
            if series.get("formatFidelity") != "EXACT":
                unresolved.append({"kind": "chart-format", "sourceId": chart["id"],
                                   "series": series["name"], "sourceFormat": series.get("numberFormat"),
                                   "fidelity": series.get("formatFidelity"),
                                   "reason": series.get("formatReason") or "source style lacks verified exact mapping"})
            y_ids.append(col_id)
            for point_index, cell in enumerate(series["sourceCells"]):
                check_id = f"{chart['id']}-s{series_index}-p{point_index + 1}"
                checks.append({"id": check_id, "chartTitle": chart["title"],
                               "series": series["name"], "sourceCell": cell,
                               "sourceCategoryCell": chart["categoryCells"][point_index],
                               "sourceFormula": series["sourceFormulas"][point_index],
                               "sourceCachedExpectation": series["cachedExpectations"][point_index]
                               if point_index < len(series["cachedExpectations"]) else None,
                               "sourceCachedExpectationAuthority": "DIAGNOSTIC_SOURCE_CACHE",
                               "referenceQualificationStatus": "PENDING_INDEPENDENT_ORACLE",
                               "targetCategoryKey": binding["points"][point_index]["targetCategoryKey"],
                               "sourceOrder": point_index + 1,
                               "liveFormula": target["formula"], "unit": series["unit"],
                               "status": "PENDING_RUNTIME_QUERY"})
                checks_by_cell[cell].append(check_id)
        columns.append({"id": order_id, "name": "Source Order", "formula": binding["orderFormula"]})
        element = {"id": element_id,
                   "kind": "donut-chart" if chart["chartType"] == "doughnut" else f"{chart['chartType']}-chart",
                   "name": chart["title"],
                   "description": {"text": marker, "visibility": "hidden"},
                   "source": {"kind": "table", "elementId": binding["sourceElementId"]},
                   "columns": columns,
                   "style": {"backgroundColor": "#FFFFFF", "borderColor": {"kind": "theme", "ref": "colors-border"},
                             "borderWidth": 1, "borderRadius": "round"}}
        if chart["chartType"] == "doughnut":
            source_geometry = chart.get("sourceGeometry") or {}
            source_hole = source_geometry.get("holePercent")
            requested = source_hole / 100 if isinstance(source_hole, int) else None
            target_radius = min(0.8, max(0.4, requested)) if requested is not None else 0.6
            element.update({
                "color": {"columnId": category_id,
                          "sort": {"direction": "ascending", "by": order_id, "aggregation": "min"}},
                "value": {"columnId": y_ids[0]}, "innerRadius": target_radius,
            })
            if requested != target_radius:
                unresolved.append({"kind": "chart-geometry", "sourceId": chart["id"],
                                   "sourceHolePercent": source_hole,
                                   "targetInnerRadius": target_radius,
                                   "reason": "native donut innerRadius supports only 0.4–0.8; source hole cannot be exact"})
            if source_geometry.get("firstSliceAngle") is not None:
                unresolved.append({"kind": "chart-geometry", "sourceId": chart["id"],
                                   "sourceFirstSliceAngle": source_geometry["firstSliceAngle"],
                                   "reason": "native donut schema has no first-slice rotation field; visual review pending"})
        else:
            element.update({"xAxis": {"columnId": category_id,
                                      "sort": {"direction": "ascending", "by": order_id, "aggregation": "min"}},
                            "yAxis": {"columnIds": y_ids}})
        elements.append(element)
        emitted_ids.append(element_id)
        emitted_sources.append(binding["sourceElementId"])
    table_bindings = [("main", bindings["mainViewColumns"]),
                      ("detail", bindings["detailViewColumns"])]
    if "editViewColumns" in bindings:
        if not isinstance(bindings["editViewColumns"], list) or not bindings["editViewColumns"]:
            raise ValueError("editViewColumns requires explicit existing edit column IDs and names")
        table_bindings.append(("edit", bindings["editViewColumns"]))
    source_references = any(column.get("sourceCell") for _role, columns in table_bindings
                            for column in columns if isinstance(column, dict))
    workbook = None
    if source_references:
        source_path = Path(plan["source"]["path"])
        if hashlib.sha256(source_path.read_bytes()).hexdigest() != plan["sourceHash"]:
            raise ValueError("presentation source changed before width emission")
        workbook = load_workbook(source_path, data_only=False, read_only=False)
    try:
        sized = {role: _sized_table_columns(columns, role, workbook)
                 for role, columns in table_bindings}
    finally:
        if workbook is not None:
            workbook.close()
    for role, element_id in (("main", bindings["mainElementId"]),
                             ("detail", bindings["detailElementId"])):
        elements.append({"id": element_id, "kind": "table", "name": f"{plan['sourceName']} {role.title()}",
                         "description": {"text": marker, "visibility": "hidden"},
                         "source": {"kind": "table", "elementId": bindings["viewSourceElementId"]},
                         "columns": sized[role]})

    patches = defaultdict(lambda: {"columnUpdates": {}, "conditionalFormats": []})
    for column in sized.get("edit", []):
        patches[bindings["editElementId"]]["columnUpdates"][column["id"]] = {
            "style": {"width": column["style"]["width"]}}
        patches[bindings["editElementId"]].setdefault("expectedColumnNames", {})[column["id"]] = column["name"]
    for rule in plan["validations"]:
        targets = _target_list(plan, bindings, "validationTargets", rule,
                               reviewed=staging_gate["status"] != "LOGICAL_FIXTURE_ONLY")
        if rule["status"] != "supported" or not targets:
            unresolved.append({"kind": "validation", "sourceId": rule["id"],
                               "reason": "unsupported source rule" if rule["status"] != "supported"
                               else "missing reviewed editable target"})
            continue
        for target in targets:
            patches[target["elementId"]]["columnUpdates"].setdefault(target["columnId"], {}).update(
                copy.deepcopy(rule["emission"]))
        unresolved.append({"kind": "validation-null-behavior", "sourceId": rule["id"],
                           "reason": "source edit/clear route and persisted value require runtime review"})
    for rule in plan["conditionalFormats"]:
        targets = _target_list(plan, bindings, "conditionalTargets", rule,
                               reviewed=staging_gate["status"] != "LOGICAL_FIXTURE_ONLY")
        emission = rule.get("emission") or {}
        domain = emission.get("domain") or {}
        invalid_mid = ("mid" in domain and
                       (not isinstance(domain["mid"], (int, float)) or isinstance(domain["mid"], bool)
                        or not math.isfinite(domain["mid"])))
        if rule["status"] != "supported" or not targets or invalid_mid:
            unresolved.append({"kind": "conditionalFormat", "sourceId": rule["id"],
                               "reason": rule.get("unresolvedReason") or
                               ("unsupported source rule" if rule["status"] != "supported"
                                else "missing reviewed target" if not targets else "invalid scale midpoint"),
                               "relativeReference": rule["relativeReference"],
                               "priority": rule["priority"]})
            continue
        for target in targets:
            native = copy.deepcopy(emission)
            native["columnIds"] = [target["columnId"]]
            if native["type"] == "backgroundScale":
                native["domain"] = {key: value.replace("[Column]", f"[{target['columnId']}]")
                                    if isinstance(value, str) else value
                                    for key, value in native["domain"].items()}
            patches[target["elementId"]]["conditionalFormats"].append((rule["priority"], native))
    for patch in patches.values():
        patch["conditionalFormats"] = [native for _priority, native in sorted(
            patch["conditionalFormats"], key=lambda item: item[0], reverse=True)]

    formula_status = {item["id"]: item.get("mappingStatus") for item in (manifest or {}).get("formulaCells", [])}
    accounting = []
    pending = []
    for item in plan["formulaProvenance"]:
        cell = item["cell"]
        owned = checks_by_cell.get(cell, [])
        mapping_status = formula_status.get(cell, "unreviewed")
        if mapping_status in {"unresolved", "cached", "unreviewed"} or not owned:
            pending.append(cell)
        accounting.append({"sourceCell": cell, "formula": item["formula"],
                           "mappingStatus": mapping_status, "pendingCheckIds": owned,
                           "status": "PENDING_DEPENDENCY_REVIEW" if cell in pending
                           else "PENDING_RUNTIME_QUERY"})
    if pending:
        unresolved.append({"kind": "formula-scope", "sourceCells": pending,
                           "scopeHash": _digest(pending), "sourceCellCount": len(pending),
                           "reason": "source formula dependency lacks approved live ownership or mapping"})
    artifact = {
        "version": PLAN_VERSION, "sourceHash": plan["sourceHash"],
        "presentationPlanHash": plan["planHash"],
        "planHash": (manifest or {}).get("planHash"),
        "stagingReceiptHash": bindings.get("stagingReceiptHash"),
        "stagingGate": staging_gate,
        "pages": pages, "elements": elements, "layout": _layout(
            page_ids, title_ids, emitted_ids, bindings["mainElementId"],
            bindings["editElementId"], bindings["detailElementId"]),
        "pageRoles": dict(zip(("overview", "edit", "detail"), page_ids)),
        "editElementId": bindings["editElementId"],
        "referencedElementIds": sorted(set([bindings["editElementId"], bindings["viewSourceElementId"]]
                                           + emitted_sources)),
        "elementPatches": dict(patches), "aggregateChecks": checks,
        "ruleProvenance": copy.deepcopy(plan["conditionalFormats"]),
        "validationProvenance": copy.deepcopy(plan["validations"]),
        "formulaAccounting": accounting, "unresolved": unresolved,
        "ownershipMarker": marker,
        "displayVerificationStatus": "PENDING_RUNTIME_READBACK",
        "restPlan": {"method": "PUT", "path": "/v2/workbooks/{workbookId}/contents",
                     "documentId": bindings["documentId"],
                     "documentVersion": bindings["documentVersion"],
                     "requires": ["complete current contents", "reviewed live point bindings",
                                  "documentVersion"]},
    }
    artifact["ownershipFingerprints"] = {
        element["id"]: _digest({key: element.get(key) for key in ("id", "kind", "name", "body", "source")})
        for element in elements}
    artifact["artifactHash"] = _digest(artifact)
    return artifact


def emit_preview(plan, bindings):
    """Render synthetic logical bindings without granting a construction gate."""
    return _render(plan, bindings, staging_gate={"status": "LOGICAL_FIXTURE_ONLY"})


def _schema_parts(node):
    """Combine properties/requirements across OpenAPI allOf object fragments."""
    parts = node.get("allOf", [node]) if isinstance(node, dict) else []
    properties, required = {}, set()
    for part in parts:
        if isinstance(part, dict):
            properties.update(part.get("properties", {}))
            required.update(part.get("required", []))
    return properties, required


def _donut_variants(node):
    if isinstance(node, dict):
        properties, _required = _schema_parts(node)
        if "donut-chart" in properties.get("kind", {}).get("enum", []):
            yield node
        for value in node.values():
            yield from _donut_variants(value)
    elif isinstance(node, list):
        for value in node:
            yield from _donut_variants(value)


def _enum_values(node):
    if isinstance(node, dict):
        for value in node.get("enum", []):
            yield value
        for key in ("oneOf", "anyOf", "allOf"):
            for item in node.get(key, []):
                yield from _enum_values(item)
    elif isinstance(node, list):
        for value in node:
            yield from _enum_values(value)


def audit_donut_openapi(openapi):
    """Check the exact native donut fields used by this emitter against OpenAPI."""
    schemas = (openapi.get("components") or {}).get("schemas") or {}
    for variant in _donut_variants(schemas):
        properties, required = _schema_parts(variant)
        color, color_required = _schema_parts(properties.get("color", {}))
        value, value_required = _schema_parts(properties.get("value", {}))
        sort, _sort_required = _schema_parts(color.get("sort", {}))
        missing = []
        for name in ("id", "kind", "source", "columns", "color", "value"):
            if name not in required:
                missing.append(f"donut required {name}")
        if "innerRadius" not in properties:
            missing.append("donut innerRadius")
        if "columnId" not in color or "columnId" not in color_required:
            missing.append("donut color.columnId")
        if "columnId" not in value or "columnId" not in value_required:
            missing.append("donut value.columnId")
        for name in ("direction", "by", "aggregation"):
            if name not in sort:
                missing.append(f"donut color.sort.{name}")
        for name, expected in (("direction", "ascending"), ("aggregation", "min")):
            if name in sort and expected not in set(_enum_values(sort[name])):
                missing.append(f"donut color.sort.{name}={expected}")
        if not missing:
            return {"status": "PASS", "missing": []}
        return {"status": "FAIL", "missing": missing}
    return {"status": "FAIL", "missing": ["native donut-chart variant"]}


def emit_approved(plan, bindings, manifest, approval, receipt, openapi):
    """Render only after the existing manifest/source/input staging gates pass."""
    if plan.get("version") != PLAN_VERSION or plan.get("planHash") != _digest(
            {key: value for key, value in plan.items() if key != "planHash"}):
        raise ValueError("presentation plan hash is invalid")
    if plan != build_source_plan(plan["source"]["path"]):
        raise ValueError("presentation plan no longer matches the audited source")
    validate_manifest(manifest, plan["source"]["path"])
    if manifest["sourceHash"] != plan["sourceHash"]:
        raise ValueError("presentation source and manifest differ")
    if (bindings.get("manifestPlanHash") != manifest["planHash"]
            or bindings.get("stagingReceiptHash") != _digest(receipt)):
        raise ValueError("presentation binding manifest/receipt hash is stale")
    projection = copy.deepcopy(manifest["planProjection"])
    projection["sourceSha256"] = plan["sourceHash"]
    projection["sourcePath"] = plan["source"]["path"]
    _source, evidence = gate_plan_staging_source(projection, source_arg=plan["source"]["path"])
    expected = gate_input_staging(manifest, approval, projection, evidence)
    if receipt != expected:
        raise ValueError("approved staging receipt does not match shared gates")
    target = manifest.get("target") or {}
    if (target.get("documentId") != bindings.get("documentId")
            or target.get("version") != bindings.get("documentVersion")):
        raise ValueError("presentation target document/version is stale")
    audit = presentation.audit_presentation_openapi(openapi)
    if audit["status"] != "PASS":
        raise ValueError("presentation OpenAPI drift: " + ", ".join(audit["missing"]))
    if any(chart["chartType"] == "doughnut" for chart in plan["charts"]):
        donut_audit = audit_donut_openapi(openapi)
        if donut_audit["status"] != "PASS":
            raise ValueError("native donut OpenAPI drift: " + ", ".join(donut_audit["missing"]))
    return _render(plan, bindings, staging_gate={"status": "APPROVED_STAGING",
                                                  "receiptHash": _digest(receipt)}, manifest=manifest)
