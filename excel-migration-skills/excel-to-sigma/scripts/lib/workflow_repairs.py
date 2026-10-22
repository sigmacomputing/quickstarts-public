"""Source-bound secondary workflows and read-only presentation repairs.

Fragments and controller programs only: no auth, execution, cached formula values,
or declaration that construction establishes published acceptance.
"""
from __future__ import annotations
import copy
import hashlib
import json
import math
import re
import xml.etree.ElementTree as ET
from decimal import Decimal,ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries,get_column_letter
try:
    from .presentation import extract_presentation
    from .native_input_wiring import _select_rows,_column_id
    from .input_sync import _digest,_approved_feature
    from .migration_manifest import build_manifest,validate_manifest
except ImportError:
    from presentation import extract_presentation
    from native_input_wiring import _select_rows,_column_id
    from input_sync import _digest,_approved_feature
    from migration_manifest import build_manifest,validate_manifest

class RepairError(ValueError):pass


def reconciliation_required(card,ledger):
    def cents(v):
        if v is None:return 0
        if isinstance(v,bool) or not isinstance(v,(int,float,Decimal)) or not math.isfinite(float(v)):
            raise RepairError('currency must be finite numeric')
        return int((Decimal(str(v))*100).quantize(Decimal(1),rounding=ROUND_HALF_UP))
    return abs(cents(card)-cents(ledger))>1


def _table_literals(wb,sheet,headers,first,last):
    ws=wb[sheet];result=[]
    for r in range(first,last+1):
        row={name:ws.cell(r,c).value for c,name in enumerate(headers,1)}
        if any(isinstance(v,str) and v.startswith('=') for v in row.values()):
            raise RepairError('reference/input extraction must not freeze formulas')
        result.append(row)
    return result


def extract_workflow_repairs(source_path):
    path=Path(source_path).resolve();extracted=extract_presentation(path)
    workflow=extracted['workflow'];wb=load_workbook(path,data_only=False)
    source={'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'name':path.name}
    result={'version':1,'workflow':workflow,'source':source,'additionalInputs':{},'referenceTables':{},
        'secondaryRows':[],'formulaCells':[],'semanticFormats':extracted['semanticFormats'],
        'conditionalFormats':extracted['conditionalFormats'],'charts':extracted['charts'],'validations':extracted['validations'],'productionStatus':'BLOCKED'}
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type=='f':
                    kind='source-audit' if ws.title=='Audit' else 'secondary-live' if ws.title in ('Rep Targets','Purchase Plan','Reconciliation','Monthly Model') else 'existing-primary-or-chart'
                    result['formulaCells'].append({'id':ws.title+'!'+cell.coordinate,'sourceHash':source['sha256'],
                        'sheet':ws.title,'address':cell.coordinate,'originalFormula':cell.value,'route':kind,
                        'verification':'UNRESOLVED','sourceNumberFormat':cell.number_format})
    if workflow=='sales':
        rows=_table_literals(wb,'Rep Targets',['Rep','Annual Target'],2,5)
        result['additionalInputs']['salesRepTargets']={'kind':'input-table','name':'Sales Rep Targets',
            'keyColumns':['Rep'],'rows':rows,'columns':[{'name':'Rep','type':'text','provenance':'Rep Targets!A2:A5'},
                {'name':'Annual Target','type':'number','provenance':'Rep Targets!B2:B5'}],
            'editMode':'published-explore','frame':[1,34,24,12],'sourceRange':'Rep Targets!A1:B5'}
        result['referenceTables']['Pipeline Stages']={'columns':['Stage','Default Probability','Forecast Category'],
            'rows':_table_literals(wb,'Assumptions',['Stage','Default Probability','Forecast Category'],2,7),
            'sourceRange':'Assumptions!A1:C7','dependency':'reference only; per-opportunity Probability remains authored'}
    elif workflow=='inventory':
        result['referenceTables']['Suppliers']={'columns':['Supplier','Standard Lead Days','Minimum Order','Buyer'],
            'rows':_table_literals(wb,'Suppliers',['Supplier','Standard Lead Days','Minimum Order','Buyer'],2,5),
            'sourceRange':'Suppliers!A1:D5','dependency':'reference only; no invented reorder dependency'}
        result['secondaryRows']=[{'SKU':wb['Inventory'].cell(r-3,1).value,'Buyer Review':wb['Purchase Plan'].cell(r,8).value,
                                 'Source Row':r} for r in range(5,41)]
    elif workflow=='expense':
        result['secondaryRows']=[{'Expense ID':wb['Expense Register'].cell(r-3,1).value,'Source Row':r} for r in range(5,53)]
        result['currencyDecision']={'sourcePrecision':2,'mappingKind':'reviewed-replacement-required',
            'sourceComparison':'Abs(Card-Ledger)>0.01','replacement':'Abs(Round(Card*100,0)-Round(Ledger*100,0))>1',
            'domain':'currency inputs at cents; preserve stored Amount/Card/Ledger values','nativeFormula':'Abs(Round([Card Amount]*100,0)-Round([Ledger Amount]*100,0)) > 1'}
    elif workflow=='forecast':
        ws=wb['Monthly Model'];result['months']=[ws.cell(5,c).value for c in range(3,15)]
        result['displayRows']=[{'Section':str(ws.cell(r,1).value or 'Derived'),'Line Item':ws.cell(r,2).value,
            'Source Row':r,'Scope':'all'} for r in range(8,29) if ws.cell(r,2).value]
    elif workflow=='research':
        result['inputAuditPolicy']='metadata spine joined to Scope=shared wide input; live Switch values, no caches'
    wb.close();result['contractHash']=_digest(result)
    return result


def validate_repair_contract(contract):
    if contract.get('version')!=1 or contract.get('productionStatus')!='BLOCKED':raise RepairError('invalid repair gate')
    if _digest({k:v for k,v in contract.items() if k!='contractHash'})!=contract.get('contractHash'):
        raise RepairError('repair contract hash changed')
    if hashlib.sha256(Path(contract['source']['path']).read_bytes()).hexdigest()!=contract['source']['sha256']:
        raise RepairError('repair source changed')


def _sql_table(eid,name,rows,connection):
    fields=list(rows[0]);return {'id':eid,'kind':'table','name':name,
        'source':{'kind':'sql','connectionId':connection,'statement':_select_rows(rows,fields)},
        'columns':[{'id':_column_id(k),'name':k,'formula':f'[Custom SQL/C{i}]'} for i,k in enumerate(fields)]}


def _project(eid,name,source,columns):
    return {'id':eid,'kind':'table','name':name,'source':{'kind':'table','elementId':source['elementId']},
        'columns':[{'id':_column_id(label),'name':label,'formula':formula} for label,formula in columns]}


def _join(eid,name,left,right,leftkey,rightkey,columns):
    return {'id':eid,'kind':'table','name':name,'source':{'kind':'join','name':'Base',
        'primarySource':{'kind':'table','elementId':left},'joins':[{'name':'Detail','joinType':'left-outer',
        'left':{'kind':'table','elementId':left},'right':{'kind':'table','elementId':right},
        'columns':[{'left':'['+leftkey+']','right':'['+rightkey+']'}]}]},
        'columns':[{'id':_column_id(label),'name':label,'formula':formula} for label,formula in columns]}


def _page(eid,title,elements):
    cursor=4;layout=f'<Page type="grid" id="{eid}" gridTemplateColumns="repeat(24, 1fr)">'
    header={'id':eid+'-title','kind':'text','body':'# '+title}
    layout+=f'<Element elementId="{header["id"]}" gridColumn="1 / 25" gridRow="1 / 4"/>'
    for element in elements:
        height=10 if element['kind']=='button' else 30
        layout+=f'<Element elementId="{element["id"]}" gridColumn="1 / 25" gridRow="{cursor} / {cursor+height}"/>'
        cursor+=height
    return {'pages':[{'id':eid,'name':title}],'elements':[header,*elements],'layout':layout+'</Page>'}


def build_workflow_repair_fragments(contract,bindings,*,connection_id):
    validate_repair_contract(contract)
    workflow=contract['workflow'];elements=[]
    calc=bindings.get('calculation',{})
    if workflow!='research' and (not calc.get('elementId') or not calc.get('name')):raise RepairError('actual calculation binding required')
    if workflow=='forecast':
        line_cols=[('Section',f'[{calc["name"]}/Section]'),('Line Item',f'[{calc["name"]}/Line Item]')]+[(m,f'SumIf([{calc["name"]}/Amount], [{calc["name"]}/Month] = {json.dumps(m)})') for m in contract['months']]
        line_summary=_project('forecast-monthly-line-results','Forecast Monthly Line Results',calc,line_cols)
        line_summary['groupings']=[{'id':'lines','groupBy':['section','line_item'],'calculations':[_column_id(m) for m in contract['months']]}]
        elements.append(line_summary)
        spine=_sql_table('forecast-pnl-rows','P&L Source Rows',contract['displayRows'],connection_id);elements.append(spine)
        section_amount=lambda section,month:f'SumIf([{calc["name"]}/Amount], [{calc["name"]}/Section] = {json.dumps(section)} and [{calc["name"]}/Month] = {json.dumps(month)})'
        cols=[('Source Row','[P&L Source Rows/Source Row]'),('Section','[P&L Source Rows/Section]'),('Line Item','[P&L Source Rows/Line Item]')]
        for month in contract['months']:
            rev,cogs,opex,below=[section_amount(s,month) for s in ('Revenue','Cost of Revenue','Operating Expenses','Below the Line')]
            base=f'Lookup([Forecast Monthly Line Results/{month}], [P&L Source Rows/Line Item], [Forecast Monthly Line Results/Line Item])'
            formula=f'If([Source Row] < 20, {base}, [Source Row] = 22, {rev}, [Source Row] = 23, {cogs}, [Source Row] = 24, {opex}, [Source Row] = 25, {below}, [Source Row] = 27, {rev}-{cogs}-{opex}, {rev}-{cogs}-{opex}-{below})'
            cols.append((month,formula))
        cols.append(('FY Total',' + '.join('['+m+']' for m in contract['months'])))
        table=_project('forecast-monthly-pnl','Monthly P&L (USD)',{'elementId':spine['id']},cols)
        table['sort']=[{'columnId':'source_row','direction':'ascending'}]
        for c in table['columns'][3:]:c['format']=copy.deepcopy(contract['semanticFormats']['Monthly Model!C8']['emission'])
        elements.append(table)
    elif workflow=='sales':
        targets=bindings.get('repTargets',{})
        if not targets.get('elementId'):raise RepairError('assigned editable Rep Targets required')
        summary=_project('sales-rep-pipeline','Weighted Pipeline by Rep',calc,[('Rep',f'[{calc["name"]}/Rep]'),
            ('Weighted Pipeline',f'Sum([{calc["name"]}/Weighted ACV])')])
        summary['groupings']=[{'id':'rep','groupBy':['rep'],'calculations':['weighted_pipeline']}];elements.append(summary)
        view=_join('sales-rep-coverage','Rep Targets and Coverage',targets['elementId'],summary['id'],'Rep','Rep',
            [('Rep','[Base/Rep]'),('Annual Target','[Base/Annual Target]'),('Weighted Pipeline','[Detail/Weighted Pipeline]'),
             ('Coverage','[Weighted Pipeline] / [Annual Target]')])
        elements.append(view)
    elif workflow=='inventory':
        meta=_sql_table('inventory-purchase-source','Purchase Plan Source',contract['secondaryRows'],connection_id);elements.append(meta)
        view=_join('inventory-purchase-plan','Purchase Plan',calc['elementId'],meta['id'],'SKU','SKU',[
            ('SKU','[Base/SKU]'),('Product','[Base/Product]'),('Supplier','[Base/Supplier]'),('Status','[Base/Status]'),
            ('Recommended Qty','[Base/Reorder Qty]'),('Unit Cost','[Base/Unit Cost]'),('Order Cost','[Recommended Qty] * [Unit Cost]'),
            ('Buyer Review','[Detail/Buyer Review]'),('Source Row','[Detail/Source Row]')])
        view['sort']=[{'columnId':'source_row','direction':'ascending'}];elements.append(view)
    elif workflow=='expense':
        cols=[('Expense ID','Expense ID'),('Employee','Employee'),('Department','Department'),('Category','Category'),
            ('Amount','Amount'),('Approval','Approval Status'),('Exception','Exception'),('Variance','Recon Variance')]
        view=_project('expense-reconciliation','Reconciliation',calc,[(a,f'[{calc["name"]}/{b}]') for a,b in cols]+[
            ('Required Action','Switch([Exception], "CLEAR", "None", "RECON", "Match card and ledger", "OVER LIMIT", "Manager review", "MISSING RECEIPT", "Attach receipt", "Complete approval")')])
        elements.append(view)
    for i,(name,table) in enumerate(contract['referenceTables'].items()):
        elements.append(_sql_table(workflow+'-reference-'+str(i),name,table['rows'],connection_id))
    dashboard=build_dashboard_view(contract,calc,monthly={'elementId':'forecast-monthly-pnl','name':'Monthly P&L (USD)'} if workflow=='forecast' else None,connection_id=connection_id)
    if dashboard:elements.extend(dashboard['elements'])
    for element in elements:
        element['description']={'visibility':'hidden','text':'excel-to-sigma:task6:'+workflow+':'+contract['source']['sha256']}
    result=_page(workflow+'-workflow',workflow.title()+' workflow',elements)
    result.update(sourceHash=contract['source']['sha256'],contractHash=contract['contractHash'],productionStatus='BLOCKED')
    return result


def build_normalized_input_audit(contract,input_spec,target,*,connection_id):
    validate_repair_contract(contract)
    if not target.get('elementId') or not re.fullmatch('[a-f0-9]{64}',str(target.get('systemEvidenceHash',''))):
        raise RepairError('reviewed physical row-ID evidence required')
    columns=input_spec['columns'];expected={c['name'] for c in columns}
    if set(target.get('columnIds',{}))!=expected or not target.get('systemColumnId'):
        raise RepairError('complete assigned input columns/system ID required')
    wb=load_workbook(contract['source']['path'],data_only=False);rows=[]
    try:
        for i,col in enumerate(columns):
            if col['name']=='Scope':continue
            source=col.get('sourceCellId') or col.get('provenance')
            if not source or '!' not in source:raise RepairError('input audit requires exact source cell provenance')
            sheet,address=source.rsplit('!',1)
            if sheet not in wb:raise RepairError('input audit source sheet missing')
            cell=wb[sheet][address]
            if cell.data_type=='f':raise RepairError('input audit cannot disguise source formula as an input')
            labels=[wb[sheet].cell(cell.row,c).value for c in range(1,cell.column)
                    if isinstance(wb[sheet].cell(cell.row,c).value,str) and wb[sheet].cell(cell.row,c).data_type!='f']
            if sheet=='FY results':labels.append(str(wb[sheet].cell(5,cell.column).value))
            rows.append({'Scope':'shared','SourceCell':source,'Label':' / '.join(labels)+' ('+source+')',
                'SourceNumberFormat':cell.number_format,'Unit':'percent' if '%' in cell.number_format else 'currency' if '$' in cell.number_format else 'number',
                'LogicalInputId':col['name'],'SourceOrder':i})
    finally:wb.close()
    if not rows:raise RepairError('input audit requires numeric source cells')
    meta=_sql_table('research-input-metadata','Research Source Input Metadata',rows,connection_id)
    switch='Switch([Base/LogicalInputId], '+', '.join(json.dumps(r['LogicalInputId'])+', [Detail/'+r['LogicalInputId']+']' for r in rows)+', Null)'
    audit=_join('research-normalized-input-audit','Research Input Audit',meta['id'],target['elementId'],'Scope','Scope',
        [(n,'[Base/'+n+']') for n in ('Scope','SourceCell','Label','SourceNumberFormat','Unit','SourceOrder')]+[
            ('RawValue',switch),('StoredRowID','[Detail/'+target['systemColumnId']+']')])
    order=next(c for c in audit['columns'] if c['name']=='SourceOrder')
    order['hidden']=True
    audit['sort']=[{'columnId':order['id'],'direction':'ascending'}]
    growth=[r for r in rows if r['SourceCell'].startswith('Model Inputs!')]
    guidance={'id':'research-input-guidance','kind':'text','body':'## Editable research assumptions\n'+ '\n'.join('- '+r['LogicalInputId']+': '+r['Label']+' ('+r['Unit']+')' for r in growth)+'\n\nThe source inputs audit below exposes every live source value and its stored row identity. Fixed gross profit and designated historical/cash-flow literals remain source inputs.'}
    result=_page('research-input-audit-normalized','Research source inputs',[guidance,meta,audit])
    result.update(metadataRows=rows,sourceHash=contract['source']['sha256'],productionStatus='BLOCKED')
    return result


def build_input_id_audit(input_spec,target):
    if not target.get('elementId') or not target.get('systemColumnId') or not target.get('systemEvidenceHash'):
        raise RepairError('input ID projection requires assigned physical-ID evidence')
    return _project('audit-'+target['elementId'],input_spec['name']+' Audit',target,
        [(c['name'],'['+input_spec['name']+'/'+c['name']+']') for c in input_spec['columns']]+[
        ('StoredRowID','['+input_spec['name']+'/'+target['systemColumnId']+']')])


def propagate_readonly_presentation(contract,element,source_cells,*,technical_source=None):
    validate_repair_contract(contract)
    if element.get('kind') not in ('table','pivot-table'):raise RepairError('read-only table required')
    result=copy.deepcopy(element);ids={c['id'] for c in result['columns']}
    if not set(source_cells)<=ids:raise RepairError('format target column does not exist')
    for column in result['columns']:
        source=source_cells.get(column['id'])
        if source:
            semantic=contract['semanticFormats'].get(source)
            if semantic:column['format']=copy.deepcopy(semantic['emission'])
    workbook=load_workbook(contract['source']['path'],data_only=False)
    try:
        rules=[]
        for rule in contract['conditionalFormats']:
            if rule.get('status')!='supported':continue
            lowc,lowr,highc,highr=range_boundaries(rule['range'])
            for cid,source in source_cells.items():
                sheet,address=source.rsplit('!',1);cell=workbook[sheet][address]
                if rule['sheet']==sheet and lowc<=cell.column<=highc and lowr<=cell.row<=highr:
                    emission=copy.deepcopy(rule['emission']);emission['columnIds']=[cid]
                    if 'domain' in emission:
                        emission['domain']={k:v.replace('[Column]','['+cid+']') if isinstance(v,str) else v for k,v in emission['domain'].items()}
                    rules.append((rule['priority'],emission))
        result.setdefault('conditionalFormats',[]).extend(r for _,r in sorted(rules,key=lambda p:p[0],reverse=True))
    finally:workbook.close()
    # Technical operation metadata is audit-only; business fields remain intact.
    owned_source=None
    if technical_source is not None:
        owned_source=technical_source.get('element',{})
        if (contract['workflow']!='expense' or technical_source.get('viewElementHash')!=_digest(element)
                or technical_source.get('elementHash')!=_digest(owned_source) or owned_source.get('kind')!='table'
                or owned_source.get('name')!='Expense Register'
                or owned_source.get('description',{}).get('text')!='excel-to-sigma:native:expense:'+contract['source']['sha256']
                or element.get('source')!={'kind':'table','elementId':owned_source.get('id')}):raise RepairError('owned technical presentation source/view binding stale or foreign')
    for col in result['columns']:
        if col.get('name') in ('Last Approval Operation ID','Applied Approval Operation IDs','StoredRowID'):
            col['hidden']=True
        elif owned_source is not None:
            reference=re.fullmatch(r'\[([^\[\]/]+)/([^\[\]/]+)\]',col.get('formula',''))
            if reference and reference[1]==owned_source['name'] and reference[2] in {'Last Approval Operation ID','Applied Approval Operation IDs'}:
                if col.get('name') not in (None,reference[2]):raise RepairError('conflicting owned technical column identity')
                col['hidden']=True
    return result


def _layout_root(layout):
    return ET.fromstring('<Layout>'+re.sub(r'<\?xml[^?]*\?>','',layout or '')+'</Layout>')


def cleanup_owned_scaffolds(contents,review):
    if not str(review.get('owner') or '').strip():raise RepairError('named cleanup review required')
    result=copy.deepcopy(contents);pages={p['id']:p for p in result['pages']};elements={e['id']:e for e in result['elements']}
    root=_layout_root(result.get('layout'))
    for eid,digest in review.get('elements',{}).items():
        if eid not in elements or _digest(elements[eid])!=digest:raise RepairError('owned scaffold element changed')
    for pid,digest in review.get('pages',{}).items():
        if pid not in pages or _digest(pages[pid])!=digest:raise RepairError('owned scaffold page changed')
        placements=[n.get('elementId') or n.get('id') for p in root.iter('Page') if p.get('id')==pid for n in p.iter('Element')]
        if any(eid not in review.get('elements',{}) for eid in placements):raise RepairError('scaffold contains foreign/unreviewed element')
        if any(elements[eid].get('kind')=='input-table' for eid in placements):raise RepairError('cannot hide active input storage scaffold')
        pages[pid]['visibility']='hidden'
    landing=review.get('landingPageId')
    if landing not in pages or pages[landing].get('visibility')=='hidden':raise RepairError('visible actual landing page required')
    result['pages']=[pages[landing]]+[p for p in result['pages'] if p['id']!=landing]
    return result


def shrink_input_layout(contents,patches):
    """Only shrink individually reviewed input placements, preserving all other XML."""
    try:
        from .native_actions import _layout_element_owners
    except ImportError:
        from native_actions import _layout_element_owners
    result=copy.deepcopy(contents);elements={e['id']:e for e in contents['elements']}
    for eid,patch in patches.items():
        element=elements.get(eid,{})
        if element.get('kind')!='input-table' or patch.get('elementHash')!=_digest(element):raise RepairError('input layout hash/kind mismatch')
        root=_layout_root(result['layout']);owners=_layout_element_owners(root,eid)
        if len(owners)!=1 or owners[0].get('id')!=patch.get('pageId'):raise RepairError('input layout owner mismatch')
        nodes=[n for n in owners[0].iter('Element') if (n.get('elementId') or n.get('id'))==eid]
        if len(nodes)!=1 or nodes[0].get('gridRow')!=patch.get('fromGridRow'):raise RepairError('input placement changed')
        old=re.fullmatch(r'(\d+) / (\d+)',patch.get('fromGridRow',''));new=re.fullmatch(r'(\d+) / (\d+)',patch.get('toGridRow',''))
        if not old or not new or int(new[1])!=int(old[1]) or not int(new[1])+14<=int(new[2])<=int(old[2]):raise RepairError('bounded same-start input height reduction requires at least14 grid rows for editable data')
        tags=[]
        for match in re.finditer(r'<Element\b[^>]*>',result['layout']):
            raw=match.group();node=ET.fromstring(raw[:-1].rstrip('/ ')+' />')
            if (node.get('elementId') or node.get('id'))==eid:tags.append((match,raw))
        if len(tags)!=1:raise RepairError('ambiguous raw placement')
        match,raw=tags[0];replacement=re.sub(r"(gridRow=[\"'])[^\"']*([\"'])",lambda m:m[1]+patch['toGridRow']+m[2],raw)
        result['layout']=result['layout'][:match.start()]+replacement+result['layout'][match.end():]
    return result


def repair_research_edit_layout(contents,*,input_element_id,guidance_element_id,page_id,input_element_hash,guidance_element_hash,layout_hash):
    """Reviewed Research edit layout: visible data rows and a separate guide.

    GridRow is start/end, not start/height. Keep all element objects and all
    unrelated layout bytes. Unknown/overlapping full-width placements block.
    """
    if layout_hash!=_digest(contents['layout']) or input_element_id==guidance_element_id:raise RepairError('research layout stale/ambiguous')
    elements={e['id']:e for e in contents['elements']}
    if (elements.get(input_element_id,{}).get('kind')!='input-table' or _digest(elements.get(input_element_id))!=input_element_hash
            or elements.get(guidance_element_id,{}).get('kind')!='text' or _digest(elements.get(guidance_element_id))!=guidance_element_hash):raise RepairError('research layout source kind/hash mismatch')
    root=_layout_root(contents['layout']);pages=[p for p in root.iter('Page') if p.get('id')==page_id]
    if len(pages)!=1:raise RepairError('research edit page ambiguous')
    desired={input_element_id:(4,18),guidance_element_id:(20,34)}
    for eid in desired:
        allnodes=[n for n in root.iter('Element') if (n.get('elementId') or n.get('id'))==eid]
        if len(allnodes)!=1 or allnodes[0] not in list(pages[0].iter('Element')):raise RepairError('research layout wrong/ambiguous owner')
    for node in pages[0].iter('Element'):
        eid=node.get('elementId') or node.get('id')
        if eid in desired:continue
        row=re.fullmatch(r'(\d+) / (\d+)',node.get('gridRow',''))
        if not row or any(int(row[1])<end and start<int(row[2]) for start,end in desired.values()):raise RepairError('research edit layout foreign/unknown placement overlap')
    result=copy.deepcopy(contents)
    for eid,(start,end) in desired.items():
        matches=[]
        for match in re.finditer(r'<Element\b[^>]*>',result['layout']):
            node=ET.fromstring(match.group()[:-1].rstrip('/ ')+' />')
            if (node.get('elementId') or node.get('id'))==eid:matches.append(match)
        if len(matches)!=1 or len(re.findall(r'gridRow=',matches[0].group()))!=1:raise RepairError('research raw grid placement ambiguous')
        match=matches[0];raw=re.sub(r'(gridRow=[\"\'])[^\"\']*([\"\'])',lambda m:m[1]+str(start)+' / '+str(end)+m[2],match.group())
        result['layout']=result['layout'][:match.start()]+raw+result['layout'][match.end():]
    return result


def extend_repair_manifest(base_manifest,receipt,contract,review):
    validate_manifest(base_manifest);validate_repair_contract(contract)
    feature=base_manifest['planProjection']['sharedInputArchitecture'].get('presentationInputFeature')
    _approved_feature(base_manifest,receipt,feature,require_assigned=True)
    if (base_manifest['sourceHash']!=contract['source']['sha256'] or review.get('status')!='APPROVED_STAGING'
            or review.get('contractHash')!=contract['contractHash'] or not review.get('owner')):
        raise RepairError('named source-bound workflow repair review required')
    arch=copy.deepcopy(base_manifest['planProjection']['sharedInputArchitecture'])
    for feature,spec in contract['additionalInputs'].items():
        if feature in arch['inputs']:raise RepairError('new input would overwrite existing spec')
        spec=copy.deepcopy(spec);spec['connectionId']=arch['targetScope']['connectionId']
        arch['inputs'][feature]=spec;arch['mappedFeatures'].append(feature);arch['featureToInput'][feature]=feature
    projection=copy.deepcopy(base_manifest['planProjection']);projection['sharedInputArchitecture']=arch
    projection['workflowRepairArchitecture']={'contract':copy.deepcopy(contract),'review':copy.deepcopy(review),'newInputPages':{f:f+'-edit' for f in contract['additionalInputs']},'productionStatus':'BLOCKED'}
    projection['formulaCells']=copy.deepcopy(base_manifest['formulaCells'])
    result=build_manifest(base_manifest['source']['path'],projection);result['target']=copy.deepcopy(base_manifest['target'])
    return result


def build_dashboard_view(contract,calculation,*,monthly=None,connection_id):
    """Translate exact source dashboard expressions using live keyed cell lookups.

    Explicit SUM argument lists become scalar addition, avoiding a bogus
    multi-argument aggregate. Source-authored literal expressions remain static.
    """
    try:
        from . import excel_expression as ex
    except ImportError:
        import excel_expression as ex
    validate_repair_contract(contract);wb=load_workbook(contract['source']['path'],data_only=False)
    workflow=contract['workflow'];sheet='Overview' if workflow=='sales' else 'Dashboard'
    if sheet not in wb:wb.close();return None
    primary={'sales':'Pipeline','inventory':'Inventory','expense':'Expense Register'}.get(workflow)
    col_names={};keys={}
    if primary:
        col_names={get_column_letter(c.column):c.value for c in wb[primary][1]}
        keys={r:wb[primary].cell(r,1).value for r in range(2,wb[primary].max_row+1)}
    columns=[];provenance=[]
    def ref(r):
        target_sheet=r.sheet or sheet
        if target_sheet==sheet:
            return '[metric_'+r.address.lower()+']'
        if primary and target_sheet==primary:
            cell=wb[primary][r.address];field=col_names[get_column_letter(cell.column)]
            return 'Lookup(['+calculation['name']+'/'+field+'], '+json.dumps(keys[cell.row])+', ['+calculation['name']+'/'+str(col_names['A'])+'])'
        if workflow=='forecast' and target_sheet=='Monthly Model':
            if not monthly:raise RepairError('computed monthly target required for dashboard')
            cell=wb[target_sheet][r.address];label=wb[target_sheet].cell(cell.row,2).value
            period='FY Total' if cell.column==15 else wb[target_sheet].cell(5,cell.column).value
            return 'Lookup(['+monthly['name']+'/'+period+'], '+json.dumps(label)+', ['+monthly['name']+'/Line Item])'
        raise RepairError('unmapped dashboard source reference '+str(r))
    def render(n):
        if isinstance(n,ex.Number):return n.source
        if isinstance(n,ex.Text):return json.dumps(n.value)
        if isinstance(n,ex.Reference):return ref(n.value)
        if isinstance(n,ex.Range):return ' + '.join(ref(r) for r in ex.expand_range(n))
        if isinstance(n,ex.Group):return '('+render(n.expression)+')'
        if isinstance(n,ex.Unary):return n.operator+render(n.operand)
        if isinstance(n,ex.Percent):return '('+render(n.operand)+' * 0.01)'
        if isinstance(n,ex.Binary):return '('+render(n.left)+' '+n.operator+' '+render(n.right)+')'
        if isinstance(n,ex.Function) and n.name=='SUM':return '('+' + '.join(render(v) for v in n.arguments)+')'
        raise RepairError('unsupported dashboard source expression')
    try:
        for row in wb[sheet]:
            if len(row)<2 or not isinstance(row[0].value,str):continue
            cell=row[1]
            if cell.value is None:continue
            if cell.data_type=='f':
                if re.fullmatch(r"=COUNT\('Expense Register'!G2:G49\)",cell.value):
                    formula=f'Count([{calculation["name"]}/Amount])';kind='live-numeric-count'
                else:
                    node=ex.parse_expression(cell.value,sheet=sheet);formula=render(node)
                    kind='source-static-expression' if not ex.extract_references(node) else 'live-source-expression'
            elif isinstance(cell.value,(int,float)):
                formula=str(cell.value);kind='source-authored-literal'
            else:continue
            col={'id':'metric_'+cell.coordinate.lower(),'name':row[0].value,'formula':formula}
            semantic=contract['semanticFormats'].get(sheet+'!'+cell.coordinate)
            if semantic:col['format']=copy.deepcopy(semantic['emission'])
            columns.append(col);provenance.append({'sourceCell':sheet+'!'+cell.coordinate,'kind':kind,'formula':cell.value})
    finally:wb.close()
    source=_sql_table(workflow+'-dashboard-scope',workflow.title()+' Dashboard Scope',[{'Scope':'shared'}],connection_id)
    result={'id':workflow+'-source-dashboard','kind':'table','name':workflow.title()+' Source Dashboard',
        'source':{'kind':'table','elementId':source['id']},'columns':columns}
    return {'elements':[source,result],'provenance':provenance}


def build_navigation(pages,*,prefix='workflow-navigation'):
    if not pages or any(not p.get('id') or not p.get('name') for p in pages):raise RepairError('actual destination pages required')
    return [{'id':prefix+'-'+str(i),'kind':'button','appearance':'filled','text':p['name'],'actions':[{'id':prefix+'-action-'+str(i),
        'trigger':'on-click','effects':[{'effect':'navigate','target':{'type':'page','page':p['id']}}],
        'successToast':{'showMessage':'hidden'}}]} for i,p in enumerate(pages)]


def build_published_input_mode_program(manifest,receipt,feature,*,review):
    spec,target,_=_approved_feature(manifest,receipt,feature,require_assigned=True)
    if (review.get('status')!='APPROVED_STAGING' or not review.get('owner')
            or review.get('sourceHash')!=manifest['sourceHash'] or review.get('planHash')!=manifest['planHash']
            or review.get('targetBindingHash')!=_digest(target) or review.get('mode')!='published-explore'):
        raise RepairError('exact restricted published-mode review required')
    return '\n'.join(['from sigma.page import Page','from sigma.inputTable import InputTable',
        f'page = sigma.workbook.get_page({target["pageId"]!r})',"if not isinstance(page, Page): raise RuntimeError('expected assigned page')",
        f'table = page.get_element({target["elementId"]!r})',"if not isinstance(table, InputTable): raise RuntimeError('expected assigned input')",
        f'if table.name() != {spec["name"]!r}: raise RuntimeError("input name changed")',
        "table.set_edit_mode('published-explore')",'print(table.edit_mode())'])+'\n'


def prepare_fragment_update(current,fragments,*,scope,source_hash,plan_hash,review):
    """Conservative new-object merge; existing edits require the higher-level adapter."""
    if (current.get('workbookId')!=scope.get('documentId') or current.get('documentVersion')!=scope.get('version')
            or not scope.get('folderId') or not scope.get('folderEvidenceId') or not scope.get('connectionId')):
        raise RepairError('fresh designated staging scope/folder evidence required')
    try:
        from .action_inventory import validate_action_inventory,action_restore_programs
    except ImportError:
        from action_inventory import validate_action_inventory,action_restore_programs
    actions=review.get('sdkActionInventory')
    validate_action_inventory(actions,current,require_recovery=True)
    doc=current.get('contents')
    if not isinstance(doc,dict) or not review.get('owner') or review.get('beforeHash')!=_digest(doc):raise RepairError('named current-content review required')
    merged=copy.deepcopy(doc);pages={p['id'] for p in doc.get('pages',[])};elements={e['id'] for e in doc.get('elements',[])}
    if len(pages)!=len(doc.get('pages',[])) or len(elements)!=len(doc.get('elements',[])):raise RepairError('duplicate source IDs')
    added_pages=[];added_elements=[]
    for fragment in fragments:
        newp=[p['id'] for p in fragment['pages']];newe=[e['id'] for e in fragment['elements']]
        if len(set(newp))!=len(newp) or len(set(newe))!=len(newe) or pages&set(newp) or elements&set(newe):raise RepairError('existing/duplicate object refused')
        pages.update(newp);elements.update(newe);added_pages+=newp;added_elements+=newe
        merged.setdefault('pages',[]).extend(copy.deepcopy(fragment['pages']));merged.setdefault('elements',[]).extend(copy.deepcopy(fragment['elements']))
        merged['layout']=merged.get('layout','')+'\n'+fragment['layout']
    proposal={'workbookId':scope['documentId'],'documentVersion':scope['version'],'scope':copy.deepcopy(scope),
        'sourceHash':source_hash,'planHash':plan_hash,'beforeHash':_digest(doc),'beforeContents':copy.deepcopy(doc),
        'sdkActions':copy.deepcopy(actions),'sdkActionRestorationArtifacts':action_restore_programs(actions),'review':copy.deepcopy(review),'addedPageIds':added_pages,'addedElementIds':added_elements,
        'dryRunRequest':{'name':current.get('name','Reviewed workflow repairs'),'folderId':scope['folderId'],'contents':merged,'dryRun':True},
        'productionStatus':'BLOCKED'}
    proposal['proposalHash']=_digest(proposal);return proposal


def prepare_repair_update(manifest,receipt,contract,current,bindings,*,scope,review):
    """Bind complete repair generation to approved actual input and calculation IDs.

    Review includes exact source-element hashes and complete physical row-ID
    evidence. Existing display changes and scaffold hiding are individually bound.
    Unknown pages/elements and all business input definitions remain intact.
    """
    try:
        from .custom_ports import _connection_ids,_source_graph_hashes
        from .native_actions import _layout_element_owners
    except ImportError:
        from custom_ports import _connection_ids,_source_graph_hashes
        from native_actions import _layout_element_owners
    validate_manifest(manifest);validate_repair_contract(contract)
    if (manifest['sourceHash']!=contract['source']['sha256'] or manifest.get('target')!={'documentId':scope['documentId'],'version':scope['version']}
        or review.get('status')!='APPROVED_STAGING' or review.get('contractHash')!=contract['contractHash']
        or review.get('planHash')!=manifest['planHash']):raise RepairError('repair source/plan/target review mismatch')
    elements={e['id']:e for e in current['contents']['elements']}
    for role,binding in bindings.items():
        element=elements.get(binding.get('elementId'))
        if (not element or element.get('kind') not in {'table','input-table'} or binding.get('name')!=element.get('name')
            or binding.get('elementHash')!=_digest(element) or review.get('sourceBindings',{}).get(role)!=binding
            or _connection_ids(element,elements)!={scope['connectionId']}):raise RepairError('source graph/title/kind/connection/hash binding mismatch')
    fragments=[];arch=manifest['planProjection']['sharedInputArchitecture'];audit_elements=[]
    if arch['targetScope']!={'connectionId':scope['connectionId'],'folderId':scope['folderId']}:raise RepairError('designated repair connection/folder differs from source-bound architecture')
    for feature in arch['inputs']:
        spec,target,_=_approved_feature(manifest,receipt,feature,require_assigned=True)
        if target['documentId']!=scope['documentId'] or target['version']!=scope['version']:raise RepairError('stale input target')
        actual=elements.get(target['elementId'])
        if not actual or actual.get('kind')!='input-table' or _connection_ids(actual,elements)!={scope['connectionId']} or [p.get('id') for p in _layout_element_owners(_layout_root(current['contents']['layout']),target['elementId'])]!=[target['pageId']]:
            raise RepairError('assigned input kind/page changed')
        system=review.get('inputSystems',{}).get(feature,{})
        if system.get('targetBindingHash')!=_digest(target) or not re.fullmatch('[a-f0-9]{64}',str(system.get('evidenceHash',''))) or not system.get('columnId'):
            raise RepairError('reviewed actual physical row ID required for every input')
        descriptor={**target,'systemColumnId':system['columnId'],'systemEvidenceHash':system['evidenceHash']}
        audit_elements.append(build_input_id_audit(spec,descriptor))
        if contract['workflow']=='research' and feature==arch.get('presentationInputFeature'):
            fragments.append(build_normalized_input_audit(contract,spec,descriptor,connection_id=scope['connectionId']))
        if feature=='salesRepTargets':
            if bindings.get('repTargets',{}).get('elementId')!=target['elementId']:raise RepairError('coverage must read the assigned Rep Targets input')
    if 'workflowOperations' in arch['inputs']:
        try:
            from .operation_outcomes import build_outcome_view
        except ImportError:
            from operation_outcomes import build_outcome_view
        spec,target,_=_approved_feature(manifest,receipt,'workflowOperations',require_assigned=True)
        audit_elements.insert(0,build_outcome_view({'elementId':target['elementId'],'name':spec['name']}))
    fragments.append(_page(contract['workflow']+'-input-ids',contract['workflow'].title()+' input IDs',audit_elements))
    if contract['workflow']!='research':fragments.append(build_workflow_repair_fragments(contract,bindings,connection_id=scope['connectionId']))
    navigation=build_navigation([*current['contents']['pages'],*[p for f in fragments for p in f['pages']]],prefix=contract['workflow']+'-navigation')
    fragments.append(_page(contract['workflow']+'-navigate',contract['workflow'].title()+' navigation',navigation))
    proposal=prepare_fragment_update(current,fragments,scope=scope,source_hash=manifest['sourceHash'],plan_hash=manifest['planHash'],review=review)
    merged=proposal['dryRunRequest']['contents'];by_id={e['id']:e for e in merged['elements']}
    for eid,mapping in review.get('readonlyFormats',{}).items():
        if eid not in elements or mapping.get('elementHash')!=_digest(elements[eid]):raise RepairError('read-only display target changed')
        technical_source=None
        if contract['workflow']=='expense':
            calc=bindings['calculation'];source=elements[calc['elementId']]
            if elements[eid].get('source')=={'kind':'table','elementId':calc['elementId']}:
                technical_source={'element':source,'elementHash':calc['elementHash'],'viewElementHash':mapping['elementHash']}
        by_id[eid]=propagate_readonly_presentation(contract,elements[eid],mapping['sourceCells'],technical_source=technical_source)
    if contract['workflow']=='expense':
        patch=review.get('currencyPatch',{});calc=bindings['calculation'];element=copy.deepcopy(by_id[calc['elementId']])
        if patch.get('decisionHash')!=_digest(contract['currencyDecision']) or patch.get('elementHash')!=calc['elementHash']:
            raise RepairError('explicit reviewed cents replacement required')
        matches=[c for c in element['columns'] if c['id']==patch.get('columnId') and c.get('name')=='Exception']
        if len(matches)!=1 or matches[0].get('formula')!=patch.get('beforeFormula'):raise RepairError('expense comparison target changed')
        olds=['Abs([Recon Variance]) > 0.01','Abs([ExpenseJoin/Card Amount] - [ExpenseJoin/Ledger Amount]) > 0.01']
        present=[v for v in olds if matches[0]['formula'].count(v)==1]
        if len(present)!=1:raise RepairError('exact raw cent comparison absent or ambiguous')
        old=present[0]
        new='Abs(Round([ExpenseJoin/Card Amount] * 100, 0) - Round([ExpenseJoin/Ledger Amount] * 100, 0)) > 1'
        if matches[0]['formula'].count(old)!=1:raise RepairError('only the exact reviewed raw cent comparison may be replaced')
        matches[0]['formula']=matches[0]['formula'].replace(old,new)
        by_id[calc['elementId']]=element
    for eid,technical in review.get('technicalVisibility',{}).items():
        if eid not in elements or technical.get('elementHash')!=_digest(elements[eid]):raise RepairError('technical-column target changed')
        element=copy.deepcopy(by_id[eid]);names={'Last Approval Operation ID','Applied Approval Operation IDs','StoredRowID'}
        ids=set(technical.get('columnIds',[]));matches=[c for c in element.get('columns',[]) if c['id'] in ids]
        if len(matches)!=len(ids) or any(c.get('name') not in names for c in matches):raise RepairError('only explicitly reviewed technical columns can be hidden')
        for column in matches:column['hidden']=True
        by_id[eid]=element
    merged['elements']=[by_id[e['id']] for e in merged['elements']]
    if review.get('inputLayout'):merged=shrink_input_layout(merged,review['inputLayout']);proposal['dryRunRequest']['contents']=merged
    if review.get('researchEditLayout'):
        patch=review['researchEditLayout'];feature=arch.get('presentationInputFeature')
        if contract['workflow']!='research' or patch.get('input_element_id')!=receipt['targetBindings'][feature]['elementId']:raise RepairError('research layout must target assigned primary input')
        merged=repair_research_edit_layout(merged,**patch);proposal['dryRunRequest']['contents']=merged
    if review.get('cleanup'):merged=cleanup_owned_scaffolds(merged,review['cleanup']);proposal['dryRunRequest']['contents']=merged
    proposal['sourceBindings']=copy.deepcopy(bindings)
    proposal['sourceGraph']=_source_graph_hashes(elements,[b['elementId'] for b in bindings.values()])
    proposal.pop('proposalHash');proposal['proposalHash']=_digest(proposal);return proposal


def finalize_repair_update(proposal,fresh,dry_run,*,consumed_proposals=(),reviewed_warning_hash=None,fresh_sdk_actions=None):
    if proposal.get('proposalHash')!=_digest({k:v for k,v in proposal.items() if k!='proposalHash'}):raise RepairError('repair proposal changed')
    if proposal['proposalHash'] in consumed_proposals:raise RepairError('consumed repair proposal refused')
    if fresh.get('workbookId')!=proposal['workbookId'] or fresh.get('documentVersion')!=proposal['documentVersion'] or _digest(fresh['contents'])!=proposal['beforeHash']:
        raise RepairError('fresh version/content conflict')
    if (dry_run.get('httpStatus')!=200 or dry_run.get('dryRun') is not True or not dry_run.get('evidenceId')
        or dry_run.get('errors') or dry_run.get('valid') is False or dry_run.get('proposalHash')!=proposal['proposalHash']):raise RepairError('actual hash-bound dry run required')
    if dry_run.get('warnings') and reviewed_warning_hash!=_digest(dry_run['warnings']):raise RepairError('dry-run warnings require review')
    try:
        from .action_inventory import verify_action_preservation
    except ImportError:
        from action_inventory import verify_action_preservation
    action_proof=verify_action_preservation(proposal['sdkActions'],fresh_sdk_actions,fresh)

    return {'method':'PUT','path':f"/v2/workbooks/{proposal['workbookId']}/contents",'body':{
        'documentVersion':fresh['documentVersion'],'contents':copy.deepcopy(proposal['dryRunRequest']['contents'])},
        'sdkActionPreflight':action_proof,'proposalHash':proposal['proposalHash'],'consumeBeforeSend':True,'productionStatus':'BLOCKED'}


def account_formula_targets(contract,targets,*,document_id,version,current_contents,audit_checks=None):
    """Account every source formula, including audit/static expressions, without caching.

    Each binding names the actual native column or explicit source audit check.
    Runtime evidence is separate; merely linking a formula cannot verify it.
    """
    validate_repair_contract(contract);ids={r['id'] for r in contract['formulaCells']}
    if current_contents.get('workbookId')!=document_id or current_contents.get('documentVersion')!=version:raise RepairError('fresh formula target contents required')
    elements={e['id']:e for e in current_contents['contents']['elements']};audit_checks=audit_checks or {}
    if set(targets)!=ids:raise RepairError('every source formula requires exactly one target/check binding')
    result=[]
    for record in contract['formulaCells']:
        binding=targets[record['id']]
        if (binding.get('documentId')!=document_id or binding.get('version')!=version or not binding.get('elementId')
            or not (binding.get('columnId') or binding.get('checkId')) or binding.get('kind') not in {'native-formula','source-audit','source-static-expression'}
            or not binding.get('sourceFormulaHash')==_digest(record['originalFormula'])):raise RepairError('stale/incomplete actual formula binding')
        element=elements.get(binding['elementId'],{})
        if binding.get('elementHash')!=_digest(element):raise RepairError('actual formula element hash mismatch')
        if binding['kind']=='source-audit':
            check=audit_checks.get(binding.get('checkId'),{})
            if check.get('sourceFormulaHash')!=binding['sourceFormulaHash'] or not check.get('reviewer') or not check.get('equation'):raise RepairError('explicit source audit equation review required')
        elif not any(c.get('id')==binding.get('columnId') and c.get('formula') for c in element.get('columns',[])):raise RepairError('actual native formula column absent')
        result.append({**record,'targetBinding':copy.deepcopy(binding),'verification':'UNRESOLVED'})
    return {'sourceHash':contract['source']['sha256'],'contractHash':contract['contractHash'],'formulas':result,'productionStatus':'BLOCKED'}


def validate_workflow_completion(manifest):
    try:
        from .acceptance import validate_production_matrix,build_workflow_reports
    except ImportError:
        from acceptance import validate_production_matrix,build_workflow_reports
    arch=manifest['planProjection'].get('workflowRepairArchitecture')
    if not arch:return True
    contract=arch['contract'];validate_repair_contract(contract)
    evidence=manifest.get('runtimeEvidence',{}).get('workflowAcceptance') or {}
    if not all(k in evidence for k in ('matrix','ledger','finalReview')):raise RepairError('mandatory actual workflow acceptance evidence missing')
    validate_production_matrix(evidence['matrix'])
    reports=build_workflow_reports(evidence['matrix'],evidence['ledger'],final_review=evidence['finalReview'])
    r=reports.get(contract['workflow'],{})
    if (r.get('verdict')!='VERIFIED_COMPLETE' or r.get('sourceHash')!=manifest['sourceHash'] or r.get('planHash')!=manifest['planHash']
        or {'documentId':r.get('documentId'),'version':r.get('version')}!=manifest.get('target')):raise RepairError('workflow acceptance incomplete or stale')
    return True


def build_secondary_input_hosts(contract):
    """New inputs receive dedicated pages; their frames cannot overlap old inputs."""
    validate_repair_contract(contract)
    return {feature:_page(feature+'-edit',spec['name'],[{'id':feature+'-guide','kind':'text',
        'body':'Editable source assumptions from '+spec['sourceRange']+'. Values drive live workflow results; preserve business keys.'}])
        for feature,spec in contract['additionalInputs'].items()}


def rebind_after_repair(manifest,receipt,proposal,readback,sdk_observations,review,*,sdk_actions=None):
    """Produce a fresh immutable input receipt after reviewed REST/SDK readback.

    It never copies a pre-write native system-column ID across a REST update.
    Unknown server normalization is rejected for independent review.
    """
    try:
        from .input_sync import build_artifact_records
        from .native_actions import _layout_element_owners
    except ImportError:
        from input_sync import build_artifact_records
        from native_actions import _layout_element_owners
    validate_manifest(manifest)
    try:
        from .action_inventory import verify_action_preservation
    except ImportError:
        from action_inventory import verify_action_preservation
    if proposal.get('nativeActionEdits'):
        verify_native_action_update_readback(proposal,readback,sdk_actions,review)
    else:verify_action_preservation(proposal['sdkActions'],sdk_actions,readback)
    if (proposal.get('proposalHash')!=_digest({k:v for k,v in proposal.items() if k!='proposalHash'})
        or proposal['sourceHash']!=manifest['sourceHash'] or proposal['planHash']!=manifest['planHash']):raise RepairError('repair proposal/source/plan changed')
    if (readback.get('workbookId')!=proposal['workbookId'] or not isinstance(readback.get('documentVersion'),int)
        or readback['documentVersion']<=proposal['documentVersion'] or readback.get('contents')!=proposal['dryRunRequest']['contents']):raise RepairError('actual repair readback differs; unknown normalization requires review')
    if (review.get('status')!='APPROVED_STAGING' or not review.get('owner') or review.get('contentsHash')!=_digest(readback['contents'])
        or review.get('sdkHash')!=_digest(sdk_observations)):raise RepairError('fresh exact contents and SDK review required')
    if sdk_observations.get('documentId')!=readback['workbookId'] or sdk_observations.get('version')!=readback['documentVersion'] or not sdk_observations.get('evidenceId'):raise RepairError('fresh version-bound SDK capture required')
    arch=manifest['planProjection']['sharedInputArchitecture'];bindings={};elements={e['id']:e for e in readback['contents']['elements']};root=_layout_root(readback['contents']['layout'])
    for feature,spec in arch['inputs'].items():
        _,old,_=_approved_feature(manifest,receipt,feature,require_assigned=True)
        element=elements.get(old['elementId'],{});observed=sdk_observations.get('tables',{}).get(feature,{})
        if element.get('kind')!='input-table' or element.get('name')!=spec['name'] or observed.get('elementId')!=element.get('id'):
            raise RepairError('assigned input identity changed')
        owners=_layout_element_owners(root,element['id'])
        if len(owners)!=1:raise RepairError('actual input page ambiguous')
        logical={c['name'] for c in spec['columns']};columns=element.get('columns',[])
        mapping={name:[c['id'] for c in columns if c.get('name')==name] for name in logical}
        if any(len(ids)!=1 for ids in mapping.values()):raise RepairError('actual input columns ambiguous/missing')
        descriptors=observed.get('columns',[])
        for name in logical:
            matches=[c for c in descriptors if c.get('id')==mapping[name][0] and c.get('name')==name]
            expected=next(c for c in spec['columns'] if c['name']==name)
            if len(matches)!=1 or matches[0].get('type')!=expected['type'] or matches[0].get('columnClass') not in {'LiteralEditableColumn','SingleSelectEditableColumn','MultiSelectEditableColumn'}:
                raise RepairError('fresh actual editable SDK schema mismatch')
        systems=[c for c in descriptors if c.get('name') not in logical]
        system=review.get('systemColumns',{}).get(feature,{})
        if len(systems)!=1 or system.get('observationHash')!=_digest(systems[0]) or not re.fullmatch('[a-f0-9]{64}',str(system.get('rowIdEvidenceHash',''))):
            raise RepairError('fresh actual system-row identity evidence required')
        binding={k:copy.deepcopy(v) for k,v in old.items() if k not in {'systemColumn','rowIdColumnId','columnContracts'}}
        binding.update(documentId=readback['workbookId'],version=readback['documentVersion'],pageId=owners[0].get('id'),columnIds={k:v[0] for k,v in mapping.items()})
        # Explicit reviewed SDK contracts may authorize a normalized FormulaColumn
        # alias or real native InputTableIDColumn. Never infer identity from name.
        if systems[0].get('columnClass')=='InputTableIDColumn':
            if systems[0].get('systemKind')!='row_id':raise RepairError('native system kind mismatch')
            binding['rowIdColumnId']=systems[0]['id']
        elif system.get('approvedAlias'):
            alias=system['approvedAlias']
            if (any(alias.get(k)!=systems[0].get(k) for k in ('id','name','columnClass','type','systemKind','formula','elementId'))
                or alias.get('rowIdBinding',{}).get('evidenceHash')!=system['rowIdEvidenceHash']):raise RepairError('approved row-ID alias differs from fresh SDK observation/evidence')
            try:
                from .input_display import _system_column_contract
            except ImportError:
                from input_display import _system_column_contract
            binding['systemColumn']=copy.deepcopy(system['approvedAlias']);_system_column_contract(binding)
        else:raise RepairError('normalized row-ID alias requires complete reviewed contract')
        if old.get('columnContracts'):
            approved=review.get('columnContracts',{}).get(feature)
            if not approved:raise RepairError('source-list conversions need fresh reviewed option/class contracts')
            binding['columnContracts']=copy.deepcopy(approved)
        bindings[feature]=binding
    updated=copy.deepcopy(manifest);updated['target']={'documentId':readback['workbookId'],'version':readback['documentVersion']}
    new_receipt={'status':'APPROVED_STAGING','sourceHash':manifest['sourceHash'],'planHash':manifest['planHash'],'targetBindings':bindings,
        'artifacts':build_artifact_records(updated,arch,bindings),'repairProposalHash':proposal['proposalHash'],'review':copy.deepcopy(review)}
    return {'manifest':updated,'receipt':new_receipt,'status':'STAGED_INPUT_BINDINGS_REVIEWED','productionStatus':'BLOCKED'}


def prepare_native_action_update(manifest,receipt,action_bindings,fragments,current,*,scope,review,sdk_actions,assigned_components):
    """Edit only compiler-produced native actions, never general element content.

    Exact restoration proof is required even for an action omitted by public GET.
    Finalize with the existing fresh/dry-run/one-use protocol, then query rows and
    capture complete SDK semantics using the declared allowed semantic changes.
    """
    try:
        from .native_actions import emit_public_native_action_fragments
        from .action_inventory import digest as action_digest
    except ImportError:
        from native_actions import emit_public_native_action_fragments
        from action_inventory import digest as action_digest
    rebuilt=emit_public_native_action_fragments(manifest,receipt,action_bindings,current_contents=current,
        sdk_actions=sdk_actions,assigned_components=assigned_components,staging_review=fragments.get('constructionReview'))
    if fragments!=rebuilt:raise RepairError('public action fragments differ from exact source compiler output')
    required={'sourceHash':manifest['sourceHash'],'planHash':manifest['planHash'],'bindingHash':action_bindings['bindingHash'],
        'fragmentHash':fragments['fragmentHash'],'assignedComponentsHash':fragments['assignedComponentsHash'],
        'sdkInventoryHash':action_digest(sdk_actions),'allowedActionEdits':fragments['allowedEditPaths']}
    if review.get('status')!='APPROVED_STAGING' or not str(review.get('owner') or '').strip() or re.match(r'(?i)^(pending|todo|tbd|unknown)\b',str(review.get('owner'))) or any(review.get(k)!=v for k,v in required.items()):
        raise RepairError('exact fresh named native action construction review required')
    semantic_changes=review.get('sdkAllowedChanges') or {}
    expected_hosts={e['hostElementId'] for e in fragments['actionEdits']}
    if set(semantic_changes)!=expected_hosts or any(set(v)!={'sequences'} or not isinstance(v['sequences'],list) for v in semantic_changes.values()):raise RepairError('complete exact declared SDK semantic changes required')
    for host in expected_hosts:
        old=sdk_actions['elements'][host]['sequences'];edits=[e for e in fragments['actionEdits'] if e['hostElementId']==host]
        edited_ids={e['sequenceId'] for e in edits};old_ids={s['id'] for s in old};expected=semantic_changes[host]['sequences']
        if len(expected)!=len(old)+len(edited_ids-old_ids):raise RepairError('SDK delta must account exactly generated sequences')
        # Preserve all unrelated sequences, including ones omitted by REST, in
        # their original order even when their host receives a public edit.
        cursor=0
        for seq in old:
            if seq['id'] in edited_ids:continue
            semantic={k:v for k,v in seq.items() if k!='id'}
            matches=[i for i in range(cursor,len(expected)) if expected[i]==semantic]
            if not matches:raise RepairError('unrelated SDK sequence cannot change/disappear on edited host')
            cursor=matches[0]+1
    evidence=review.get('readbackEvidenceHashes') or {}
    if set(evidence)!={'fullRows','schemaDisplay','sourceGraph'} or any(not re.fullmatch('[a-f0-9]{64}',str(h)) for h in evidence.values()):
        raise RepairError('fresh complete rows/schema/display/source evidence hashes required')
    proposal=prepare_fragment_update(current,[],scope=scope,source_hash=manifest['sourceHash'],plan_hash=manifest['planHash'],
        review={**review,'sdkActionInventory':sdk_actions})
    merged=proposal['dryRunRequest']['contents'];elements={e['id']:e for e in merged['elements']}
    seen=set()
    for edit in fragments['actionEdits']:
        key=(edit['hostElementId'],edit['sequenceId'])
        if key in seen:raise RepairError('duplicate generated action edit')
        seen.add(key);host=elements[edit['hostElementId']];actions=host.setdefault('actions',[])
        matches=[i for i,a in enumerate(actions) if a['id']==edit['sequenceId']]
        if edit['beforeActionHash'] is None:
            if matches:raise RepairError('new action already exists')
            actions.append(copy.deepcopy(edit['afterAction']))
        elif len(matches)==1 and _digest(actions[matches[0]])==edit['beforeActionHash']:
            actions[matches[0]]=copy.deepcopy(edit['afterAction'])
        else:raise RepairError('existing action semantic pointer stale')
    # Assert exact data/layout preservation independently of fragment generation.
    before={e['id']:e for e in current['contents']['elements']}
    for eid,element in elements.items():
        if {k:v for k,v in element.items() if k!='actions'}!={k:v for k,v in before[eid].items() if k!='actions'}:
            raise RepairError('native action adapter attempted a non-action edit')
    proposal['nativeActionEdits']=copy.deepcopy(fragments['allowedEditPaths'])
    proposal['publicSelectionAuthorities']=copy.deepcopy(fragments['selectionAuthorities'])
    proposal['nativeActionFragmentHash']=fragments['fragmentHash']
    proposal['nativeActionFragments']=copy.deepcopy(fragments)
    proposal['sdkAllowedChangesReview']=copy.deepcopy(review.get('sdkAllowedChanges',{}))
    proposal['limitations']=['Initial construction is disabled; enabling requires separate exact disabled-proof and configuration reviews.',
        'Trigger CountIf and dictionary/single-row update are serialized staging safeguards with a TOCTOU race; no transaction or affected-row count.',
        'Public GET can omit native Prepare; exact reviewed missing-only restoration and actual SDK post-readback remain mandatory.']
    proposal.pop('proposalHash');proposal['proposalHash']=_digest(proposal)
    return proposal


def verify_native_action_update_readback(proposal,readback,sdk_actions,review):
    """Actual post-REST proof for precisely declared native action deltas only."""
    try:
        from .action_inventory import verify_action_preservation,digest as action_digest,qualify_public_selection_construction
    except ImportError:
        from action_inventory import verify_action_preservation,digest as action_digest,qualify_public_selection_construction
    if proposal.get('proposalHash')!=_digest({k:v for k,v in proposal.items() if k!='proposalHash'}):raise RepairError('native action proposal hash stale')
    if (readback.get('workbookId')!=proposal['workbookId'] or not isinstance(readback.get('documentVersion'),int)
            or readback['documentVersion']<=proposal['documentVersion'] or readback.get('contents')!=proposal['dryRunRequest']['contents']):raise RepairError('exact fresh native action REST readback required')
    expected_hosts={e['hostElementId'] for e in proposal.get('nativeActionEdits',[])}
    changes=proposal.get('sdkAllowedChangesReview') or {}
    if not expected_hosts or set(changes)!=expected_hosts:raise RepairError('complete exact declared SDK semantic changes required')
    if (review.get('status')!='APPROVED_STAGING' or not str(review.get('owner') or '').strip() or re.match(r'(?i)^(pending|todo|tbd|unknown)\b',str(review.get('owner')))
            or review.get('proposalHash')!=proposal['proposalHash'] or review.get('contentsHash')!=_digest(readback['contents'])
            or review.get('actionInventoryHash')!=action_digest(sdk_actions) or review.get('allowedSemanticChangesHash')!=action_digest(changes)):
        raise RepairError('named exact actual SDK semantic readback review required')
    fragments=proposal['nativeActionFragments']
    # Construction authority bridges only the exact known literal/typed binding;
    # it deliberately cannot authorize business activation or claim cardinality.
    construction_review={'status':'APPROVED_STAGING','owner':review['owner'],'documentId':readback['workbookId'],
        'version':readback['documentVersion'],'contentsHash':_digest(readback['contents']),'captureHash':action_digest(sdk_actions),
        **{k:fragments[k] for k in ('fragmentHash','sourceHash','planHash','bindingHash')}}
    canonical=qualify_public_selection_construction(sdk_actions,readback,fragments,construction_review)
    result=verify_action_preservation(proposal['sdkActions'],canonical,readback,allowed_changes=changes)
    result.update(status='VERIFIED_NATIVE_ACTION_CONSTRUCTION_RUNTIME_UNRESOLVED',productionStatus='BLOCKED',
        activation=fragments.get('activation','disabled'))
    return result
