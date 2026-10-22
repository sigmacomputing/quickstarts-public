"""Reviewed, source-bound ordinary input tables; no target calls or writes."""
from __future__ import annotations

import copy
import datetime as dt
import hashlib
import json
import re
from pathlib import Path

from openpyxl.utils.cell import get_column_letter, range_boundaries

try:
    from .input_sync import plan_seed
    from .source_workbook import _bounds, build_source_model
except ImportError:
    from input_sync import plan_seed
    from source_workbook import _bounds, build_source_model


PLAN_VERSION = 2
LEGACY_PLAN_VERSION = 1
_TYPES = {"text", "number", "datetime", "boolean"}
_SAFE = re.compile(r"^[A-Za-z][A-Za-z0-9_-]*$")
_CELL = re.compile(r"^[A-Z]{1,3}[1-9][0-9]*$")


def _digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":"),
                                     ensure_ascii=False).encode("utf-8")).hexdigest()


def _value(cell):
    if cell.get("dataType") == "e":
        raise ValueError("Excel error cell cannot be an entered value")
    if cell.get("dataType") not in {"n", "s", "b", "d", "inlineStr"}:
        raise ValueError(f"unsupported entered source cell dataType: {cell.get('dataType')}")
    value = cell["value"]
    if isinstance(value, dict):
        if value.get("type") in {"date", "datetime"}:
            return value["iso"]
        raise ValueError(f"unsupported source cell type: {value.get('type')}")
    if value is None or type(value) in {str, int, float, bool}:
        return value
    raise ValueError("unsupported source cell value")


def _matches(value, kind):
    if value is None:
        return True
    if kind == "text":
        return type(value) is str
    if kind == "number":
        return type(value) in {int, float}
    if kind == "boolean":
        return type(value) is bool
    if kind == "datetime":
        if not isinstance(value, str):
            return False
        try:
            if re.fullmatch(r"\d{4}-\d\d-\d\d", value):
                dt.date.fromisoformat(value)
            elif re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d(?:\.\d+)?(?:[+-]\d\d:\d\d|Z)?", value):
                dt.datetime.fromisoformat(value.replace("Z", "+00:00"))
            else:
                return False
        except ValueError:
            return False
        return True
    return False


def _selection(model, decision):
    sheet = decision.get("sheet")
    if not isinstance(sheet, str) or sheet not in {item["name"] for item in model["sheets"] if item["kind"] == "worksheet"}:
        raise ValueError("unknown source sheet")
    table_name, block = decision.get("table"), decision.get("range")
    if bool(table_name) == bool(block):
        raise ValueError("select exactly one formal table or bounded cell block")
    if table_name:
        matches = [item for item in model["tables"] if item["sheet"] == sheet and item["name"] == table_name]
        if len(matches) != 1 or not matches[0]["bodyRange"]:
            raise ValueError("unknown or empty source table")
        item = matches[0]
        left, top, right, bottom = range_boundaries(item["bodyRange"])
        headers = item["headers"]
        source_range = item["range"]
    else:
        try:
            left, top, right, bottom = range_boundaries(block)
        except (TypeError, ValueError) as exc:
            raise ValueError("invalid bounded cell block") from exc
        if None in (left, top, right, bottom) or left > right or top > bottom:
            raise ValueError("invalid bounded cell block")
        headers = decision.get("headers")
        if not isinstance(headers, list) or len(headers) != right - left + 1:
            raise ValueError("cell block requires explicit headers")
        source_range = block
    if any(not isinstance(name, str) or not name for name in headers) or len(set(headers)) != len(headers):
        raise ValueError("source headers must be unique and nonempty")
    return sheet, headers, (left, top, right, bottom), source_range


def _build_replacement(model, decision, connection_id):
    feature = decision.get("featureId")
    if not isinstance(feature, str) or not _SAFE.fullmatch(feature):
        raise ValueError("replacement featureId must be stable and safe")
    review = decision.get("replacementReview") or {}
    if (decision.get("reviewed") is not True or review.get("approved") is not True
            or not str(review.get("basis") or "").strip()
            or not str(decision.get("replacementFor") or "").strip()
            or not str(decision.get("displayName") or "").strip()):
        raise ValueError("replacement requires explicit reviewed source and target label")
    keys, literal, types = (decision.get(name) for name in ("keyColumns", "literalColumns", "columnTypes"))
    if (not isinstance(keys, list) or len(keys) != 1 or not isinstance(literal, list)
            or len(literal) != 2 or keys[0] not in literal or not isinstance(types, dict)
            or set(types) != set(literal) or any(kind not in _TYPES for kind in types.values())):
        raise ValueError("replacement requires one typed natural key and one typed value")
    value_name = next(name for name in literal if name != keys[0])
    rows, bindings = decision.get("suppliedRows"), decision.get("replacementBindings")
    if (not isinstance(rows, list) or not rows or not isinstance(bindings, list)
            or len(bindings) != len(rows)):
        raise ValueError("replacement suppliedRows and source bindings must match")
    cells = {f"{cell['sheet']}!{cell['address']}": cell for cell in model["cells"]}
    provenance, obligations, seen_keys, seen_formulas = {}, [], set(), set()
    for index, (row, binding) in enumerate(zip(rows, bindings)):
        if (not isinstance(row, dict) or set(row) != set(literal)
                or not isinstance(binding, dict) or set(binding.get("key") or {}) != set(keys)
                or binding["key"] != {key: row[key] for key in keys}
                or set(binding.get("sourceCells") or {}) != set(literal)):
            raise ValueError("replacement row and source bindings are inconsistent")
        for name in literal:
            if row[name] is None or not _matches(row[name], types[name]):
                raise ValueError("replacement row disagrees with reviewed type")
        key_token = _digest([[type(row[key]).__name__, row[key]] for key in keys])
        if key_token in seen_keys:
            raise ValueError("duplicate replacement business key")
        seen_keys.add(key_token)
        key_cell_id = binding["sourceCells"][keys[0]]
        value_cell_id = binding["sourceCells"][value_name]
        key_cell, value_cell = cells.get(key_cell_id), cells.get(value_cell_id)
        if (not key_cell or key_cell["kind"] != "literal"
                or _value(key_cell) != binding.get("sourceLabel")):
            raise ValueError("replacement key source label is stale")
        if (not value_cell or value_cell["kind"] != "formula"
                or value_cell.get("expression") != binding.get("originalFormula")
                or value_cell_id in seen_formulas):
            raise ValueError("replacement source formula is stale or duplicated")
        seen_formulas.add(value_cell_id)
        provenance[str(index)] = dict(binding["sourceCells"])
        obligations.append({"featureId": feature, "sourceCellId": value_cell_id,
                            "sourceRow": int(re.search(r"\d+$", value_cell["address"]).group()),
                            "sourceColumn": value_name, "expression": binding["originalFormula"],
                            "dependencies": value_cell.get("dependencies") or [],
                            "mappingStatus": "reviewed-replacement-input", "replacementRowKey": binding["key"]})
    spec = {"kind": "input-table", "name": decision["displayName"],
            "keyColumns": keys, "connectionId": connection_id, "editMode": "published-explore",
            "sourceBinding": decision["replacementFor"], "rows": copy.deepcopy(rows),
            "columns": [{"name": name, "type": types[name], "provenance": "reviewed-replacement"}
                        for name in literal], "sourceCells": provenance,
            "sourceRowNumbers": [item["sourceRow"] for item in obligations],
            "sourceHeaders": literal, "derivedColumns": [], "overrideCells": {},
            "replacementBindings": copy.deepcopy(bindings), "replacementReview": copy.deepcopy(review)}
    return feature, spec, obligations, {}


def _build_reshape_feature(source_path, model, decision, connection_id):
    """Recompute a reviewed unpivot from exact source bytes before admitting it as input."""
    if (not isinstance(decision, dict) or decision.get("reviewed") is not True
            or decision.get("role") != "input"):
        raise ValueError("reshape input requires an explicit reviewed input role")
    feature = decision.get("featureId")
    if not isinstance(feature, str) or not _SAFE.fullmatch(feature):
        raise ValueError("reshape featureId must be stable and safe")
    reshape = decision.get("reshape")
    expected = {"kind", "sheet", "table", "keyColumns", "valueColumns", "axisName", "valueName"}
    if (not isinstance(reshape, dict) or reshape.get("kind") != "unpivot"
            or set(reshape) not in (expected, expected | {"axisValues"})
            or decision.get("overrideCells") or decision.get("emptySchema")):
        raise ValueError("reshape must explicitly select one reviewed unpivot")
    try:
        from .source_reshape import unpivot_source_table
    except ImportError:
        from source_reshape import unpivot_source_table
    result = unpivot_source_table(
        source_path, sheet=reshape["sheet"], table=reshape["table"],
        key_columns=reshape["keyColumns"], value_columns=reshape["valueColumns"],
        axis_name=reshape["axisName"], value_name=reshape["valueName"],
        axis_values=reshape.get("axisValues"))
    columns = result["keyColumns"] + [reshape["valueName"]]
    if (result["sourceSha256"] != model["sourceSha256"]
            or decision.get("keyColumns") != result["keyColumns"]
            or decision.get("literalColumns") != columns
            or decision.get("columnTypes") != result["columnTypes"]):
        raise ValueError("reviewed reshape keys, columns or types disagree with source")
    spec = {"kind": "input-table", "name": decision.get("displayName") or f"{reshape['table']} Long",
            "keyColumns": result["keyColumns"], "connectionId": connection_id,
            "editMode": "published-explore", "sourceBinding": result["sourceBinding"],
            "rows": result["rows"],
            "columns": [{"name": name, "type": result["columnTypes"][name], "provenance": "source-reshape"}
                        for name in columns],
            "sourceCells": result["sourceCells"],
            "sourceCellFormats": result["sourceCellFormats"],
            "sourceCellStates": result["sourceCellStates"],
            "sourceRowNumbers": result["sourceRowNumbers"],
            "sourceHeaders": reshape["keyColumns"] + reshape["valueColumns"],
            "axisMapping": result["axisMapping"], "reshape": copy.deepcopy(reshape),
            "derivedColumns": [], "overrideCells": {}}
    return feature, spec, [], {}


def _build_record_feature(model, decision, connection_id):
    """Bind reviewed business records whose typed fields occupy arbitrary cells."""
    allowed = {"featureId", "displayName", "role", "reviewed", "keyColumns",
               "literalColumns", "columnTypes", "recordBindings", "recordIdentity",
               "immutableReview", "reviewedBlankColumns"}
    if set(decision) - allowed:
        raise ValueError("source records cannot contain supplied values or formula cache fields")
    if decision.get("reviewed") is not True or decision.get("role") not in {"input", "reference"}:
        raise ValueError("source records require a reviewed input or reference role")
    immutable_review = decision.get("immutableReview")
    if decision["role"] == "reference" and (not isinstance(immutable_review, dict)
            or immutable_review.get("reviewed") is not True
            or immutable_review.get("sourceSha256") != model["sourceSha256"]
            or not str(immutable_review.get("basis") or "").strip()):
        raise ValueError("immutable source reference requires exact source hash and review basis")
    feature = decision.get("featureId")
    if not isinstance(feature, str) or not _SAFE.fullmatch(feature):
        raise ValueError("featureId must be a stable safe identifier")
    keys, literal, types = (decision.get(name) for name in ("keyColumns", "literalColumns", "columnTypes"))
    if (not isinstance(keys, list) or not keys or len(set(keys)) != len(keys)
            or not isinstance(literal, list) or not literal or len(set(literal)) != len(literal)
            or not set(keys) <= set(literal) or not isinstance(types, dict)
            or set(types) != set(literal) or any(kind not in _TYPES for kind in types.values())):
        raise ValueError("source record columns, keys and explicit types must agree")
    bindings = decision.get("recordBindings")
    if not isinstance(bindings, list) or not bindings:
        raise ValueError("source record bindings must be nonempty")
    identity = decision.get("recordIdentity")
    identity_column = None
    if identity is not None:
        if (not isinstance(identity, dict) or identity.get("kind") != "frozen-source-record"
                or identity.get("reviewed") is not True
                or identity.get("sourceSha256") != model["sourceSha256"]
                or not str(identity.get("reviewBasis") or "").strip()
                or identity.get("column") not in literal
                or types[identity["column"]] != "text"
                or keys != [identity["column"]]):
            raise ValueError("frozen source record identity requires exact source hash and review basis")
        identity_column = identity["column"]
    if decision.get("overrideCells") or decision.get("reshape") or decision.get("emptySchema"):
        raise ValueError("source record bindings cannot mix reshape or override selections")
    worksheet_names = {sheet["name"] for sheet in model["sheets"] if sheet["kind"] == "worksheet"}
    cells = {f"{cell['sheet']}!{cell['address']}": cell for cell in model["cells"]}
    expected_fields = set(literal) - ({identity_column} if identity_column else set())
    rows, provenance, formats, states, row_numbers, seen_cells, seen_keys = [], {}, {}, {}, [], set(), set()
    for index, binding in enumerate(bindings):
        if isinstance(binding, dict) and set(binding) - {"sourceCells", "sourceRecordId"}:
            raise ValueError("source record binding cannot contain supplied values or formula cache fields")
        fields = binding.get("sourceCells") if isinstance(binding, dict) else None
        if not isinstance(fields, dict) or set(fields) != expected_fields:
            raise ValueError("source record requires one source cell per typed field")
        row, source_cells, source_formats, source_states = {}, {}, {}, {}
        for name in literal:
            if name == identity_column:
                continue
            cell_id = fields[name]
            if not isinstance(cell_id, str) or "!" not in cell_id:
                raise ValueError("source record cell ID must name a worksheet cell")
            sheet, address = cell_id.rsplit("!", 1)
            if sheet not in worksheet_names or not _CELL.fullmatch(address):
                raise ValueError("source record cell ID must name a worksheet cell")
            _bounds(address)
            if cell_id in seen_cells:
                raise ValueError("duplicate source cell binding")
            seen_cells.add(cell_id)
            cell = cells.get(cell_id)
            if cell and cell["kind"] == "formula":
                raise ValueError(f"formula cannot be seeded as literal entered data: {cell_id}")
            value = _value(cell) if cell else None
            if not _matches(value, types[name]):
                raise ValueError(f"source record type disagrees with reviewed type: {cell_id}")
            row[name] = value
            source_cells[name] = cell_id
            source_formats[name] = cell.get("numberFormat") if cell else None
            source_states[name] = ("blank" if cell is None else
                                   "empty-text" if value == "" else "literal")
        if identity_column:
            source_rows = {(cell_id.rsplit("!", 1)[0], int(re.search(r"\d+$", cell_id).group()))
                           for cell_id in source_cells.values()}
            if len(source_rows) != 1:
                raise ValueError("frozen source record identity requires one physical source row")
            source_sheet, source_row = next(iter(source_rows))
            expected_id = f"{source_sheet}!row:{source_row}"
            if binding.get("sourceRecordId") != expected_id:
                raise ValueError("frozen source record identity disagrees with physical row")
            row[identity_column] = expected_id
        elif "sourceRecordId" in binding:
            raise ValueError("source record identity lacks reviewed identity contract")
        for name in keys:
            if row[name] is None or row[name] == "":
                raise ValueError(f"missing typed source record key {name!r}")
        key_token = _digest([[type(row[name]).__name__, row[name]] for name in keys])
        if key_token in seen_keys:
            raise ValueError("duplicate typed business key")
        seen_keys.add(key_token)
        rows.append(row)
        provenance[str(index)] = source_cells
        formats[str(index)] = source_formats
        states[str(index)] = source_states
        first_cell = next(iter(source_cells.values()))
        row_numbers.append(int(re.search(r"\d+$", first_cell).group()))
    for name in expected_fields:
        if all(row[name] is None for row in rows) and name not in (decision.get("reviewedBlankColumns") or []):
            raise ValueError(f"all-blank source record field {name!r} requires reviewed type")
    spec = {"kind": "input-table", "name": decision.get("displayName") or feature,
            "keyColumns": keys, "connectionId": connection_id, "editMode": "published-explore",
            "sourceBinding": "reviewed-source-records", "rows": rows,
            "columns": [{"name": name, "type": types[name],
                         "provenance": "frozen-source-record" if name == identity_column else "source"}
                        for name in literal],
            "sourceCells": provenance, "sourceCellFormats": formats,
            "sourceCellStates": states, "sourceRowNumbers": row_numbers,
            "sourceHeaders": literal, "derivedColumns": [], "overrideCells": {},
            "recordIdentity": copy.deepcopy(identity) if identity else None,
            "immutableReview": copy.deepcopy(immutable_review) if immutable_review else None}
    return feature, spec, [], {}


def _build_feature(model, decision, connection_id, source_path):
    if isinstance(decision, dict) and decision.get("role") == "replacement":
        return _build_replacement(model, decision, connection_id)
    if isinstance(decision, dict) and "recordBindings" in decision:
        return _build_record_feature(model, decision, connection_id)
    if isinstance(decision, dict) and "reshape" in decision:
        return _build_reshape_feature(source_path, model, decision, connection_id)
    if not isinstance(decision, dict) or decision.get("reviewed") is not True or decision.get("role") not in {"input", "reference"}:
        raise ValueError("feature requires a reviewed input or maintained reference role")
    feature = decision.get("featureId")
    if not isinstance(feature, str) or not _SAFE.fullmatch(feature):
        raise ValueError("featureId must be a stable safe identifier")
    sheet, headers, bounds, source_range = _selection(model, decision)
    keys = decision.get("keyColumns")
    literal = decision.get("literalColumns")
    types = decision.get("columnTypes")
    if (not isinstance(keys, list) or not keys or len(keys) != len(set(keys))
            or not isinstance(literal, list) or not literal or len(literal) != len(set(literal))
            or not set(keys) <= set(literal) or not set(literal) <= set(headers)
            or not isinstance(types, dict) or set(types) != set(literal)
            or any(kind not in _TYPES for kind in types.values())):
        raise ValueError("input columns, typed keys, and explicit types must match source schema")
    overrides = decision.get("overrideCells") or {}
    if not isinstance(overrides, dict) or any(name not in literal or name in keys or not isinstance(cells, list)
                                                for name, cells in overrides.items()):
        raise ValueError("editable override mappings must name non-key literal columns")
    override_set = {name: set(cells) for name, cells in overrides.items()}
    if any(len(cells) != len(override_set[name]) for name, cells in overrides.items()):
        raise ValueError("duplicate editable override mapping")
    left, top, right, bottom = bounds
    cells = {(item["sheet"], item["address"]): item for item in model["cells"]}
    if decision.get("emptySchema") is True:
        if (set(decision.get("reviewedBlankColumns") or []) != set(literal)
                or not str(decision.get("typeReviewEvidence") or "").strip()):
            raise ValueError("empty schema requires reviewed types and typeReviewEvidence")
        if any((sheet, f"{get_column_letter(col)}{row}") in cells
               for row in range(top, bottom + 1) for col in range(left, right + 1)):
            raise ValueError("emptySchema cannot suppress populated source cells")
        return feature, {"kind": "input-table", "name": decision.get("displayName") or decision.get("table") or feature,
                         "keyColumns": keys, "connectionId": connection_id,
                         "editMode": "published-explore", "sourceBinding": f"{sheet}!{source_range}",
                         "rows": [], "columns": [{"name": name, "type": types[name], "provenance": "source"}
                                                 for name in literal],
                         "sourceCells": {}, "sourceCellFormats": {}, "sourceCellStates": {},
                         "sourceRowNumbers": [], "sourceHeaders": headers,
                         "derivedColumns": [], "overrideCells": {}, "emptySchema": True,
                         "typeReviewEvidence": decision["typeReviewEvidence"]}, [], {}
    rows, source_cells, source_formats, source_states, obligations, key_tokens = [], {}, {}, {}, [], set()
    override_records = {name: [] for name in overrides}
    for row_number in range(top, bottom + 1):
        row, provenance, formats, states, row_overrides = {}, {}, {}, {}, {}
        for offset, name in enumerate(headers):
            address = f"{get_column_letter(left + offset)}{row_number}"
            cell = cells.get((sheet, address))
            cell_id = f"{sheet}!{address}"
            if cell and cell["kind"] == "formula":
                obligations.append({"featureId": feature, "sourceCellId": cell_id,
                                    "sourceRow": row_number, "sourceColumn": name,
                                    "expression": cell.get("expression"),
                                    "dependencies": cell.get("dependencies") or [],
                                    "mappingStatus": "live-calculation-required"})
                if name in literal and name not in override_set:
                    raise ValueError(f"formula cannot be seeded as literal entered data: {cell_id}")
                if name in override_set and address in override_set[name]:
                    raise ValueError(f"formula cell cannot be an editable override: {cell_id}")
                continue
            if name not in literal:
                if cell and cell["kind"] == "literal":
                    raise ValueError(f"unmapped literal source cell: {cell_id}")
                continue
            if name in override_set and cell is None:
                if address in override_set[name]:
                    raise ValueError(f"editable override has no literal source cell: {cell_id}")
                continue
            if cell is None:
                value = None
            else:
                value = _value(cell)
            if not _matches(value, types[name]):
                raise ValueError(f"source cell type disagrees with reviewed type: {cell_id}")
            if name in override_set and address not in override_set[name]:
                raise ValueError(f"literal override lacks explicit editable mapping: {cell_id}")
            if name in override_set:
                row_overrides[name] = (value, cell_id)
                continue
            row[name] = value
            provenance[name] = cell_id
            formats[name] = cell.get("numberFormat") if cell else None
            states[name] = ("blank" if cell is None else
                            "empty-text" if value == "" else "literal")
        key = []
        for name in keys:
            if name not in row or row[name] is None or row[name] == "":
                raise ValueError(f"missing typed key {name!r} in {sheet} row {row_number}")
            key.append([type(row[name]).__name__, row[name]])
        token = _digest(key)
        if token in key_tokens:
            raise ValueError("duplicate typed business key")
        key_tokens.add(token)
        rows.append(row)
        source_cells[str(len(rows) - 1)] = provenance
        source_formats[str(len(rows) - 1)] = formats
        source_states[str(len(rows) - 1)] = states
        for name, (value, cell_id) in row_overrides.items():
            override_records[name].append({
                "row": {**{key_name: row[key_name] for key_name in keys}, name: value},
                "sourceCells": {**{key_name: provenance[key_name] for key_name in keys}, name: cell_id},
                "sourceCellFormats": {**{key_name: formats[key_name] for key_name in keys},
                                      name: cells[(sheet, cell_id.rsplit('!', 1)[1])].get("numberFormat")},
                "sourceCellStates": {**{key_name: states[key_name] for key_name in keys},
                                     name: "empty-text" if value == "" else "literal"},
                "sourceRow": row_number})
    for name, addresses in override_set.items():
        seen = {item["sourceCells"][name].split("!")[1] for item in override_records[name]}
        if addresses != seen:
            raise ValueError("editable override mappings are incomplete or foreign")
        if not any(item["sourceColumn"] == name for item in obligations):
            raise ValueError(f"editable override column has no source formula: {name}")
    main_literal = [name for name in literal if name not in overrides]
    for name in main_literal:
        if all(row.get(name) is None for row in rows) and name not in (decision.get("reviewedBlankColumns") or []):
            raise ValueError(f"all-blank column {name!r} requires explicit reviewed type")
    if not rows:
        raise ValueError("source input selection has no rows")
    display = decision.get("displayName") or decision.get("table") or feature
    override_features = {name: f"{feature}__override_{_digest(name)[:8]}" for name in sorted(overrides)}
    derived = [name for name in headers if name not in main_literal]
    spec = {"kind": "input-table", "name": display,
            "keyColumns": keys, "connectionId": connection_id, "editMode": "published-explore",
            "sourceBinding": f"{sheet}!{source_range}", "rows": rows,
            "columns": [{"name": name, "type": types[name], "provenance": "source"} for name in main_literal],
            "sourceCells": source_cells, "sourceCellFormats": source_formats,
            "sourceCellStates": source_states,
            "sourceRowNumbers": list(range(top, bottom + 1)),
            "sourceHeaders": headers, "derivedColumns": derived,
            "derivedMappings": {name: {"sourceFormulaCells": [item["sourceCellId"] for item in obligations
                                                                  if item["sourceColumn"] == name],
                                       "overrideInputFeature": override_features.get(name)}
                                for name in derived},
            "overrideCells": overrides, "overrideInputFeatures": override_features}
    extra_inputs = {}
    for name, override_feature in override_features.items():
        records = override_records[name]
        if not records or all(item["row"][name] is None for item in records):
            raise ValueError(f"editable override column has no entered values: {name}")
        extra_inputs[override_feature] = {
            "kind": "input-table", "name": f"{display} — {name} Overrides",
            "keyColumns": keys, "connectionId": connection_id, "editMode": "published-explore",
            "sourceBinding": f"{sheet}!{source_range}",
            "rows": [item["row"] for item in records],
            "columns": [{"name": column, "type": types[column], "provenance": "source"}
                        for column in keys + [name]],
            "sourceCells": {str(index): item["sourceCells"] for index, item in enumerate(records)},
            "sourceCellFormats": {str(index): item["sourceCellFormats"] for index, item in enumerate(records)},
            "sourceCellStates": {str(index): item["sourceCellStates"] for index, item in enumerate(records)},
            "sourceRowNumbers": [item["sourceRow"] for item in records],
            "sourceHeaders": headers, "derivedColumns": [], "overrideCells": {name: overrides[name]},
            "overrideFor": {"featureId": feature, "sourceColumn": name,
                            "sourceFormulaCells": spec["derivedMappings"][name]["sourceFormulaCells"]}}
    return feature, spec, obligations, extra_inputs


def build_source_input_plan(source_path, decisions, *, connection_id, folder_id) -> dict:
    """Build reviewed ordinary input architecture and source-cell reverse mappings."""
    if isinstance(decisions, (str, Path)):
        decisions = json.loads(Path(decisions).read_text(encoding="utf-8"))
    if not isinstance(decisions, dict) or not isinstance(decisions.get("features"), list):
        raise ValueError("source input decisions require features")
    if not connection_id or not folder_id:
        raise ValueError("target connection and folder scope are required")
    model = build_source_model(source_path)
    if decisions.get("sourceSha256") != model["sourceSha256"]:
        raise ValueError("source hash does not match reviewed decisions")
    features = decisions["features"]
    if not features:
        raise ValueError("at least one reviewed source feature is required")
    inputs, obligations = {}, []
    for decision in features:
        feature, spec, formulas, extras = _build_feature(model, decision, connection_id, source_path)
        if set([feature, *extras]) & set(inputs):
            raise ValueError("duplicate featureId or generated override input")
        inputs[feature] = spec
        inputs.update(extras)
        obligations.extend(formulas)
    record_cells = {cell for spec in inputs.values() if spec.get("sourceBinding") == "reviewed-source-records"
                    for row in spec["sourceCells"].values() for cell in row.values()}
    if record_cells:
        seen_bindings = set()
        for spec in inputs.values():
            for row in spec.get("sourceCells", {}).values():
                for cell in row.values():
                    if cell in record_cells and cell in seen_bindings:
                        raise ValueError(f"ambiguous duplicate source record cell binding: {cell}")
                    seen_bindings.add(cell)
    selected_formulas = {item["sourceCellId"] for item in obligations}
    for cell in model["cells"]:
        if cell["kind"] != "formula":
            continue
        source_cell = f"{cell['sheet']}!{cell['address']}"
        if source_cell not in selected_formulas:
            obligations.append({"featureId": None, "sourceCellId": source_cell,
                                "sourceRow": int(re.search(r"\d+$", cell["address"]).group()),
                                "sourceColumn": None, "expression": cell.get("expression"),
                                "dependencies": cell.get("dependencies") or [],
                                "mappingStatus": "live-calculation-required"})
    architecture = {"workflows": ["ordinary-source-inputs"],
                    "targetScope": {"connectionId": connection_id, "folderId": folder_id},
                    "mappedFeatures": sorted(inputs),
                    "featureToInput": {feature: feature for feature in sorted(inputs)},
                    "inputs": {feature: inputs[feature] for feature in sorted(inputs)},
                    "acceptanceChecks": [f"{feature}-edit-persist-restore" for feature in sorted(inputs)],
                    "nativeWiring": {"elements": [
                        {"id": f"input_{re.sub('[^A-Za-z0-9_]', '_', feature)}_{_digest(feature)[:8]}",
                         "kind": "input-table", "name": inputs[feature]["name"],
                         "inputMode": "explore", "source": {"kind": "empty", "connectionId": connection_id},
                         "columns": [{"id": f"c{index}", "name": column["name"], "type": column["type"]}
                                     for index, column in enumerate(inputs[feature]["columns"])]}
                        for feature in sorted(inputs)], "consumers": {}, "validations": {}}}
    logical = {"inputPlanVersion": PLAN_VERSION, "sourceSha256": model["sourceSha256"],
               "sourceName": model["sourceName"], "decisions": copy.deepcopy(decisions),
               "formulaObligations": sorted(obligations, key=lambda item: item["sourceCellId"]),
               "sharedInputArchitecture": architecture,
               "status": "REVIEWED_PLAN_PENDING_STAGING"}
    return {**logical, "inputPlanHash": _digest(logical)}


def validate_source_input_plan(plan, source_path=None) -> None:
    """Reject altered plans and, when supplied, changed source bytes or semantics."""
    if not isinstance(plan, dict) or plan.get("inputPlanVersion") not in {LEGACY_PLAN_VERSION, PLAN_VERSION}:
        raise ValueError("unsupported source input plan")
    legacy = plan["inputPlanVersion"] == LEGACY_PLAN_VERSION
    if legacy:
        if any("recordBindings" in feature for feature in plan.get("decisions", {}).get("features", [])):
            raise ValueError("version-one source input plan cannot contain record bindings")
        if any("sourceCellFormats" in spec or "sourceCellStates" in spec
               for spec in plan.get("sharedInputArchitecture", {}).get("inputs", {}).values()):
            raise ValueError("version-one source input plan cannot contain version-two metadata")
    logical = {key: value for key, value in plan.items() if key != "inputPlanHash"}
    if _digest(logical) != plan.get("inputPlanHash"):
        raise ValueError("source input plan hash is stale")
    if source_path is not None:
        if hashlib.sha256(Path(source_path).read_bytes()).hexdigest() != plan["sourceSha256"]:
            raise ValueError("source hash is stale")
        scope = plan["sharedInputArchitecture"]["targetScope"]
        rebuilt = build_source_input_plan(source_path, plan["decisions"],
                                          connection_id=scope["connectionId"], folder_id=scope["folderId"])
        if legacy:
            rebuilt["inputPlanVersion"] = LEGACY_PLAN_VERSION
            for spec in rebuilt["sharedInputArchitecture"]["inputs"].values():
                spec.pop("sourceCellFormats", None)
                spec.pop("sourceCellStates", None)
            rebuilt["inputPlanHash"] = _digest({key: value for key, value in rebuilt.items()
                                                 if key != "inputPlanHash"})
        if rebuilt != plan:
            raise ValueError("source input plan differs from source bytes")


def _normalized_datetime_key(value, name, date_only):
    if value is None:
        return value  # The seed planner reports missing business keys.
    if isinstance(value, str):
        if not _matches(value, "datetime"):
            raise ValueError(f"invalid datetime key: {name}")
        if "T" not in value:
            return value
        value = dt.datetime.fromisoformat(value.replace("Z", "+00:00"))
    if isinstance(value, dt.datetime):
        if value.tzinfo is not None:
            value = value.astimezone(dt.timezone.utc).replace(tzinfo=None)
        if date_only:
            if value.time() not in {dt.time.min, dt.time(12)}:
                raise ValueError(f"queried date-only key has unexpected time: {name}")
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    raise ValueError(f"invalid datetime key: {name}")


def plan_source_seed(plan, feature_id, current_rows, *, row_id_column, prior_seed=None) -> dict:
    """Delegate three-way, row-ID-aware reconciliation to the existing planner."""
    validate_source_input_plan(plan)
    inputs = plan["sharedInputArchitecture"]["inputs"]
    if feature_id not in inputs:
        raise ValueError("unknown source input feature")
    spec = inputs[feature_id]
    date_keys = {name for name in spec["keyColumns"]
                 if next(column["type"] for column in spec["columns"] if column["name"] == name) == "datetime"}
    date_only_keys = {name for name in date_keys if spec["rows"] and all(
        isinstance(row.get(name), str) and re.fullmatch(r"\d{4}-\d\d-\d\d", row[name])
        for row in spec["rows"])}

    def normalized_row(row):
        item = dict(row)
        for name in date_keys:
            if name in item:
                item[name] = _normalized_datetime_key(item[name], name, name in date_only_keys)
        return item

    def key_token(row):
        return tuple((type(row.get(name)), row.get(name)) for name in spec["keyColumns"])

    # Match canonical instants using the full business key, then restore the
    # reviewed spelling. Source rows and persisted seed keys remain unchanged.
    source_keys = {}
    if date_keys:
        for row in spec["rows"]:
            token = key_token(normalized_row(row))
            if token in source_keys:
                raise ValueError("duplicate source business key after datetime normalization")
            source_keys[token] = {name: row[name] for name in date_keys}
    normalized_current = []
    for current in current_rows:
        item = normalized_row(current)
        if date_keys:
            item.update(source_keys.get(key_token(item), {}))
        normalized_current.append(item)
    return plan_seed(spec["rows"], normalized_current, spec["keyColumns"], row_id_column,
                     prior_seed=prior_seed,
                     column_types={column["name"]: column["type"] for column in spec["columns"]})
