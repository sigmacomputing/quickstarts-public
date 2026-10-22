"""Explicit, read-only wide-to-long source table reshape with cell provenance."""
from __future__ import annotations

import datetime as dt
import json

from openpyxl.utils.cell import get_column_letter, range_boundaries

try:
    from .source_workbook import build_source_model
    from .source_inputs import _value
except ImportError:
    from source_workbook import build_source_model
    from source_inputs import _value


def _source_type(cell):
    if cell is None or cell.get("value") is None:
        return None
    _value(cell)
    value = cell["value"]
    if isinstance(value, dict) and value.get("type") in {"date", "datetime"}:
        return "datetime"
    if type(value) is bool:
        return "boolean"
    if type(value) in {int, float}:
        return "number"
    if type(value) is str:
        return "text"
    raise ValueError("unsupported unpivot source type")


def _axis_type(value):
    if type(value) is str:
        try:
            dt.date.fromisoformat(value)
            return "datetime"
        except ValueError:
            return "text"
    if type(value) is bool:
        return "boolean"
    if type(value) in {int, float}:
        return "number"
    raise ValueError("unsupported axis value")


def unpivot_source_table(source_path, *, sheet, table, key_columns, value_columns,
                         axis_name, value_name, axis_values=None) -> dict:
    """Unpivot selected columns; never infer calendar semantics from headers."""
    model = build_source_model(source_path)
    matches = [item for item in model["tables"] if item["sheet"] == sheet and item["name"] == table]
    if len(matches) != 1 or not matches[0]["bodyRange"]:
        raise ValueError("unknown or empty source table")
    spec = matches[0]
    headers = spec["headers"]
    if (not isinstance(key_columns, list) or not key_columns or not isinstance(value_columns, list)
            or not value_columns or len(key_columns) != len(set(key_columns))
            or len(value_columns) != len(set(value_columns))):
        raise ValueError("explicit unique key and value columns are required")
    if set(key_columns) & set(value_columns):
        raise ValueError("key/value column overlap")
    if not set(key_columns + value_columns) <= set(headers):
        raise ValueError("unknown key or value column")
    if not isinstance(axis_name, str) or not axis_name or not isinstance(value_name, str) or not value_name:
        raise ValueError("axis and value names are required")
    if len({axis_name, value_name, *key_columns}) != len(key_columns) + 2:
        raise ValueError("axis/value names overlap key columns")
    if axis_values is None:
        if any(str(name).strip()[:3].lower() in {"jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"}
               for name in value_columns):
            raise ValueError("ambiguous month axis requires an explicit year mapping")
        axis_values = {name: name for name in value_columns}
    if not isinstance(axis_values, dict) or set(axis_values) != set(value_columns):
        raise ValueError("axis mapping must exactly cover value columns")
    tokens = [json.dumps([type(axis_values[name]).__name__, axis_values[name]], sort_keys=True)
              for name in value_columns]
    if len(set(tokens)) != len(tokens):
        raise ValueError("duplicate resulting axis values")
    cells = {(item["sheet"], item["address"]): item for item in model["cells"]}
    left, top, _right, bottom = range_boundaries(spec["bodyRange"])
    rows, source_cells, source_formats, source_states, row_numbers, seen = [], {}, {}, {}, [], set()
    column_types = {name: set() for name in key_columns + [axis_name, value_name]}
    for row_number in range(top, bottom + 1):
        base, base_cells, base_formats, base_states = {}, {}, {}, {}
        for name in key_columns:
            address = f"{get_column_letter(left + headers.index(name))}{row_number}"
            cell = cells.get((sheet, address))
            if cell is None or cell["kind"] != "literal":
                raise ValueError("unpivot key requires a literal source value")
            value = _value(cell)
            column_types[name].add(_source_type(cell))
            if value is None or value == "":
                raise ValueError("unpivot key is missing")
            base[name] = value; base_cells[name] = f"{sheet}!{address}"
            base_formats[name] = cell.get("numberFormat")
            base_states[name] = "empty-text" if value == "" else "literal"
        for name in value_columns:
            address = f"{get_column_letter(left + headers.index(name))}{row_number}"
            cell = cells.get((sheet, address))
            if cell and cell["kind"] == "formula":
                raise ValueError("formula measure cannot be frozen as input")
            value = _value(cell) if cell else None
            column_types[value_name].add(_source_type(cell))
            axis = axis_values[name]
            if axis is None or not isinstance(axis, (str, int, float, bool)):
                raise ValueError("unsupported axis value")
            row = {**base, axis_name: axis, value_name: value}
            column_types[axis_name].add(_axis_type(axis))
            key = json.dumps([[type(row[key]).__name__, row[key]] for key in key_columns + [axis_name]],
                             sort_keys=True, ensure_ascii=False)
            if key in seen:
                raise ValueError("duplicate resulting keys")
            seen.add(key)
            header_address = f"{get_column_letter(left + headers.index(name))}{top-1}"
            header_cell = cells.get((sheet, header_address))
            source_cells[str(len(rows))] = {**base_cells, axis_name: f"{sheet}!{header_address}",
                                             value_name: f"{sheet}!{address}"}
            source_formats[str(len(rows))] = {**base_formats,
                                              axis_name: header_cell.get("numberFormat") if header_cell else None,
                                              value_name: cell.get("numberFormat") if cell else None}
            source_states[str(len(rows))] = {**base_states,
                                             axis_name: "reviewed-axis-mapping",
                                             value_name: ("blank" if cell is None else
                                                          "empty-text" if value == "" else "literal")}
            row_numbers.append(row_number)
            rows.append(row)
    resolved_types = {}
    for name, observed in column_types.items():
        concrete = observed - {None}
        if len(concrete) != 1:
            raise ValueError(f"ambiguous or inconsistent unpivot column type: {name}")
        resolved_types[name] = concrete.pop()
    return {"reshapeVersion": 1, "sourceSha256": model["sourceSha256"],
            "sourceBinding": f"{sheet}!{spec['range']}", "keyColumns": key_columns + [axis_name],
            "axisMapping": axis_values, "columnTypes": resolved_types,
            "rows": rows, "sourceCells": source_cells,
            "sourceCellFormats": source_formats, "sourceCellStates": source_states,
            "sourceRowNumbers": row_numbers}
