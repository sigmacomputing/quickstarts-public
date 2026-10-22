#!/usr/bin/env python3
"""Build the reviewed XLSX companion. Does not compile, attach, or execute VBA."""
import argparse
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]


def build_fixture(path):
    model=json.loads((ROOT/'tests/fixtures/task5-controlled-model.json').read_text())
    wb=Workbook();wb.remove(wb.active)
    for title in ('Headcount','Assumptions','Alloc','Actuals','Forecast Inputs','Line Definitions','Forecast','P&L'):
        wb.create_sheet(title)
    def table(sheet, headers, rows, name, start=1):
        ws=wb[sheet]
        for c,h in enumerate(headers,1):ws.cell(start,c,h)
        for r,row in enumerate(rows,start+1):
            for c,value in enumerate(row,1):ws.cell(r,c,value)
        from openpyxl.utils import get_column_letter
        t=Table(displayName=name,ref=f'A{start}:{get_column_letter(len(headers))}{start+len(rows)}')
        t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(t)
    table('Headcount',['Department','Headcount'],list(zip(['Eng','Sales','G&A','Support'],[40,30,20,10])),'ControlledHeadcount')
    table('Alloc',['Department','Allocation'],[(d,-999) for d in ['Eng','Sales','G&A','Support']],'ControlledAllocation')
    wb['Assumptions']['A5']='Subscription monthly growth';wb['Assumptions']['B5']=.032
    wb['Assumptions']['A12']='Overhead';wb['Assumptions']['B12']=100000
    table('Actuals',['Line Item','Month','Section','Amount'],
          [[r[k] for k in ('Line Item','Month','Section','Amount')] for r in model['actuals']],'ControlledActuals')
    fields=['Line Item','Method','Monthly Growth','Jul-26','Aug-26','Sep-26','Oct-26','Nov-26','Dec-26']
    table('Forecast Inputs',fields,[[r[k] for k in fields] for r in model['inputs']],'ControlledForecastInputs',5)
    wb['Forecast Inputs']['C6']="=Assumptions!B5"
    table('Line Definitions',['Line Item','Section'],[[r['Line Item'],r['Section']] for r in model['definitions']],'ControlledDefinitions')
    ws=wb['Forecast'];ws.append(['Line Item','Section','Method','June Base',*fields[3:]])
    for i,row in enumerate(model['inputs'],2):
        source=i+4
        ws.cell(i,1,row['Line Item']);ws.cell(i,2,model['definitions'][i-2]['Section'])
        ws.cell(i,3,f"='Forecast Inputs'!B{source}")
        ws.cell(i,4,f'=SUMIFS(Actuals!D2:D163,Actuals!A2:A163,A{i},Actuals!B2:B163,"Jun-26")')
        for n,col in enumerate('EFGHIJ',1):
            manual='DEFGHI'[n-1]
            ws[f'{col}{i}']=f'=IF(C{i}="manual",\'Forecast Inputs\'!{manual}{source},IF(C{i}="flat",D{i},D{i}*(1+\'Forecast Inputs\'!C{source})^{n}))'
    # Explicit net-income identity at the legacy objective address.
    pnl=wb['P&L'];pnl['M33']='Annual Net Income'
    actual='SUMIFS(Actuals!D2:D163,Actuals!C2:C163,"{section}")'
    parts=[]
    for section,sign in [('Revenue',1),('Cost of Revenue',-1),('Operating Expenses',-1),('Below the Line',-1)]:
        term=actual.format(section=section)+'+'+'+'.join(f'SUMIFS(Forecast!{c}2:{c}13,Forecast!B2:B13,"{section}")' for c in 'EFGHIJ')
        parts.append(('+' if sign==1 else '-')+'('+term+')')
    pnl['N33']='='+''.join(parts).lstrip('+')
    for sheet in wb:
        sheet.freeze_panes='A2'
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width=26
    wb.save(path)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--out',type=Path,required=True)
    args=parser.parse_args();build_fixture(args.out)
